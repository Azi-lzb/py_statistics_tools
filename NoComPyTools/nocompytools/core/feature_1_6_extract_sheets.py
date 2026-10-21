from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string

from .io_excel import clear_sheet_cache, list_sheet_names, read_sheet_2d
from .logger import get_logger


_INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")
_CELL_RE = re.compile(r"^([A-Za-z]+)(\d+)$")
_CELL_RANGE_RE = re.compile(r"^([A-Za-z]+)(\d+):([A-Za-z]+)(\d+)$")
_COL_RANGE_RE = re.compile(r"^([A-Za-z]+):([A-Za-z]+)$")
_ROW_RANGE_RE = re.compile(r"^(\d+):(\d+)$")
_FULL_AREA_WORDS = {"", "整表", "全部", "全表"}


def _normalize_range_spec(text: str) -> str:
    if not text:
        return ""
    s = text.replace(" ", "").replace("\t", "")
    for ch in ("，", "、", "；", ";"):
        s = s.replace(ch, ",")
    s = s.replace("：", ":")
    return s


def _token_to_index(token: str, is_row: bool) -> int:
    token = token.strip()
    if not token:
        return 0
    if token.isdigit():
        return int(token)
    if is_row:
        return 0
    try:
        return column_index_from_string(token.upper())
    except Exception:
        return 0


def _parse_indexes(spec: str, max_value: int, is_row: bool) -> list[int]:
    norm = _normalize_range_spec(spec)
    if not norm or max_value <= 0:
        return list(range(1, max_value + 1))
    result: list[int] = []
    for seg in norm.split(","):
        seg = seg.strip()
        if not seg:
            continue
        if ":" in seg:
            parts = seg.split(":", 1)
            start = _token_to_index(parts[0], is_row)
            end = _token_to_index(parts[1], is_row)
            if start <= 0 or end <= 0:
                continue
            if start > end:
                start, end = end, start
            start = max(start, 1)
            end = min(end, max_value)
            result.extend(range(start, end + 1))
        else:
            idx = _token_to_index(seg, is_row)
            if 1 <= idx <= max_value:
                result.append(idx)
    if not result:
        return list(range(1, max_value + 1))
    # 保序去重
    seen: set[int] = set()
    out: list[int] = []
    for v in result:
        if v not in seen:
            seen.add(v)
            out.append(v)
    return out


def _build_exact_name(keywords: list[str]) -> str:
    return "-".join(k for k in keywords if k.strip())


def _sheet_matches_all(name: str, keywords: list[str]) -> bool:
    name_low = name.lower()
    for k in keywords:
        k = k.strip()
        if not k:
            continue
        if k.lower() not in name_low:
            return False
    return True


def _find_matched_sheets(wb, keywords: list[str]) -> list:
    matched = []
    seen: set[str] = set()
    exact = _build_exact_name(keywords)
    if exact and exact in wb.sheetnames:
        matched.append(wb[exact])
        seen.add(exact)
    for ws in wb.worksheets:
        if ws.title in seen:
            continue
        if _sheet_matches_all(ws.title, keywords):
            matched.append(ws)
            seen.add(ws.title)
    return matched


def _split_keyword_groups(text: str) -> list[list[str]]:
    if not text:
        return []
    groups = []
    normalized = str(text).replace("；", ";")
    for group_text in normalized.split(";"):
        options = []
        for option in re.split(r"[,，|]", group_text):
            option = option.strip()
            if option:
                options.append(option)
        if options:
            groups.append(options)
    return groups


def _name_matches_keyword_options(name: str, keyword_text: str) -> bool:
    groups = _split_keyword_groups(keyword_text)
    if not groups:
        return True
    name_low = name.lower()
    return all(any(option.lower() in name_low for option in group) for group in groups)


def _workbook_matches(src: Path, task: dict) -> bool:
    return _name_matches_keyword_options(src.name, str(task.get("workbook_keywords") or ""))


def _name_matches_exclude(name: str, task: dict) -> bool:
    exclude_keywords = str(task.get("exclude_keywords") or "")
    if not exclude_keywords.strip():
        return False
    return _name_matches_keyword_options(name, exclude_keywords)


def _find_matched_sheets_for_task(wb, task: dict) -> list:
    sheet_keywords = str(task.get("sheet_keywords") or "")
    matched = []
    seen: set[str] = set()
    for ws in wb.worksheets:
        if ws.title in seen:
            continue
        if _name_matches_exclude(ws.title, task):
            continue
        if _name_matches_keyword_options(ws.title, sheet_keywords):
            matched.append(ws)
            seen.add(ws.title)
    return matched


def _find_matched_sheet_names_for_task(sheet_names: list[str], task: dict) -> list[str]:
    sheet_keywords = str(task.get("sheet_keywords") or "")
    return [
        name for name in sheet_names
        if not _name_matches_exclude(name, task)
        and _name_matches_keyword_options(name, sheet_keywords)
    ]


def _sanitize_sheet_name(name: str) -> str:
    cleaned = _INVALID_SHEET_CHARS.sub("_", name)
    return cleaned[:31] if len(cleaned) > 31 else cleaned


def _unique_sheet_name(wb, base: str) -> str:
    base = _sanitize_sheet_name(base)
    if base not in wb.sheetnames:
        return base
    root = base[:28] if len(base) > 28 else base
    for i in range(2, 101):
        cand = f"{root}_{i}"
        if cand not in wb.sheetnames:
            return cand
    return f"{root}_{datetime.now().strftime('%H%M%S')}"


def _read_ws_values(ws) -> list[list]:
    max_r = ws.max_row or 0
    max_c = ws.max_column or 0
    if max_r == 0 or max_c == 0:
        return []
    return [[ws.cell(row=r, column=c).value for c in range(1, max_c + 1)]
            for r in range(1, max_r + 1)]


def _extract_values(data: list[list], rows_spec: str, cols_spec: str,
                    full_extract: bool) -> list[list]:
    if not data:
        return []
    max_r = len(data)
    max_c = max((len(row) for row in data), default=0)
    if max_c == 0:
        return []

    if full_extract or (not rows_spec and not cols_spec):
        rows_idx = list(range(1, max_r + 1))
        cols_idx = list(range(1, max_c + 1))
    else:
        rows_idx = _parse_indexes(rows_spec, max_r, is_row=True)
        cols_idx = _parse_indexes(cols_spec, max_c, is_row=False)

    out: list[list] = []
    for r in rows_idx:
        src_row = data[r - 1] if r - 1 < len(data) else []
        new_row = [src_row[c - 1] if c - 1 < len(src_row) else None for c in cols_idx]
        out.append(new_row)
    return out


def _is_full_area_spec(area_spec: str) -> bool:
    return str(area_spec or "").strip() in _FULL_AREA_WORDS


def _split_area_specs(area_spec: str) -> list[str]:
    text = str(area_spec or "").strip()
    if _is_full_area_spec(text):
        return [""]
    parts = [part.strip() for part in text.replace("；", ";").split(";") if part.strip()]
    return parts or [""]


def _cell_to_row_col(token: str) -> tuple[int, int] | None:
    m = _CELL_RE.fullmatch(token.strip())
    if not m:
        return None
    col = column_index_from_string(m.group(1).upper())
    row = int(m.group(2))
    return row, col


def _parse_area_indexes(area_spec: str, max_row: int, max_col: int) -> tuple[list[int], list[int]] | None:
    area = str(area_spec or "").strip().replace(" ", "").replace("\t", "").replace("：", ":")
    if _is_full_area_spec(area):
        return list(range(1, max_row + 1)), list(range(1, max_col + 1))
    if max_row <= 0 or max_col <= 0:
        return [], []

    cell_range = _CELL_RANGE_RE.fullmatch(area)
    if cell_range:
        start_col = column_index_from_string(cell_range.group(1).upper())
        start_row = int(cell_range.group(2))
        end_col = column_index_from_string(cell_range.group(3).upper())
        end_row = int(cell_range.group(4))
        row_start, row_end = sorted((start_row, end_row))
        col_start, col_end = sorted((start_col, end_col))
        rows = range(max(row_start, 1), min(row_end, max_row) + 1)
        cols = range(max(col_start, 1), min(col_end, max_col) + 1)
        return list(rows), list(cols)

    col_range = _COL_RANGE_RE.fullmatch(area)
    if col_range:
        start_col = column_index_from_string(col_range.group(1).upper())
        end_col = column_index_from_string(col_range.group(2).upper())
        col_start, col_end = sorted((start_col, end_col))
        cols = range(max(col_start, 1), min(col_end, max_col) + 1)
        return list(range(1, max_row + 1)), list(cols)

    row_range = _ROW_RANGE_RE.fullmatch(area)
    if row_range:
        start_row, end_row = sorted((int(row_range.group(1)), int(row_range.group(2))))
        rows = range(max(start_row, 1), min(end_row, max_row) + 1)
        return list(rows), list(range(1, max_col + 1))

    cell = _cell_to_row_col(area)
    if cell:
        row, col = cell
        if 1 <= row <= max_row and 1 <= col <= max_col:
            return [row], [col]
        return [], []

    return None


def _extract_area_values(data: list[list], area_spec: str) -> list[list] | None:
    if not data:
        return []
    max_row = len(data)
    max_col = max((len(row) for row in data), default=0)
    if max_col == 0:
        return []
    parsed = _parse_area_indexes(area_spec, max_row, max_col)
    if parsed is None:
        return None
    rows_idx, cols_idx = parsed
    values: list[list] = []
    for r in rows_idx:
        src_row = data[r - 1] if r - 1 < len(data) else []
        values.append([src_row[c - 1] if c - 1 < len(src_row) else None for c in cols_idx])
    return values


def _stitch_area_values(area_values: list[list[list]], stitch_mode: str) -> list[list]:
    if not area_values:
        return []
    mode = str(stitch_mode or "").strip()
    if mode == "纵向追加":
        out: list[list] = []
        for values in area_values:
            out.extend(values)
        return out

    max_rows = max((len(values) for values in area_values), default=0)
    widths = [max((len(row) for row in values), default=0) for values in area_values]
    out = []
    for row_idx in range(max_rows):
        row: list = []
        for values, width in zip(area_values, widths):
            if row_idx < len(values):
                src_row = values[row_idx]
                row.extend(src_row + [None] * (width - len(src_row)))
            else:
                row.extend([None] * width)
        out.append(row)
    return out


def _extract_configured_values(data: list[list], area_spec: str, stitch_mode: str) -> list[list] | None:
    parts = _split_area_specs(area_spec)
    area_values: list[list[list]] = []
    for part in parts:
        values = _extract_area_values(data, part)
        if values is None:
            return None
        area_values.append(values)
    if len(area_values) == 1:
        return area_values[0]
    return _stitch_area_values(area_values, stitch_mode)


def _compact(values: list[list]) -> list[list]:
    """去除全空行与全空列。"""
    def _is_empty(v) -> bool:
        return v is None or (isinstance(v, str) and v.strip() == "")

    rows_kept = [row for row in values if any(not _is_empty(v) for v in row)]
    if not rows_kept:
        return []
    ncol = max(len(r) for r in rows_kept)
    # padding
    rows_kept = [r + [None] * (ncol - len(r)) for r in rows_kept]
    col_has_data = [any(not _is_empty(r[c]) for r in rows_kept) for c in range(ncol)]
    return [[r[c] for c in range(ncol) if col_has_data[c]] for r in rows_kept]


def _write_values(ws_out, values: list[list]) -> None:
    for i, row in enumerate(values, start=1):
        for j, v in enumerate(row, start=1):
            if v is not None:
                ws_out.cell(row=i, column=j, value=v)


def _is_blank_spec(spec: str) -> bool:
    return str(spec or "").strip() == ""


def _is_full_sheet_task(task: dict) -> bool:
    return _is_full_area_spec(str(task.get("area_spec") or ""))


def _output_path_for_task(task: dict, output_dir: Path, default_name: str) -> Path:
    folder = str(task.get("output_folder") or "").strip()
    if folder:
        base_dir = Path(folder)
        if not base_dir.is_absolute():
            base_dir = output_dir / base_dir
    else:
        base_dir = output_dir
    name = str(task.get("output_name") or "").strip() or default_name
    path = Path(name)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    if path.is_absolute():
        return path
    return base_dir / path


def _remove_default_blank(wb) -> None:
    if len(wb.sheetnames) > 1:
        first = wb.worksheets[0]
        if first.max_row == 1 and first.max_column == 1 and first.cell(1, 1).value is None:
            wb.remove(first)


def run_extract_sheets(target_wbs: Iterable[Path], tasks: list[dict],
                       output_dir: Path, log_dir: Path,
                       append_existing: bool = False,
                       default_output_name: str | None = None) -> dict:
    logger = get_logger("feature_1_6_extract_sheets", log_dir)
    if not tasks:
        logger.warning("'工作表提取'配置无启用任务")
        return {"files": 0, "extracted_sheets": 0, "output_files": 0,
                "failed_files": 0, "paths": []}

    ts_default = default_output_name or f"提取结果_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    output_dir.mkdir(parents=True, exist_ok=True)

    out_wbs: dict[Path, Workbook] = {}
    stat = {"files": 0, "extracted_sheets": 0, "output_files": 0,
            "failed_files": 0, "paths": []}

    for src in target_wbs:
        try:
            logger.info(f"打开源文件: {src}")
            sheet_names = list_sheet_names(src)
            if not sheet_names:
                raise ValueError("未读取到工作表")
        except Exception as e:
            stat["failed_files"] += 1
            logger.error(f"打开失败: {src} - {e}")
            continue

        stat["files"] += 1
        try:
            for task in tasks:
                if not _workbook_matches(src, task):
                    continue
                if _name_matches_exclude(src.name, task):
                    logger.info(f"  [skip] {src.name} 命中排除关键字")
                    continue
                matched_names = _find_matched_sheet_names_for_task(sheet_names, task)
                if not matched_names:
                    continue
                full_sheet = _is_full_sheet_task(task)

                path = _output_path_for_task(task, output_dir, ts_default)
                if path not in out_wbs:
                    if append_existing and path.exists():
                        out_wbs[path] = load_workbook(path)
                    else:
                        out_wbs[path] = Workbook()
                wb_out = out_wbs[path]

                for source_sheet_name in matched_names:
                    sheet_title = _unique_sheet_name(wb_out, source_sheet_name)
                    ws_out = wb_out.create_sheet(title=sheet_title)
                    try:
                        frame = read_sheet_2d(src, source_sheet_name)
                        data = frame.astype(object).where(frame.notna(), None).values.tolist()
                    except Exception as e:
                        wb_out.remove(ws_out)
                        stat["failed_files"] += 1
                        logger.error(f"读取失败: {src} | {source_sheet_name} - {e}")
                        continue

                    if full_sheet:
                        if not any(value is not None for row in data for value in row):
                            wb_out.remove(ws_out)
                            logger.info(f"  [skip] {src.name} | {source_sheet_name} 无数据")
                            continue
                        _write_values(ws_out, data)
                    else:
                        values = _extract_configured_values(
                            data, task.get("area_spec", ""), task.get("stitch_mode", ""),
                        )
                        if values is None:
                            wb_out.remove(ws_out)
                            logger.warning(
                                f"  [skip] 配置行{task.get('row_no')} 提取区域无效: "
                                f"{task.get('area_spec')}"
                            )
                            continue
                        values = _compact(values)
                        if not values:
                            wb_out.remove(ws_out)
                            logger.info(f"  [skip] {src.name} | {source_sheet_name} 无数据")
                            continue
                        _write_values(ws_out, values)
                    stat["extracted_sheets"] += 1
                    logger.info(
                        f"  提取: {src.name} | {source_sheet_name} -> {path.name}[{sheet_title}]"
                    )
        finally:
            clear_sheet_cache()

    for path, wb_out in out_wbs.items():
        _remove_default_blank(wb_out)
        if (len(wb_out.worksheets) == 1
                and wb_out.active.max_row == 1
                and wb_out.active.max_column == 1
                and wb_out.active.cell(1, 1).value is None):
            wb_out.close()
            continue
        path.parent.mkdir(parents=True, exist_ok=True)
        wb_out.save(path)
        wb_out.close()
        stat["output_files"] += 1
        stat["paths"].append(str(path))
        logger.info(f"已保存: {path}")

    return stat
