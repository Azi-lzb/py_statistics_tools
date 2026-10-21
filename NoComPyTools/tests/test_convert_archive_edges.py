from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from nocompytools.core.config_xlsx import ARCHIVE_TYPE_CONFIG_COLS, SHEET_ARCHIVE_TYPE_CONFIG
from nocompytools.core import convert_33_34_35_37 as convert


def _archive_config(path: Path, include: str = "", exclude: str = "") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_ARCHIVE_TYPE_CONFIG
    ws.append(ARCHIVE_TYPE_CONFIG_COLS)
    ws.append(["Y", "贷款", "贷款类", "enabled", include, exclude])
    ws.append(["N", "禁用", "禁用类", "disabled", "", ""])
    wb.save(path)
    wb.close()
    return path


def _touch(path: Path, text: str = "x") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def test_archive_text_and_date_helpers_cover_supported_separators() -> None:
    assert convert._norm(None) == ""
    assert convert._norm(float("nan")) == ""
    assert convert._norm(pd.NA) == ""
    assert convert._norm("  x  ") == "x"
    assert convert._split_tokens(" a；b,c ;  ") == ["a", "b", "c"]
    assert convert._norm(["x"]) == "['x']"
    assert convert._is_enabled("启用") and convert._is_enabled("TRUE")
    assert not convert._is_enabled("否")
    assert convert._parse_archive_exts(".XLSX； docx, pdf") == {".xlsx", ".docx", ".pdf"}
    assert convert._parse_archive_exts("") == set()
    assert convert._parse_archive_exts(" ; . ") == set()
    assert convert._archive_date_from_name("monthly_2026-2-3_report.xlsx") == "202602"
    assert convert._archive_date_from_name("20261231_summary.xlsx") == "202612"
    assert convert._archive_date_from_name("2026年7月报告.xlsx") == "202607"
    assert convert._archive_date_from_name("202614_report.xlsx") == ""
    assert convert._safe_folder_name(" <贷款>/2026. ") == "_贷款__2026"
    assert convert._archive_type_from_name("普惠贷款表.xlsx", [("贷款", "贷款"), ("普惠", "普惠")]) == "贷款"
    assert convert._archive_type_from_name("无匹配.xlsx", []) == ""


def test_archive_candidate_collection_deduplicates_and_ignores_lock_files(tmp_path: Path) -> None:
    root = tmp_path / "source"
    first = _touch(root / "nested" / "one.xlsx")
    _touch(root / "~$locked.xlsx")
    standalone = _touch(tmp_path / "two.docx")
    missing = tmp_path / "missing.xlsx"

    files, invalid = convert._iter_archive_candidates([root, first, standalone, missing])

    assert set(files) == {first, standalone}
    assert invalid == [missing]


def test_archive_filter_load_handles_missing_corrupt_and_configured_workbooks(tmp_path: Path) -> None:
    missing = tmp_path / "missing.xlsx"
    corrupt = _touch(tmp_path / "corrupt.xlsx", "not a workbook")
    empty = tmp_path / "empty.xlsx"
    Workbook().save(empty)
    cfg = _archive_config(tmp_path / "config.xlsx", include="xlsx;docx", exclude="docx")
    default = (set(), set(), "未配置后缀筛选，按全部文件处理")

    assert convert._load_archive_ext_filter(None) == default
    assert convert._load_archive_ext_filter(missing) == default
    assert convert._load_archive_ext_filter(corrupt) == default
    assert convert._load_archive_ext_filter(empty) == default
    assert convert._load_archive_ext_filter(cfg) == ({".xlsx", ".docx"}, {".docx"}, "")
    assert convert._archive_ext_skip_status(Path("a.docx"), set(), {".docx"}) == "跳过：后缀被排除"
    assert convert._archive_ext_skip_status(Path("a.pdf"), {".xlsx"}, set()) == "跳过：后缀不在归档范围"
    assert convert._archive_ext_skip_status(Path("a.xlsx"), {".xlsx"}, set()) == ""


def test_archive_preview_combines_invalid_paths_duplicate_names_and_extension_filters(tmp_path: Path) -> None:
    cfg = _archive_config(tmp_path / "config.xlsx", include="xlsx", exclude="")
    first = _touch(tmp_path / "a" / "20260601_贷款表.xlsx")
    second = _touch(tmp_path / "b" / "20260601_贷款表.xlsx")
    excluded = _touch(tmp_path / "b" / "20260601_install.exe")
    missing = tmp_path / "missing.xlsx"

    stat = convert.run_archive_by_date_preview(
        [first, second, excluded, missing], tmp_path / "archive", tmp_path / "out", tmp_path / "logs", cfg
    )

    plan = pd.read_excel(stat["output"], sheet_name="归档计划", dtype=str).fillna("")
    assert stat["files"] == 3 and stat["invalid"] == 1 and stat["skipped_ext"] == 1
    assert set(plan["状态"]) == {"待归档", "跳过：后缀不在归档范围", "跳过：非文件"}
    targets = plan.loc[plan["状态"] == "待归档", "目标文件路径"].tolist()
    assert len(set(targets)) == 2
    assert any(path.endswith("_1.xlsx") for path in targets)


def test_archive_plan_copy_validates_columns_and_copies_repeated_targets(tmp_path: Path) -> None:
    invalid_plan = tmp_path / "invalid.xlsx"
    pd.DataFrame([{"文件名": "x.xlsx"}]).to_excel(invalid_plan, sheet_name="归档计划", index=False)
    with pytest.raises(ValueError, match="缺少必要列"):
        convert.run_archive_plan_copy(invalid_plan, tmp_path / "out", tmp_path / "logs")

    source1 = _touch(tmp_path / "source1.xlsx", "first")
    source2 = _touch(tmp_path / "source2.xlsx", "second")
    target = tmp_path / "archive" / "loan.xlsx"
    plan = tmp_path / "plan.xlsx"
    pd.DataFrame([
        {"源文件路径": str(source1), "目标文件路径": str(target), "状态": "待归档"},
        {"源文件路径": str(source2), "目标文件路径": str(target), "状态": "未识别类型"},
        {"源文件路径": str(source2), "目标文件路径": "", "状态": "待归档"},
    ]).to_excel(plan, sheet_name="归档计划", index=False)

    result = convert.run_archive_plan_copy(plan, tmp_path / "out", tmp_path / "logs")

    assert (result["copied"], result["skip"], result["failed"]) == (2, 1, 0)
    assert target.read_text(encoding="utf-8") == "first"
    assert (target.parent / "loan_1.xlsx").read_text(encoding="utf-8") == "second"


def test_archive_copy_records_filesystem_failures(tmp_path: Path, monkeypatch) -> None:
    source = _touch(tmp_path / "source.xlsx")
    row = {
        "源文件路径": str(source),
        "目标文件路径": str(tmp_path / "archive" / "copy.xlsx"),
        "状态": "待归档",
    }
    monkeypatch.setattr(convert.shutil, "copy2", lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("disk full")))

    result = convert._copy_archive_rows([row], tmp_path / "out", tmp_path / "logs")

    assert result["copied"] == 0 and result["failed"] == 1 and result["skip"] == 0
    output = pd.read_excel(result["output"], sheet_name="归档执行结果", dtype=str).fillna("")
    assert output.loc[0, "执行状态"] == "失败：disk full"


def test_archive_copy_suffixes_repeated_targets_and_records_the_note(tmp_path: Path) -> None:
    first = _touch(tmp_path / "first.xlsx", "first")
    second = _touch(tmp_path / "second.xlsx", "second")
    target = tmp_path / "archive" / "copy.xlsx"
    result = convert._copy_archive_rows([
        {"源文件路径": str(first), "目标文件路径": str(target), "状态": "待归档"},
        {"源文件路径": str(second), "目标文件路径": str(target), "状态": "未识别日期"},
    ], tmp_path / "out", tmp_path / "logs")

    assert result["copied"] == 2 and result["failed"] == 0
    output = pd.read_excel(result["output"], sheet_name="归档执行结果", dtype=str).fillna("")
    assert output.loc[1, "实际目标路径"].endswith("copy_1.xlsx")
    assert output.loc[1, "备注"] == "目标同名：已自动后缀"


def test_archive_plan_copy_records_copy_failure(tmp_path: Path, monkeypatch) -> None:
    source = _touch(tmp_path / "source.xlsx")
    plan = tmp_path / "plan.xlsx"
    pd.DataFrame([{
        "源文件路径": str(source), "目标文件路径": str(tmp_path / "target.xlsx"), "状态": "待归档"
    }]).to_excel(plan, sheet_name="归档计划", index=False)
    monkeypatch.setattr(convert.shutil, "copy2", lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("read only")))

    result = convert.run_archive_plan_copy(plan, tmp_path / "out", tmp_path / "logs")

    assert result["copied"] == 0 and result["failed"] == 1
    output = pd.read_excel(result["output"], sheet_name="归档执行结果", dtype=str).fillna("")
    assert output.loc[0, "执行状态"] == "失败：read only"


def test_submission_check_ignores_invalid_files_and_requires_an_enabled_rule(tmp_path: Path) -> None:
    config = tmp_path / "submission.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "提交检查配置"
    ws.append(["是否启用", "日期关键词", "县区关键词", "机构关键词", "备用关键词", "命中次数", "命中的文件名", "备注"])
    ws.append(["Y", "2026,2025", "惠城；惠阳", "建行", "存款", "", "", "rule"])
    ws.append(["N", "", "", "", "", "", "", "disabled"])
    wb.save(config)
    wb.close()
    matching = _touch(tmp_path / "2026_惠城_建行_存款.xlsx")
    partial = _touch(tmp_path / "2026_建行_存款.xlsx")

    assert convert._file_matches_submission_rule(matching.name, {
        "groups": [["2026", "2025"], ["惠城", "惠阳"], ["建行"], ["存款"]]
    })
    assert not convert._file_matches_submission_rule(partial.name, {
        "groups": [["2026"], ["惠城", "惠阳"], ["建行"], ["存款"]]
    })
    result = convert.run_submission_check_48(
        config, [matching, partial, tmp_path / "missing.xlsx", tmp_path], tmp_path / "out", tmp_path / "logs"
    )
    assert result["files"] == 2 and result["unmatched"] == 1

    wb = load_workbook(config)
    wb.active["A2"] = "N"
    wb.save(config)
    wb.close()
    with pytest.raises(ValueError, match="没有启用的检查项"):
        convert.run_submission_check_48(config, [], tmp_path / "out", tmp_path / "logs")
