from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
import xlwt
from openpyxl import Workbook

from nocompytools.core.config_xlsx import CONFIG_RENAME_COLS, FILE_SORT_CONFIG_COLS, SHEET_CONFIG_RENAME, SHEET_FILE_SORT_CONFIG
from nocompytools.core import convert_33_34_35_37 as convert


def _config(
    path: Path,
    rename_rows: list[dict[str, object]] | None = None,
    sort_rows: list[tuple[object, object, object, object]] | None = None,
) -> Path:
    wb = Workbook()
    rename = wb.active
    rename.title = SHEET_CONFIG_RENAME
    rename.append(CONFIG_RENAME_COLS)
    for row in rename_rows or []:
        rename.append([row.get(column, "") for column in CONFIG_RENAME_COLS])
    sorting = wb.create_sheet(SHEET_FILE_SORT_CONFIG)
    sorting.append(FILE_SORT_CONFIG_COLS)
    for row in sort_rows or []:
        sorting.append(row)
    wb.save(path)
    wb.close()
    return path


def _workbook(path: Path, cells: dict[str, object] | None = None, sheets: list[str] | None = None) -> None:
    wb = Workbook()
    wb.active.title = (sheets or ["sheet1"])[0]
    for title in (sheets or ["sheet1"])[1:]:
        wb.create_sheet(title)
    for address, value in (cells or {}).items():
        wb.worksheets[0][address] = value
    wb.save(path)
    wb.close()


def _touch(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"fixture")
    return path


def test_sort_prefix_helpers_and_rule_loader_filter_rows(tmp_path: Path) -> None:
    assert convert._sort_prefix_digits("001.0") == "1"
    assert convert._sort_prefix_digits("A-02B") == "02"
    assert convert._sort_prefix_digits("") == ""
    assert convert._sort_prefix_width("001.0") == 3
    assert convert._sort_prefix_width("A-02B") == 2
    assert convert._normalize_sort_prefix("2", 3) == "002"
    assert convert._normalize_sort_prefix("", 3) == ""
    assert convert._strip_sort_prefix("12 old-name.xlsx") == "old-name.xlsx"
    assert convert._strip_sort_prefix("1234_old-name.xlsx") == "1234_old-name.xlsx"

    cfg = _config(tmp_path / "config.xlsx", sort_rows=[
        ("Y", "1", " first；alpha ", "one"),
        ("是", "20.0", "second", "two"),
        ("Y", "", "missing prefix", "skip"),
        ("N", "300", "disabled", "skip"),
        ("Y", "", "", "empty"),
    ])
    rules = convert._load_file_sort_rules(cfg)
    assert [rule["prefix"] for rule in rules] == ["01", "20"]
    assert rules[0]["keywords"] == ["first", "alpha"]
    assert convert._match_file_sort_rule("ALPHA.xlsx", rules) == rules[0]
    assert convert._match_file_sort_rule("other.xlsx", rules) is None


def test_sort_prefix_add_and_remove_cover_unmatched_same_name_and_collision(tmp_path: Path) -> None:
    cfg = _config(tmp_path / "config.xlsx", sort_rows=[("Y", "1", "key", "")])
    same = _touch(tmp_path / "1_key.xlsx")
    collision = _touch(tmp_path / "key2.xlsx")
    _touch(tmp_path / "1_key2.xlsx")
    unmatched = _touch(tmp_path / "other.xlsx")
    missing = tmp_path / "missing.xlsx"

    added = convert.run_file_sort_prefix_49(cfg, [same, collision, unmatched, missing], tmp_path / "out", tmp_path / "logs")

    assert (added["renamed"], added["unmatched"], added["skipped"], added["failed"]) == (2, 1, 1, 0)
    assert same.exists() and (tmp_path / "1_key2_1.xlsx").exists()
    assert unmatched.exists()

    removed = convert.run_remove_file_sort_prefix_410(
        cfg, [tmp_path / "1_key.xlsx", tmp_path / "1_key2_1.xlsx", unmatched, missing],
        tmp_path / "out", tmp_path / "logs",
    )
    assert (removed["renamed"], removed["skipped"], removed["failed"]) == (2, 2, 0)
    assert (tmp_path / "key.xlsx").exists() and (tmp_path / "key2_1.xlsx").exists()


def test_sort_prefix_rename_failure_is_reported(tmp_path: Path, monkeypatch) -> None:
    cfg = _config(tmp_path / "config.xlsx", sort_rows=[("Y", "2", "key", "")])
    source = _touch(tmp_path / "key.xlsx")
    monkeypatch.setattr(convert.os, "rename", lambda *_args: (_ for _ in ()).throw(OSError("locked")))

    result = convert.run_file_sort_prefix_49(cfg, [source], tmp_path / "out", tmp_path / "logs")

    assert result["renamed"] == 0 and result["failed"] == 1
    assert source.exists()

    with pytest.raises(ValueError, match="没有启用且有效"):
        convert.run_file_sort_prefix_49(_config(tmp_path / "no-rules.xlsx"), [], tmp_path / "out", tmp_path / "logs")


def test_remove_sort_prefix_preserves_a_matching_content_prefix(tmp_path: Path) -> None:
    cfg = _config(
        tmp_path / "config.xlsx",
        rename_rows=[{"内容重命名启用": "Y", "内容重命名取值规则": "sheet1@A1"}],
        sort_rows=[("Y", "1", "key", "")],
    )
    source = tmp_path / "A_1_key.xlsx"
    _workbook(source, {"A1": "A"})

    result = convert.run_remove_file_sort_prefix_410(cfg, [source], tmp_path / "out", tmp_path / "logs")

    assert result["renamed"] == 1
    assert (tmp_path / "A_key.xlsx").is_file()


def test_remove_sort_prefix_without_rules_or_content_config_is_safe(tmp_path: Path) -> None:
    no_rules = _config(tmp_path / "no-rules.xlsx")
    with pytest.raises(ValueError, match="没有启用且有效"):
        convert.run_remove_file_sort_prefix_410(no_rules, [], tmp_path / "out", tmp_path / "logs")

    cfg = _config(tmp_path / "config.xlsx", sort_rows=[("Y", "3", "key", "")])
    source = _touch(tmp_path / "3_key.xls")
    result = convert.run_remove_file_sort_prefix_410(cfg, [source], tmp_path / "out", tmp_path / "logs")
    assert result["renamed"] == 1 and (tmp_path / "key.xls").is_file()


def test_remove_sort_prefix_reads_legacy_content_with_calamine_and_skips_errors(tmp_path: Path, monkeypatch) -> None:
    cfg = _config(
        tmp_path / "config.xlsx",
        rename_rows=[{"内容重命名启用": "Y", "内容重命名取值规则": "sheet1@A1"}],
        sort_rows=[("Y", "3", "key", "")],
    )
    legacy = tmp_path / "A_3_key.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("sheet1")
    sheet.write(0, 0, "A")
    workbook.save(str(legacy))
    corrupt = _touch(tmp_path / "B_3_key.xlsx")

    result = convert.run_remove_file_sort_prefix_410(cfg, [legacy, corrupt], tmp_path / "out", tmp_path / "logs")

    assert result["renamed"] == 1 and result["skipped"] == 1
    assert (tmp_path / "A_key.xls").is_file()


def test_batch_file_rename_handles_all_skip_and_collision_reasons(tmp_path: Path, monkeypatch) -> None:
    cfg = _config(tmp_path / "config.xlsx", rename_rows=[
        {"简称": "甲行", "全称": "甲银行", "代码": "001", "全称(代码映射)": "甲银行",
         "键": "数据日期", "值": "2026-10", "原文件名片段": "old", "新文件名片段": "new"},
        {"简称": "2026-10", "全称": "甲银行", "代码": "001", "全称(代码映射)": "甲银行"},
        {"简称": "乙行", "全称": "乙银行", "代码": "", "全称(代码映射)": "乙银行"},
        {"键": "报表名称", "值": "月报"},
    ])
    source = _touch(tmp_path / "甲行.xlsx")
    existing = _touch(tmp_path / "2026-10 001 甲银行 月报.xlsx")
    no_code = _touch(tmp_path / "乙行.xlsx")
    unknown = _touch(tmp_path / "unknown.xlsx")
    missing = tmp_path / "missing.xlsx"

    result = convert.run_batch_rename_33(cfg, [source, existing, no_code, unknown, missing], tmp_path / "logs")

    assert result == {"ok": 1, "skip": 4}
    assert existing.exists()
    assert (tmp_path / "2026-10 001 甲银行 月报_1.xlsx").exists()


def test_batch_file_rename_records_os_errors(tmp_path: Path, monkeypatch) -> None:
    cfg = _config(tmp_path / "config.xlsx", rename_rows=[
        {"简称": "甲行", "全称": "甲银行", "代码": "001", "全称(代码映射)": "甲银行",
         "键": "数据日期", "值": "2026", "原文件名片段": "old", "新文件名片段": "new"},
        {"键": "报表名称", "值": "月报"},
    ])
    source = _touch(tmp_path / "甲行.xlsx")
    monkeypatch.setattr(convert.os, "rename", lambda *_args: (_ for _ in ()).throw(OSError("locked")))

    result = convert.run_batch_rename_33(cfg, [source], tmp_path / "logs")

    assert result == {"ok": 0, "skip": 1} and source.exists()


def test_python_sheet_rename_skips_invalid_names_and_unsupported_inputs(tmp_path: Path) -> None:
    cfg = _config(tmp_path / "config.xlsx", rename_rows=[
        {"原表名": "same", "新表名": "same"},
        {"原表名": "invalid", "新表名": "bad/name"},
        {"原表名": "long", "新表名": "x" * 32},
        {"原表名": "duplicate", "新表名": "taken"},
    ])
    source = tmp_path / "book.xlsm"
    _workbook(source, sheets=["same", "invalid", "long", "duplicate", "taken"])
    other = _touch(tmp_path / "input.xls")
    missing = tmp_path / "missing.xlsx"

    result = convert.run_batch_rename_sheet_37(cfg, [source, other, missing], tmp_path / "logs")

    assert result == {"workbooks": 1, "rename_ok": 0, "skip": 6}
    assert convert._sheet_name_invalid("same", {"same"}, "same")
    assert convert._sheet_name_invalid("bad/name", set(), "old")
    assert convert._sheet_name_invalid("x" * 32, set(), "old")
    assert convert._sheet_name_invalid("taken", {"taken"}, "old")
    assert not convert._sheet_name_invalid("new", set(), "old")
    with pytest.raises(ValueError, match="缺少原表名/新表名"):
        convert.run_batch_rename_sheet_37(_config(tmp_path / "empty.xlsx"), [], tmp_path / "logs")


def test_content_spec_parsing_and_cell_reads_cover_missing_values(tmp_path: Path) -> None:
    assert convert._parse_sheet_cell_spec(" sheet1@a1；sheet2@B2 ") == [("sheet1", "A1"), ("sheet2", "B2")]
    for invalid in ("", "sheet1", "@A1", "sheet1@", "sheet1@A0"):
        with pytest.raises(ValueError):
            convert._parse_sheet_cell_spec(invalid)

    source = tmp_path / "book.xlsx"
    _workbook(source, {"A1": "<A>  B", "A2": None}, sheets=["sheet1"])
    assert convert._read_cells_openpyxl(source, [("sheet1", "A1"), ("missing", "C3"), ("sheet1", "A2")]) == [
        "_A_ B", "missing!C3", "sheet1!A2"
    ]
    assert convert._content_prefix_items([" A ", "", "a/b"]) == ["A", "a_b"]
    assert convert._already_renamed_by_content("A_B_file.xlsx", ["A", "B"])
    assert not convert._already_renamed_by_content("file.xlsx", [])
    assert convert._content_prefix(["A", "B"]) == "A_B_"
    assert convert._remove_content_prefix("A_B_A_file.xlsx", ["A", "B"]) == "file.xlsx"
    assert convert._remove_content_prefix("file.xlsx", []) is None
    assert convert._remove_content_prefix_keep_sort_prefix("003_A_file.xlsx", ["A"], ["003"]) == "003_file.xlsx"


def test_content_rename_handles_invalid_sources_and_rename_collision(tmp_path: Path) -> None:
    cfg = _config(tmp_path / "config.xlsx", rename_rows=[
        {"内容重命名启用": "Y", "内容重命名取值规则": "sheet1@A1"}
    ])
    source = tmp_path / "book.xlsx"
    collision = tmp_path / "A_book.xlsx"
    _workbook(source, {"A1": "A"})
    _workbook(collision, {"A1": "occupied"})
    invalid = _touch(tmp_path / "bad.xlsx")
    other = _touch(tmp_path / "other.txt")
    missing = tmp_path / "missing.xlsx"

    result = convert.run_batch_rename_by_content_46(cfg, [source, collision, invalid, other, missing], tmp_path / "logs")

    assert result["ok"] == 2 and result["skip"] == 3
    assert (tmp_path / "A_book_1.xlsx").is_file()
    assert (tmp_path / "occupied_A_book.xlsx").is_file()






def test_local_rename_rejects_empty_maps_and_reports_os_errors(tmp_path: Path, monkeypatch) -> None:
    empty = _config(tmp_path / "empty.xlsx")
    with pytest.raises(ValueError, match="缺少 M/N 列"):
        convert.run_batch_rename_local_map_48(empty, [], tmp_path / "logs")

    cfg = _config(tmp_path / "config.xlsx", rename_rows=[{"原文件名片段": "old", "新文件名片段": "new"}])
    source = _touch(tmp_path / "old.pdf")
    unmatched = _touch(tmp_path / "other.pdf")
    monkeypatch.setattr(convert.os, "rename", lambda *_args: (_ for _ in ()).throw(OSError("locked")))
    result = convert.run_batch_rename_local_map_48(cfg, [source, unmatched, tmp_path / "missing.pdf"], tmp_path / "logs")
    assert result == {"ok": 0, "skip": 3, "rules": 1}




def test_content_rename_and_unrename_legacy_xls_uses_calamine(tmp_path: Path) -> None:
    cfg = _config(tmp_path / "config.xlsx", rename_rows=[
        {"内容重命名启用": "Y", "内容重命名取值规则": "sheet1@A1"}
    ])
    source = tmp_path / "legacy.xls"
    book = xlwt.Workbook()
    book.add_sheet("sheet1").write(0, 0, "Alpha")
    book.save(str(source))
    renamed = convert.run_batch_rename_by_content_46(cfg, [source], tmp_path / "logs")
    marked = tmp_path / "Alpha_legacy.xls"
    assert renamed["ok"] == 1 and marked.is_file()
    restored = convert.run_batch_unrename_by_content_47(cfg, [marked], tmp_path / "logs")
    assert restored["ok"] == 1 and source.is_file()
