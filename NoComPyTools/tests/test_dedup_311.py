from datetime import date, datetime

from openpyxl import Workbook

from nocompytools.core.dedup_311 import _append_unique_source_to_target, _norm


def test_dedup_key_ignores_time_component():
    expected = "2026-01-31"

    assert _norm("2026-01-31") == expected
    assert _norm(date(2026, 1, 31)) == expected
    assert _norm(datetime(2026, 1, 31, 0, 0, 0)) == expected
    assert _norm(datetime(2026, 1, 31, 12, 30, 45)) == expected


def test_append_treats_text_date_and_excel_datetime_as_duplicate():
    src_wb = Workbook()
    src_ws = src_wb.active
    src_ws.append(["数据日期", "机构"])
    src_ws.append(["2026-01-31", "建设银行"])

    tgt_wb = Workbook()
    tgt_ws = tgt_wb.active
    tgt_ws.append(["数据日期", "机构"])
    tgt_ws.append([datetime(2026, 1, 31, 16, 20, 0), "建设银行"])

    appended, filtered, _ = _append_unique_source_to_target(src_ws, tgt_ws, [1, 2])

    assert appended == 0
    assert filtered == 1
    assert tgt_ws.max_row == 2

    src_wb.close()
    tgt_wb.close()
