from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from nocompytools.core.dedup_311 import (
    DedupTask,
    _append_unique_source_to_target,
    _build_existing_key_set,
    _dedup_tasks_append,
    _dedup_tasks_check_delete,
    _delete_duplicates,
    _fmt_header_items,
    _fmt_position_diff,
    _get_key_cols_by_comment,
    _header_mismatch_detail,
    _headers_same,
    _is_key_blank,
    _mark_duplicates,
    _norm,
    _used_col_count,
    run_append_by_config,
    run_check_by_comment,
    run_check_by_config,
    run_delete_by_comment,
    run_delete_by_config,
    run_precheck_by_config,
)


def _book(path: Path, rows: list[list[object]], sheet: str = "Data") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()
    return path


def _task(
    row_no: int,
    source: Path | None,
    sheet: str = "Data",
    keys: list[int] | None = None,
    target: Path | None = None,
    target_sheet: str | None = None,
    enabled: bool = True,
) -> DedupTask:
    return DedupTask(row_no, enabled, source, sheet, keys or [], target, target_sheet)


def test_key_normalization_used_columns_and_comment_key_selection() -> None:
    assert _norm(None) == ""
    assert _norm(datetime(2026, 4, 2, 12, 0)) == "2026-04-02"
    assert _norm(date(2026, 4, 2)) == "2026-04-02"
    assert _norm("2026年4月2日 12:30") == "2026-04-02"
    assert _norm("2026-02-31") == "2026-02-31"
    assert _is_key_blank(("", ""))
    assert _norm(" ") == ""
    assert not _is_key_blank((" ",))
    assert not _is_key_blank(("x", ""))

    wb = Workbook()
    ws = wb.active
    ws.append(["id", "value", None])
    ws["A1"].comment = Comment("源数据 标识列", "test")
    ws.append([1, "a"])
    assert _used_col_count(ws) == 3
    assert _get_key_cols_by_comment(ws) == [1]
    ws["A1"].comment = None
    assert _get_key_cols_by_comment(ws) == [1, 2, 3]
    wb.close()

    assert _used_col_count(SimpleNamespace(
        calculate_dimension=lambda: (_ for _ in ()).throw(RuntimeError("dimension error")), max_column=4
    )) == 4


def test_duplicate_helpers_skip_blank_rows_and_keep_first_duplicate(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["id", "value"])
    ws.append([1, "first"])
    ws.append([None, None])
    ws.append([1, "duplicate"])
    ws.append([2, "other"])
    assert _mark_duplicates(ws, [1], used_col=2) == 1
    assert ws["A4"].fill.fgColor.rgb in ("FFFF0000", "00FF0000")
    assert _delete_duplicates(ws, [1]) == 1
    assert ws.max_row == 4
    wb.close()

    source = _book(tmp_path / "source.xlsx", [["id", "value"], [1, "a"], [2, "b"], [2, "duplicate"], [None, None]])
    target = Workbook()
    target_ws = target.active
    target_ws.append(["id", "value"])
    target_ws.append([1, "old"])
    existing = _build_existing_key_set(target_ws, [1])
    assert existing == {("1",)}
    src = load_workbook(source, data_only=True)["Data"]
    appended, filtered, updated = _append_unique_source_to_target(src, target_ws, [1], existing)
    assert (appended, filtered) == (1, 2)
    assert updated == {("1",), ("2",)}
    assert list(target_ws.values) == [("id", "value"), (1, "old"), (2, "b")]
    assert _build_existing_key_set(Workbook().active, [1]) == set()
    target.close()
    src.parent.close()


def test_task_coalescing_and_header_formatting_helpers() -> None:
    first = _task(2, Path("same.xlsx"), keys=[1])
    duplicate = _task(3, Path("same.xlsx"), keys=[1])
    disabled = _task(4, Path("disabled.xlsx"), enabled=False)
    different_keys = _task(5, Path("same.xlsx"), keys=[2])
    checked, merged = _dedup_tasks_check_delete([disabled, first, duplicate, different_keys])
    assert checked == [first, different_keys] and merged == 1

    append_a = _task(6, Path("src.xlsx"), keys=[1], target=Path("target.xlsx"), target_sheet="Rows")
    append_b = _task(7, Path("src.xlsx"), keys=[1], target=Path("target.xlsx"), target_sheet="Rows")
    append_c = _task(8, Path("src.xlsx"), keys=[1], target=Path("target.xlsx"), target_sheet="Other")
    appended, merged = _dedup_tasks_append([disabled, append_a, append_b, append_c])
    assert appended == [append_a, append_c] and merged == 1

    wb = Workbook()
    src = wb.active
    src.append(["A", "B"])
    same = wb.create_sheet("same")
    same.append(["A", "B", "extra"])
    diff = wb.create_sheet("diff")
    diff.append(["B", "A"])
    assert _headers_same(src, same)
    assert not _headers_same(src, diff)
    empty = wb.create_sheet("empty")
    empty.delete_rows(1, empty.max_row)
    assert not _headers_same(src, empty)
    detail = _header_mismatch_detail(src, diff)
    assert detail == {
        "src_count": 2, "target_count": 2, "source_only": [], "target_only": [],
        "position_diff": [(1, "A", "B"), (2, "B", "A")],
    }
    assert _fmt_header_items([]) == "-"
    assert _fmt_header_items(["a", "b", "c"], limit=2) == "a;b...(+1)"
    assert _fmt_position_diff([]) == "-"
    assert "...(+1)" in _fmt_position_diff([(1, "a", "x"), (2, "b", "y")], limit=1)
    wb.close()


def test_comment_modes_validate_paths_and_merge_duplicate_tasks(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        run_check_by_comment(tmp_path / "missing.xlsx", tmp_path / "logs")
    with pytest.raises(ValueError, match="仅支持"):
        run_delete_by_comment(_book(tmp_path / "bad.txt", [["x"]]), tmp_path / "logs")

    source = _book(tmp_path / "source.xlsx", [["id", "value"], [1, "a"], [1, "duplicate"], [None, None]])
    wb = load_workbook(source)
    wb["Data"]["A1"].comment = Comment("源数据 标识列", "test")
    wb.create_sheet("Unmarked")["A1"] = "ignored"
    wb.save(source)
    wb.close()
    checked = run_check_by_comment(source, tmp_path / "logs")
    assert checked == {"hit_sheets": 1, "marked_rows": 1}
    deleted = run_delete_by_comment(source, tmp_path / "logs")
    assert deleted == {"hit_sheets": 1, "deleted_rows": 1}

    missing = _task(10, tmp_path / "missing.xlsx")
    unsupported_file = _book(tmp_path / "unsupported.txt", [["id"]])
    invalid_sheet = _task(11, source, "Missing", [1])
    duplicate = _task(12, source, "Data", [1])
    duplicate_task = _task(13, source, "Data", [1])
    disabled = _task(14, source, "Data", [1], enabled=False)
    check = run_check_by_config(
        [duplicate, duplicate_task, invalid_sheet, missing, _task(15, unsupported_file), disabled], tmp_path / "logs"
    )
    assert check["task_ok"] == 1 and check["task_skip"] == 3

    delete = run_delete_by_config([duplicate, duplicate_task, invalid_sheet, missing], tmp_path / "logs")
    assert delete["task_ok"] == 1 and delete["task_skip"] == 2


def test_append_by_config_creates_targets_reuses_caches_and_skips_bad_inputs(tmp_path: Path, capsys) -> None:
    source = _book(tmp_path / "source.xlsx", [
        ["id", "value"], [1, "a"], [2, "b"], [2, "duplicate"], [None, None],
    ])
    target = tmp_path / "target.xlsx"
    wb = Workbook()
    rows = wb.active
    rows.title = "Rows"
    rows.append(["id", "value"])
    rows.append([1, "old"])
    mismatch = wb.create_sheet("Mismatch")
    mismatch.append(["other", "header"])
    wb.create_sheet("Keep")["A1"] = "preserve"
    wb.save(target)
    wb.close()

    tasks = [
        _task(2, source, keys=[1], target=target, target_sheet="Rows"),
        _task(3, source, keys=[1], target=target, target_sheet="Rows"),
        _task(4, source, keys=[1], target=target, target_sheet="NewSheet"),
        _task(5, source, keys=[1], target=tmp_path / "fresh.xlsx", target_sheet="Fresh"),
        _task(6, source, keys=[1], target=target, target_sheet="Mismatch"),
        _task(7, source, keys=[1], target=None, target_sheet="MissingTarget"),
        _task(8, source, sheet="Missing", keys=[1], target=target, target_sheet="MissingSourceSheet"),
        _task(9, tmp_path / "missing.xlsx", keys=[1], target=target, target_sheet="MissingSource"),
        _task(10, _book(tmp_path / "unsupported.txt", [["id"]]), keys=[1], target=target, target_sheet="Unsupported"),
    ]

    result = run_append_by_config(tasks, tmp_path / "logs")

    assert result["task_ok"] == 3 and result["task_skip"] == 5
    assert result["appended_rows"] == 5 and result["deleted_rows"] == 4
    assert "header_mismatch" in capsys.readouterr().out
    out = load_workbook(target, data_only=True)
    try:
        assert out.sheetnames == ["Rows", "Mismatch", "Keep", "NewSheet"]
        assert out["Rows"].max_row == 3
        assert out["NewSheet"].cell(1, 1).value == "id"
        assert out["Keep"]["A1"].value == "preserve"
    finally:
        out.close()
    fresh = load_workbook(tmp_path / "fresh.xlsx", data_only=True)
    try:
        assert fresh.sheetnames == ["Fresh"]
        assert fresh["Fresh"]["A1"].value == "id"
    finally:
        fresh.close()


def test_precheck_reports_disabled_missing_unsupported_corrupt_and_empty_keys(tmp_path: Path) -> None:
    valid = _book(tmp_path / "valid.xlsx", [["id"], [1]])
    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_text("not xlsx", encoding="utf-8")
    unsupported = _book(tmp_path / "unsupported.txt", [["id"]])
    tasks = [
        _task(1, valid, enabled=False),
        _task(2, None),
        _task(3, valid, sheet=""),
        _task(4, tmp_path / "missing.xlsx"),
        _task(5, unsupported),
        _task(6, valid, sheet="Missing"),
        _task(7, corrupt),
        _task(8, valid, keys=[]),
    ]

    result = run_precheck_by_config(tasks, tmp_path / "logs")

    assert result == {"ok": 1, "warn": 1, "fail": 6}
