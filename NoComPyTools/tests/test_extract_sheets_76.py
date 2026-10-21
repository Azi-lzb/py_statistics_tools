from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
import xlwt
from openpyxl import Workbook, load_workbook

from nocompytools.core.config_xlsx import SHEET_EXTRACT_CONFIG, load_extract_tasks
from nocompytools.core.config_init import initialize_or_repair_config
from nocompytools.core import feature_1_6_extract_sheets as extract_module
from nocompytools.core.feature_1_6_extract_sheets import run_extract_sheets


NEW_EXTRACT_HEADERS = [
    "是否启用",
    "工作簿关键字",
    "工作表关键字",
    "排除关键字",
    "提取区域",
    "多区域拼接方式",
    "输出文件夹",
    "文件名",
    "备注",
]


def _write_config(path: Path, rows: list[dict]) -> None:
    df = pd.DataFrame(rows, columns=NEW_EXTRACT_HEADERS)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=SHEET_EXTRACT_CONFIG, index=False)


def _write_source(path: Path, sheet_name: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for r in range(1, 7):
        for c in range(1, 7):
            ws.cell(r, c, f"R{r}C{c}")
    wb.save(path)
    wb.close()


def _write_xls_source(path: Path) -> None:
    wb = xlwt.Workbook()
    ws = wb.add_sheet("源工作表")
    ws.write_merge(0, 0, 0, 1, "标题", xlwt.easyxf("font: bold on"))
    ws.write(1, 1, "B2")
    ws.write(3, 3, 42)
    ws.col(0).width = 5000
    wb.save(str(path))


def _read_values(path: Path, sheet_name: str) -> list[list]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    values = [
        [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for r in range(1, ws.max_row + 1)
    ]
    wb.close()
    return values


def test_extract_sheets_new_config_horizontal_multi_area_and_output_folder(tmp_path: Path):
    cfg = tmp_path / "config.xlsx"
    src = tmp_path / "2026存款表.xlsx"
    out = tmp_path / "output"
    logs = tmp_path / "logs"
    _write_source(src, "本外币明细")
    _write_config(cfg, [{
        "是否启用": "Y",
        "工作簿关键字": "存款,贷款",
        "工作表关键字": "本外币,人民币",
        "排除关键字": "",
        "提取区域": "A1:B2;D1:E2",
        "多区域拼接方式": "横向拼接",
        "输出文件夹": "子目录",
        "文件名": "横向结果",
        "备注": "",
    }])

    tasks = load_extract_tasks(cfg)
    stat = run_extract_sheets([src], tasks, out, logs)

    result = out / "子目录" / "横向结果.xlsx"
    assert stat["extracted_sheets"] == 1
    assert stat["paths"] == [str(result)]
    assert _read_values(result, "本外币明细") == [
        ["R1C1", "R1C2", "R1C4", "R1C5"],
        ["R2C1", "R2C2", "R2C4", "R2C5"],
    ]


def test_extract_sheets_vertical_multi_area(tmp_path: Path):
    cfg = tmp_path / "config.xlsx"
    src = tmp_path / "2026专项二批.xlsx"
    out = tmp_path / "output"
    logs = tmp_path / "logs"
    _write_source(src, "涉农贷款")
    _write_config(cfg, [{
        "是否启用": "是",
        "工作簿关键字": "专项二批,商业银行共享二批",
        "工作表关键字": "涉农",
        "排除关键字": "",
        "提取区域": "A1:B2;A5:B6",
        "多区域拼接方式": "纵向追加",
        "输出文件夹": "",
        "文件名": "纵向结果",
        "备注": "",
    }])

    tasks = load_extract_tasks(cfg)
    stat = run_extract_sheets([src], tasks, out, logs)

    result = out / "纵向结果.xlsx"
    assert stat["extracted_sheets"] == 1
    assert _read_values(result, "涉农贷款") == [
        ["R1C1", "R1C2"],
        ["R2C1", "R2C2"],
        ["R5C1", "R5C2"],
        ["R6C1", "R6C2"],
    ]


def test_partial_extract_uses_calamine_without_openpyxl_source(tmp_path: Path, monkeypatch):
    source = tmp_path / "2026 report.xlsx"
    output_dir = tmp_path / "output"
    _write_source(source, "源工作表")
    monkeypatch.setattr(
        extract_module, "load_workbook",
        lambda *_args, **_kwargs: pytest.fail("区域提取不应打开源工作簿的 openpyxl 对象"),
    )

    stat = run_extract_sheets([source], [{
        "workbook_keywords": "2026", "sheet_keywords": "源工作",
        "exclude_keywords": "", "area_spec": "A1:B2", "output_name": "fast",
    }], output_dir, tmp_path / "logs")

    assert stat["extracted_sheets"] == 1
    assert _read_values(output_dir / "fast.xlsx", "源工作表") == [
        ["R1C1", "R1C2"], ["R2C1", "R2C2"],
    ]


def test_full_extract_reads_xls_values_without_source_openpyxl(tmp_path: Path, monkeypatch):
    source = tmp_path / "2026 report.xls"
    output_dir = tmp_path / "output"
    _write_xls_source(source)
    monkeypatch.setattr(
        extract_module, "load_workbook",
        lambda *_args, **_kwargs: pytest.fail("整表提取不应使用 openpyxl 读取源文件"),
    )

    stat = run_extract_sheets([source], [{
        "workbook_keywords": "2026", "sheet_keywords": "源工作",
        "exclude_keywords": "", "area_spec": "整表", "output_name": "full-values",
    }], output_dir, tmp_path / "logs")

    result = output_dir / "full-values.xlsx"
    assert stat["extracted_sheets"] == 1
    wb = load_workbook(result, data_only=True)
    try:
        ws = wb["源工作表"]
        assert ws["A1"].value == "标题"
        assert ws["B2"].value == "B2"
        assert ws["D4"].value == 42
        assert not ws.merged_cells.ranges
        assert ws.column_dimensions["A"].width != 50
    finally:
        wb.close()


def test_extract_sheets_skips_disabled_and_non_matching_workbook(tmp_path: Path):
    cfg = tmp_path / "config.xlsx"
    src = tmp_path / "2026存款表.xlsx"
    out = tmp_path / "output"
    logs = tmp_path / "logs"
    _write_source(src, "本外币明细")
    _write_config(cfg, [
        {
            "是否启用": "N",
            "工作簿关键字": "存款",
            "工作表关键字": "本外币",
            "排除关键字": "",
            "提取区域": "A1:B2",
            "多区域拼接方式": "",
            "输出文件夹": "",
            "文件名": "禁用结果",
            "备注": "",
        },
        {
            "是否启用": "Y",
            "工作簿关键字": "贷款",
            "工作表关键字": "本外币",
            "排除关键字": "",
            "提取区域": "A1:B2",
            "多区域拼接方式": "",
            "输出文件夹": "",
            "文件名": "不应生成",
            "备注": "",
        },
    ])

    tasks = load_extract_tasks(cfg)
    stat = run_extract_sheets([src], tasks, out, logs)

    assert len(tasks) == 1
    assert stat["extracted_sheets"] == 0
    assert not (out / "不应生成.xlsx").exists()


def test_extract_sheets_keywords_support_and_groups_with_or_options(tmp_path: Path):
    cfg = tmp_path / "config.xlsx"
    src_ok = tmp_path / "2026专项二批_商业银行共享二批.xlsx"
    src_missing = tmp_path / "2026专项二批.xlsx"
    out = tmp_path / "output"
    logs = tmp_path / "logs"
    _write_source(src_ok, "人民币涉农贷款")
    _write_source(src_missing, "人民币涉农贷款")
    _write_config(cfg, [{
        "是否启用": "Y",
        "工作簿关键字": "专项二批;商业银行共享二批",
        "工作表关键字": "本外币,人民币;涉农",
        "排除关键字": "",
        "提取区域": "A1:B1",
        "多区域拼接方式": "",
        "输出文件夹": "",
        "文件名": "关键词结果",
        "备注": "",
    }])

    tasks = load_extract_tasks(cfg)
    stat = run_extract_sheets([src_ok, src_missing], tasks, out, logs)

    result = out / "关键词结果.xlsx"
    assert stat["extracted_sheets"] == 1
    assert _read_values(result, "人民币涉农贷款") == [["R1C1", "R1C2"]]


def test_extract_sheets_exclude_keywords_skip_workbook_and_sheet(tmp_path: Path):
    cfg = tmp_path / "config.xlsx"
    src_ok = tmp_path / "2026正式_商业银行共享二批.xlsx"
    src_excluded_file = tmp_path / "2026草稿_商业银行共享二批.xlsx"
    src_excluded_sheet = tmp_path / "2026正式_专项二批.xlsx"
    out = tmp_path / "output"
    logs = tmp_path / "logs"
    _write_source(src_ok, "人民币涉农贷款")
    _write_source(src_excluded_file, "人民币涉农贷款")
    _write_source(src_excluded_sheet, "人民币涉农草稿")
    _write_config(cfg, [{
        "是否启用": "Y",
        "工作簿关键字": "2026正式,2026草稿",
        "工作表关键字": "人民币;涉农",
        "排除关键字": "草稿,测试",
        "提取区域": "A1:B1",
        "多区域拼接方式": "",
        "输出文件夹": "",
        "文件名": "排除结果",
        "备注": "",
    }])

    tasks = load_extract_tasks(cfg)
    stat = run_extract_sheets([src_ok, src_excluded_file, src_excluded_sheet], tasks, out, logs)

    result = out / "排除结果.xlsx"
    assert stat["extracted_sheets"] == 1
    assert _read_values(result, "人民币涉农贷款") == [["R1C1", "R1C2"]]


def test_extract_sheets_append_existing_output_keeps_previous_sheets(tmp_path: Path):
    cfg_first = tmp_path / "config_first.xlsx"
    cfg_second = tmp_path / "config_second.xlsx"
    src_first = tmp_path / "2026第一批.xlsx"
    src_second = tmp_path / "2026第二批.xlsx"
    out = tmp_path / "output"
    logs = tmp_path / "logs"
    _write_source(src_first, "第一张表")
    _write_source(src_second, "第二张表")
    base_row = {
        "是否启用": "Y",
        "工作簿关键字": "2026",
        "工作表关键字": "",
        "排除关键字": "",
        "提取区域": "A1:B1",
        "多区域拼接方式": "",
        "输出文件夹": "",
        "文件名": "合并结果",
        "备注": "",
    }
    _write_config(cfg_first, [base_row])
    _write_config(cfg_second, [base_row])

    first_tasks = load_extract_tasks(cfg_first)
    second_tasks = load_extract_tasks(cfg_second)
    run_extract_sheets([src_first], first_tasks, out, logs)
    run_extract_sheets([src_second], second_tasks, out, logs, append_existing=True)

    result = out / "合并结果.xlsx"
    wb = load_workbook(result, data_only=True)
    try:
        assert "第一张表" in wb.sheetnames
        assert "第二张表" in wb.sheetnames
        assert wb["第一张表"]["A1"].value == "R1C1"
        assert wb["第二张表"]["A1"].value == "R1C1"
    finally:
        wb.close()


def test_config_repair_creates_new_extract_sheet_headers(tmp_path: Path):
    cfg = tmp_path / "config.xlsx"

    initialize_or_repair_config(cfg)

    headers = pd.read_excel(cfg, sheet_name=SHEET_EXTRACT_CONFIG, nrows=0).columns.tolist()
    assert headers[:len(NEW_EXTRACT_HEADERS)] == NEW_EXTRACT_HEADERS


def test_config_repair_adds_extract_sheet_header_comments(tmp_path: Path):
    cfg = tmp_path / "config.xlsx"

    initialize_or_repair_config(cfg)

    wb = load_workbook(cfg)
    ws = wb[SHEET_EXTRACT_CONFIG]
    comments = {ws.cell(1, col).value: ws.cell(1, col).comment.text for col in range(1, ws.max_column + 1)}
    wb.close()

    assert "分号/中文分号=必须同时命中" in comments["工作簿关键字"]
    assert "逗号/中文逗号/|=同组任选" in comments["工作表关键字"]
    assert "命中则跳过" in comments["排除关键字"]
    assert "支持 A1:D20、A:D、1:10" in comments["提取区域"]
    assert "纵向追加=上下追加" in comments["多区域拼接方式"]
