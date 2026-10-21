from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from nocompytools.core import feature_1_6_extract_sheets as extract


def _data():
    return [["a", "b", "c"], ["d", "e", "f"], ["g", "h", "i"]]


def test_range_and_keyword_parsing_helpers():
    assert extract._normalize_range_spec("") == ""
    assert extract._normalize_range_spec(" A1；B2、C3：D4 ") == "A1,B2,C3:D4"
    assert extract._token_to_index(" 7 ", True) == 7
    assert extract._token_to_index("B", False) == 2
    assert extract._token_to_index("B", True) == 0
    assert extract._token_to_index("?", False) == 0
    assert extract._parse_indexes("", 3, True) == [1, 2, 3]
    assert extract._parse_indexes("9:2,2,4,invalid", 4, False) == [2, 3, 4]
    assert extract._parse_indexes("bad", 3, True) == [1, 2, 3]
    assert extract._parse_indexes("1", 0, False) == []
    assert extract._build_exact_name(["", "A", "B"]) == "A-B"
    assert extract._sheet_matches_all("2026 Profit Report", ["profit", "report"])
    assert not extract._sheet_matches_all("2026 Profit Report", ["profit", "balance"])


def test_sheet_selection_and_keyword_groups():
    wb = Workbook()
    wb.active.title = "A-B"
    wb.create_sheet("A-B detail")
    wb.create_sheet("A-B draft")
    matched = extract._find_matched_sheets(wb, ["A", "B"])
    assert [ws.title for ws in matched] == ["A-B", "A-B detail", "A-B draft"]
    assert extract._split_keyword_groups("") == []
    assert extract._split_keyword_groups("a, b; c|d；e") == [["a", "b"], ["c", "d"], ["e"]]
    assert extract._name_matches_keyword_options("2026 report", "2025,2026;report")
    assert not extract._name_matches_keyword_options("2026 report", "2025;report")
    assert extract._workbook_matches(Path("2026 report.xlsx"), {"workbook_keywords": "2026"})
    assert not extract._workbook_matches(Path("2026 report.xlsx"), {"workbook_keywords": "2025"})
    assert not extract._name_matches_exclude("report", {"exclude_keywords": ""})
    assert extract._name_matches_exclude("report draft", {"exclude_keywords": "final,draft"})
    task = {"sheet_keywords": "detail", "exclude_keywords": "draft;temp"}
    assert [ws.title for ws in extract._find_matched_sheets_for_task(wb, task)] == ["A-B detail"]
    wb.close()


def test_sheet_names_and_data_extraction_helpers():
    wb = Workbook()
    wb.active.title = "Valid"
    assert extract._sanitize_sheet_name("A/B:*?[]\\long name") == "A_B______long name"
    assert extract._unique_sheet_name(wb, "A/B") == "A_B"
    wb.create_sheet("A_B")
    assert extract._unique_sheet_name(wb, "A/B") == "A_B_2"
    long_name = "x" * 40
    wb.create_sheet(long_name[:31])
    assert extract._unique_sheet_name(wb, long_name) == "x" * 28 + "_2"
    assert extract._read_ws_values(SimpleNamespace(max_row=0, max_column=2, cell=None)) == []

    data = _data()
    assert extract._extract_values([], "", "", False) == []
    assert extract._extract_values([[], []], "", "", True) == []
    assert extract._extract_values(data, "1:2", "B:C", False) == [["b", "c"], ["e", "f"]]
    assert extract._extract_values([["a"], ["b", "c"]], "1:2", "A:C", False) == [
        ["a", None], ["b", "c"]
    ]
    wb.close()


@pytest.mark.parametrize(
    ("area", "expected"),
    [
        ("整表", ([1, 2, 3], [1, 2, 3])),
        ("C3:A1", ([1, 2, 3], [1, 2, 3])),
        ("B:C", ([1, 2, 3], [2, 3])),
        ("3:1", ([1, 2, 3], [1, 2, 3])),
        ("B2", ([2], [2])),
        ("B20", ([], [])),
        ("invalid", None),
    ],
)
def test_area_parser_supports_cells_rows_columns_and_invalid_specs(area, expected):
    assert extract._parse_area_indexes(area, 3, 3) == expected


def test_area_stitch_compact_and_output_path_helpers(tmp_path):
    data = _data()
    assert extract._is_full_area_spec("")
    assert extract._is_full_area_spec(" 全表 ")
    assert not extract._is_full_area_spec("A1")
    assert extract._split_area_specs("整表") == [""]
    assert extract._split_area_specs("A1:B2；D1:D3") == ["A1:B2", "D1:D3"]
    assert extract._extract_area_values([], "A1") == []
    assert extract._extract_area_values([[], []], "A1") == []
    assert extract._extract_area_values(data, "invalid") is None
    assert extract._extract_area_values(data, "A1:B2") == [["a", "b"], ["d", "e"]]
    assert extract._stitch_area_values([], "横向拼接") == []
    assert extract._stitch_area_values([[[1], [2]], [[3, 4]]], "纵向追加") == [[1], [2], [3, 4]]
    assert extract._stitch_area_values([[[1], [2]], [[3, 4]]], "横向拼接") == [
        [1, 3, 4], [2, None, None]
    ]
    assert extract._extract_configured_values(data, "A1:B2", "") == [["a", "b"], ["d", "e"]]
    assert extract._extract_configured_values(data, "A1:A2;bad", "") is None
    assert extract._extract_configured_values(data, "A1:A2;C1:C3", "横向拼接") == [
        ["a", "c"], ["d", "f"], [None, "i"]
    ]
    assert extract._compact([[None, ""], ["", "  "]]) == []
    assert extract._compact([[None, "x", None], [" ", "y", ""]]) == [["x"], ["y"]]
    assert extract._is_blank_spec("  ")
    assert not extract._is_blank_spec("A1")
    assert extract._is_full_sheet_task({"area_spec": "全部"})
    assert not extract._is_full_sheet_task({"area_spec": "A1"})
    assert extract._output_path_for_task({}, tmp_path, "default") == tmp_path / "default.xlsx"
    assert extract._output_path_for_task(
        {"output_folder": "child", "output_name": "result.csv"}, tmp_path, "default"
    ) == tmp_path / "child" / "result.xlsx"
    absolute = tmp_path / "outside" / "book.xlsx"
    assert extract._output_path_for_task(
        {"output_folder": "ignored", "output_name": str(absolute)}, tmp_path, "default"
    ) == absolute


def test_full_sheet_values_and_default_sheet_cleanup(tmp_path):
    dst_wb = Workbook()
    dst = dst_wb.active
    extract._write_values(dst, [[None, "Title"], [10, None, 20]])
    path = tmp_path / "extract_values.xlsx"
    dst_wb.save(path)
    check_wb = load_workbook(path, data_only=True)
    check = check_wb.active
    assert check["B1"].value == "Title"
    assert check["A2"].value == 10 and check["C2"].value == 20
    assert not check.merged_cells.ranges
    cleanup_wb = Workbook()
    cleanup_wb.create_sheet("real")
    extract._remove_default_blank(cleanup_wb)
    assert cleanup_wb.sheetnames == ["real"]
    one_sheet = Workbook()
    extract._remove_default_blank(one_sheet)
    assert len(one_sheet.sheetnames) == 1
    dst_wb.close()
    check_wb.close()
    path.unlink(missing_ok=True)
    cleanup_wb.close()
    one_sheet.close()


def test_extract_workflow_empty_invalid_unmatched_and_blank_inputs(tmp_path):
    out, logs = tmp_path / "out", tmp_path / "logs"
    assert extract.run_extract_sheets([], [], out, logs)["failed_files"] == 0

    missing = tmp_path / "missing.xlsx"
    failed = extract.run_extract_sheets([missing], [{"workbook_keywords": ""}], out, logs)
    assert failed["failed_files"] == 1

    src = tmp_path / "2026 data.xlsx"
    wb = Workbook()
    wb.active.title = "Report"
    wb.active["A1"] = "value"
    wb.create_sheet("Report detail")
    wb.create_sheet("Draft")
    wb.save(src)
    wb.close()

    no_match = extract.run_extract_sheets([src], [{
        "workbook_keywords": "other", "sheet_keywords": "Report", "exclude_keywords": "",
        "area_spec": "A1", "output_name": "unused",
    }], out, logs)
    assert no_match["extracted_sheets"] == 0

    invalid_area = extract.run_extract_sheets([src], [{
        "workbook_keywords": "2026", "sheet_keywords": "Report", "exclude_keywords": "Draft",
        "area_spec": "not-an-area", "output_name": "invalid",
    }], out, logs)
    assert invalid_area["extracted_sheets"] == 0
    assert invalid_area["output_files"] == 0
    assert not (out / "invalid.xlsx").exists()

    blank = tmp_path / "blank.xlsx"
    Workbook().save(blank)
    blank_result = extract.run_extract_sheets([blank], [{
        "workbook_keywords": "blank", "sheet_keywords": "", "exclude_keywords": "",
        "area_spec": "整表", "output_name": "blank-result",
    }], out, logs)
    assert blank_result["extracted_sheets"] == 0
    assert not (out / "blank-result.xlsx").exists()
