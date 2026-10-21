from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from nocompytools.core import (
    feature_1_1_split_village_bank as split_village,
    feature_1_2_normalize_institution as normalize_institution,
    feature_1_3_remove_foreign_bank as remove_foreign,
    feature_1_4_fx_header as fx_header,
    feature_1_5_region_sum_check as region_sum,
    feature_1_8_adjust_rural_loan as loan_adjust,
    feature_1_9_adjust_rural_loan_quarterly as loan_quarterly,
)


def _save_book(path: Path, fill) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    fill(wb)
    wb.save(path)
    wb.close()
    return path


def test_split_village_bank_matrix_and_sheet_helpers(tmp_path: Path) -> None:
    assert split_village._to_num(None) == 0
    assert split_village._to_num(" 2.5 ") == 2.5
    assert split_village._to_num("bad") == 0
    assert split_village._yoy(0, 0) == 0
    assert split_village._yoy(2, 0) == "∞"
    assert split_village._yoy(15, 10) == 50
    assert split_village._detect_sheet_type("本外币") == "本外币"
    assert split_village._detect_sheet_type("人民币") == "人民币"
    assert split_village._detect_sheet_type("外汇") == "外汇"
    assert split_village._detect_sheet_type("其他") is None

    current = tmp_path / "current.xlsx"
    previous = tmp_path / "previous.xlsx"

    def create(wb, amount):
        for kind in split_village.INDICATORS:
            ws = wb.create_sheet(f"{kind}村镇银行")
            for row in (5, 6):
                ws[f"C{row}"] = amount
                ws[f"D{row}"] = "d"
                ws[f"E{row}"] = "e"
                ws[f"CI{row}"] = amount * 2
                ws[f"CJ{row}"] = "j"
                ws[f"CK{row}"] = "k"

    _save_book(current, lambda wb: create(wb, 20))
    _save_book(previous, lambda wb: create(wb, 10))
    matrix = split_village.compute_village_bank_matrix(current, previous)
    assert len(matrix) == 6 and matrix[0] == [20, "d", "e", 100, 40, "j", "k", 100]

    wb = Workbook()
    ws = wb.active
    ws.append(["机构", "1", "2"])
    ws.append(["村镇银行", 0, 0])
    assert split_village._find_village_bank_row(ws) == 2
    assert split_village._find_village_bank_row(Workbook().active) is None
    split_village._fill_rows(ws, 2, [[1] * 8, [2] * 8, [3] * 8, [4] * 8, [5] * 8, [6] * 8], "外汇")
    assert ws.cell(2, 2).value == 3 and ws.cell(3, 2).value == 6
    wb.close()


def test_split_village_bank_process_and_empty_calculation(tmp_path: Path) -> None:
    target = tmp_path / "target.xlsx"

    def make_target(wb):
        branch = wb.create_sheet("报表分机构本外币")
        branch.append(["机构", "a"])
        branch.append(["村镇银行", 1])
        wb.create_sheet("报表分机构其他").append(["机构"])
        wb.create_sheet("not-a-branch").append(["机构"])

    _save_book(target, make_target)
    stat = split_village._process_one_workbook(target, [[i] * 8 for i in range(1, 7)], SimpleNamespace(info=lambda *_: None))
    assert stat == {"branch_sheets": 2, "filled_sheets": 1, "skipped_sheets": 1}
    wb = load_workbook(target, data_only=True)
    try:
        assert wb["报表分机构本外币"]["A2"].value == split_village.SPLIT_ROW_1_NAME
        assert wb["报表分机构本外币"]["A3"].value == split_village.SPLIT_ROW_2_NAME
    finally:
        wb.close()

    current = tmp_path / "zero.xlsx"
    previous = tmp_path / "zero-prev.xlsx"
    _save_book(current, lambda wb: wb.create_sheet("no indicator"))
    _save_book(previous, lambda wb: wb.create_sheet("no indicator"))
    empty = split_village.run_split_village_bank(current, previous, [], tmp_path / "logs")
    assert empty["matrix_empty"] is True and empty["files"] == 0


def test_institution_normalization_and_foreign_row_helpers(tmp_path: Path) -> None:
    mapping = {"ABC Bank": "Standard Bank", "Branch": "Branch Standard"}
    keys = [(key.lower(), key) for key in mapping]
    assert normalize_institution._norm_name("ABC Bank", mapping, keys) == "Standard Bank"
    assert normalize_institution._norm_name("abc bank office", mapping, keys) == "Standard Bank"
    assert normalize_institution._norm_name("unmatched", mapping, keys) == "unmatched"

    wb = Workbook()
    ws = wb.active
    ws.append(["id", "name"])
    ws.append([None, "Branch office"])
    ws.append([None, " "])
    assert normalize_institution._find_last_row(ws, 2) == 2
    assert normalize_institution._process_sheet(ws, mapping, keys) == 1
    assert ws["B2"].value == "Branch Standard"
    assert normalize_institution._process_sheet(Workbook().active, mapping, keys) == 0
    wb.close()

    clean = Workbook().active
    clean.append(["row", "x"])
    clean.append(["Local Bank", 1])
    clean.append(["Foreign A", 2])
    clean.append(["Foreign B", 3])
    clean.append(["Local Bank", 4])
    assert remove_foreign._find_last_row_colA(clean) == 5
    assert remove_foreign._delete_foreign_rows(clean, {"Foreign A", "Foreign B"}) == 2
    assert clean.max_row == 3
    assert remove_foreign._delete_foreign_rows(Workbook().active, {"Foreign A"}) == 0
    assert remove_foreign._build_copy_path(Path("report.XLSX")) == Path("report(金融局).XLSX")


def test_fx_header_helpers_and_region_sum_numeric_edges(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.oddHeader.left.text = "金额（万元）"
    ws.evenHeader.center.text = "固定标题"
    assert fx_header._process_sheet(ws) is True
    assert "万美元" in ws.oddHeader.left.text
    assert fx_header._process_sheet(Workbook().active) is False
    wb.close()

    assert region_sum._to_num_or_none(None) == 0
    assert region_sum._to_num_or_none("-") == 0
    assert region_sum._to_num_or_none("2.5") == 2.5
    assert region_sum._to_num_or_none("bad") is None
    empty = Workbook().active
    assert region_sum._check_one_sheet(empty) == (1, "工作表为空或没有数据")
    empty.append(["序号", "地区"])
    empty.append([1, "惠阳区", 2])
    assert region_sum._check_one_sheet(empty) == (1, "未找到'惠州市'行")

    ws = Workbook().active
    ws.append(["序号", "地区", "金额"])
    ws.append([None, "表头"])
    ws.append([None, "惠州市", "bad"])
    ws.append([None, "惠阳区", "bad"])
    errors, detail = region_sum._check_one_sheet(ws)
    assert errors == 0 and "未找到" in detail

    short = Workbook().active
    short.append(["short"])
    short.append([1, None])
    assert region_sum._check_one_sheet(short) == (1, "未找到'惠州市'行")
    no_data_area = Workbook().active
    no_data_area.append(["id", "惠州市"])
    assert region_sum._check_one_sheet(no_data_area) == (1, "未找到有效数据区域")

    ws = Workbook().active
    ws.append(["id", "地区", "金额", None, None])
    ws.append([None, "header", None, None, None])
    ws.append([None, "惠州市", 0, 5, 8])
    ws.append([None, "惠阳区", 1, "bad"])
    errors, detail = region_sum._check_one_sheet(ws)
    assert errors == 2
    assert "列C" in detail and "列E" in detail and "未找到" in detail


def test_region_sum_check_reports_unreadable_sources(tmp_path: Path) -> None:
    broken = tmp_path / "broken.xlsx"
    broken.write_text("not an xlsx", encoding="utf-8")
    result = region_sum.run_region_sum_check([broken], tmp_path / "out", tmp_path / "logs")
    assert result["failed_files"] == 1 and result["files"] == 0
    assert Path(result["path"]).is_file()


def test_loan_adjustment_helper_maps_and_missing_sheet_errors(tmp_path: Path) -> None:
    assert loan_adjust._normalize("  A\t B  ") == "A B"
    assert loan_adjust._to_num("1,200") == 1200
    assert loan_adjust._to_num("-") == 0
    ws = Workbook().active
    ws.cell(3, 2, " 机构甲 ")
    ws.cell(3, 3, "3.5")
    ws.cell(4, 2, "机构甲")
    values, rows, duplicates = loan_adjust._build_inst_map(ws)
    assert values == {"机构甲": 3.5} and rows == {"机构甲": 3} and duplicates == ["机构甲"]
    assert loan_adjust._find_header_col(ws, "missing") == 0
    assert loan_adjust._find_city_row(ws, "missing") == 0
    wb = Workbook()
    wb.create_sheet("惠州市-涉农贷款分机构")
    wb.create_sheet("惠州市备用-涉农贷款分机构")
    with pytest.raises(ValueError, match="多个目标工作表"):
        loan_adjust._find_sheet(wb, loan_adjust.INST_SHEET_KEYS)
    wb.close()

    current = tmp_path / "current.xlsx"
    previous = tmp_path / "previous.xlsx"
    _save_book(current, lambda book: book.create_sheet("unrelated"))
    _save_book(previous, lambda book: book.create_sheet("unrelated"))
    result = loan_adjust.run_adjust_rural_loan(current, previous, tmp_path / "logs")
    assert result["error"] and "未找到" in result["error"]


def test_quarterly_helper_duplicate_and_region_resolution() -> None:
    assert loan_quarterly._region_from_sheet_name("惠阳区-涉农贷款分机构表") == "惠阳区"
    assert loan_quarterly._region_from_sheet_name("not-an-institution-sheet") == ""
    ws = Workbook().active
    ws.append(["idx", "institution", "value"])
    ws.append([])
    ws.append([1, "机构甲", 10])
    ws.append([2, "机构甲", 11])
    values, rows, dups = loan_quarterly._build_institution_map(ws)
    assert values == {"机构甲": 10} and rows == {"机构甲": 3} and dups == ["机构甲"]
    wb = Workbook()
    wb.create_sheet("惠阳区-涉农贷款分机构表")
    wb.create_sheet("惠阳区 涉农贷款分机构表")
    with pytest.raises(ValueError, match="重复地区"):
        loan_quarterly._institution_sheets(wb)
    wb.close()
    assert loan_quarterly._find_area_sheet(Workbook()) is None

    assert loan_quarterly._last_row(Workbook().active) == 0
    assert loan_quarterly._last_col(Workbook().active) == 0
    unnamed = Workbook()
    unnamed.create_sheet("-涉农贷款分机构表")
    with pytest.raises(ValueError, match="识别地区"):
        loan_quarterly._institution_sheets(unnamed)
    unnamed.close()


def test_quarterly_adjustment_reports_missing_previous_region(tmp_path: Path) -> None:
    current = tmp_path / "current.xlsx"
    previous = tmp_path / "previous.xlsx"

    def make_current(wb):
        inst = wb.create_sheet("惠州市-涉农贷款分机构表")
        inst.append(["序号", "机构名称", "贷款"])
        inst.append([])
        inst.append([1, "机构甲", 10])
        area = wb.create_sheet("涉农贷款分地区统计表")
        area.append(["地区", "涉农贷款比上月"])
        area.append(["惠州市", 1])

    def make_previous(wb):
        inst = wb.create_sheet("惠州市-涉农贷款分机构表")
        inst.append(["序号", "机构名称", "贷款"])
        inst.append([])
        inst.append([1, "机构甲", 9])

    _save_book(current, make_current)
    _save_book(previous, make_previous)
    result = loan_quarterly.run_adjust_rural_loan_quarterly(current, previous, tmp_path / "logs")
    assert result["error"] is None and result["consistent"] is True
    assert result["regions"][0]["matched"] == 1 and result["regions"][0]["fill_total"] == 1

    missing_previous = tmp_path / "missing-previous.xlsx"
    _save_book(missing_previous, lambda wb: wb.create_sheet("unrelated"))
    error = loan_quarterly.run_adjust_rural_loan_quarterly(current, missing_previous, tmp_path / "logs")
    assert "上期文件未找到" in error["error"]

