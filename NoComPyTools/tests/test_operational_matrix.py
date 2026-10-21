"""Functional coverage for menu features that are not covered by the timeline e2e tests."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from nocompytools.core.config_init import initialize_or_repair_config
from nocompytools.core.config_xlsx import (
    DEDUP_TASK_COLS,
    PRINT_CONFIG_COLS,
    SHEET_DEDUP_TASK,
    SHEET_PRINT_CONFIG,
    load_dedup_tasks,
    load_print_tasks,
)
from nocompytools.core.dedup_311 import (
    DedupTask,
    run_append_by_config,
    run_check_by_comment,
    run_check_by_config,
    run_delete_by_comment,
    run_delete_by_config,
    run_precheck_by_config,
)
from nocompytools.core.feature_1_1_split_village_bank import run_split_village_bank
from nocompytools.core.feature_1_2_normalize_institution import run_normalize_institution
from nocompytools.core.feature_1_3_remove_foreign_bank import run_remove_foreign_bank
from nocompytools.core.feature_1_4_fx_header import run_fx_header_fix
from nocompytools.core.feature_1_5_region_sum_check import run_region_sum_check
from nocompytools.core.feature_1_8_adjust_rural_loan import run_adjust_rural_loan
from nocompytools.core.print_310 import (
    run_print_config_all_modes,
    run_print_config_precheck,
    run_print_fast_by_comment,
    run_print_keep_by_comment,
)
from nocompytools.core.summary_321_322 import run_summary_by_comment, run_summary_by_usedrange


def _save(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def _dedup_book(path: Path, *, comments: bool = False) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["机构", "金额"])
    if comments:
        ws["A1"].comment = Comment("源数据；标识列", "test")
    ws.append(["甲行", 10])
    ws.append(["甲行", 11])
    ws.append([None, None])
    return _save(wb, path)


def _print_source(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "报表"
    ws.append([0, "标题"])
    ws.append(["甲行", 5])
    ws["A1"].comment = Comment("打印 保留源格式", "test")
    ws["A1"].comment = Comment("打印区域 1；打印 保留源格式", "test")
    ws["B2"].comment = Comment("打印区域#1", "test")
    ws["B1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws.column_dimensions["A"].width = 18
    return _save(wb, path)


def _region_book(path: Path, *, mismatch: bool = False) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "资产分地区"
    ws.append(["序号", "地区", "指标甲"])
    ws.append([None, "地区", "指标甲"])
    ws.append([None, "惠州市", 101 if mismatch else 100])
    values = {
        "惠阳区": 40,
        "博罗县": 30,
        "惠东县": 10,
        "龙门县": 5,
        "惠州市本部": 10,
        "大亚湾": 5,
    }
    for name, value in values.items():
        ws.append([None, name, value])
    return _save(wb, path)


def test_dedup_comment_check_then_delete_marks_only_duplicate_rows(tmp_path: Path) -> None:
    path = _dedup_book(tmp_path / "comment.xlsx", comments=True)

    checked = run_check_by_comment(path, tmp_path / "logs")
    assert checked == {"hit_sheets": 1, "marked_rows": 1}
    wb = load_workbook(path)
    ws = wb["Data"]
    assert ws["A3"].fill.fill_type == "solid"
    assert ws["A3"].fill.fgColor.rgb in ("FFFF0000", "00FF0000")
    assert ws["A2"].fill.fill_type is None
    wb.close()

    deleted = run_delete_by_comment(path, tmp_path / "logs")
    assert deleted == {"hit_sheets": 1, "deleted_rows": 1}
    wb = load_workbook(path, data_only=True)
    assert [wb["Data"].cell(row, 1).value for row in range(2, 4)] == ["甲行", None]
    wb.close()


def test_dedup_config_check_delete_precheck_and_append(tmp_path: Path) -> None:
    source = _dedup_book(tmp_path / "source.xlsx")
    task = DedupTask(2, True, source, "Data", [1])

    precheck = run_precheck_by_config([task], tmp_path / "logs")
    assert precheck["ok"] == 1 and precheck["warn"] == 0 and precheck["fail"] == 0
    checked = run_check_by_config([task], tmp_path / "logs")
    assert checked["task_ok"] == 1 and checked["marked_rows"] == 1
    deleted = run_delete_by_config([task], tmp_path / "logs")
    assert deleted["task_ok"] == 1 and deleted["deleted_rows"] == 1

    append_source = tmp_path / "append_source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["机构", "金额"])
    ws.append(["甲行", 99])
    ws.append(["乙行", 20])
    ws.append(["乙行", 21])
    _save(wb, append_source)
    target = tmp_path / "target.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(["机构", "金额"])
    ws.append(["甲行", 10])
    _save(wb, target)

    append_task = DedupTask(3, True, append_source, "Data", [1], target, "汇总")
    appended = run_append_by_config([append_task], tmp_path / "logs")
    assert appended["task_ok"] == 1
    assert appended["appended_rows"] == 1
    assert appended["deleted_rows"] == 2
    wb = load_workbook(target, data_only=True)
    assert list(wb["汇总"].values) == [("机构", "金额"), ("甲行", 10), ("乙行", 20)]
    wb.close()

    missing = DedupTask(4, True, source, "Missing", [1])
    invalid = DedupTask(5, True, tmp_path / "absent.xlsx", "Data", [1])
    precheck = run_precheck_by_config([missing, invalid], tmp_path / "logs")
    assert precheck == {"ok": 0, "warn": 0, "fail": 2}


def test_dedup_task_config_loader_resolves_relative_paths(tmp_path: Path) -> None:
    cfg = tmp_path / "config" / "config.xlsx"
    initialize_or_repair_config(cfg)
    source = cfg.parent.parent / "input" / "source.xlsx"
    source.parent.mkdir(parents=True)
    _dedup_book(source)
    with pd.ExcelWriter(cfg, engine="openpyxl", mode="a", if_sheet_exists="overlay") as writer:
        pd.DataFrame(
            [["是", "input/source.xlsx", "Data", "1;2", "output/target.xlsx", "汇总", "1", ""]],
            columns=DEDUP_TASK_COLS,
        ).to_excel(writer, sheet_name=SHEET_DEDUP_TASK, index=False)
    tasks = load_dedup_tasks(cfg)
    assert len(tasks) == 1
    assert tasks[0]["source_wb"] == source.resolve()
    assert tasks[0]["target_wb"] == (cfg.parent.parent / "output/target.xlsx").resolve()
    assert tasks[0]["key_cols"] == [1, 2]


def test_print_comment_modes_copy_expected_area_and_page_setup(tmp_path: Path) -> None:
    source = _print_source(tmp_path / "source.xlsx")
    out = tmp_path / "output"

    fast = run_print_fast_by_comment([source], out, tmp_path / "logs", hide_zero=True)
    assert fast["workbooks_hit"] == 1 and fast["sheets_hit"] == 1 and fast["saved_files"] == 1
    fast_file = next(out.glob("*_按批注打印（快速复制）.xlsx"))
    wb = load_workbook(fast_file, data_only=True)
    ws = wb["报表"]
    assert ws["A1"].value is None
    assert ws["B1"].value == "标题"
    assert ws.print_area == "'报表'!$A$1:$B$2"
    wb.close()

    kept = run_print_keep_by_comment([source], out, tmp_path / "logs")
    assert kept["workbooks_hit"] == 1 and kept["sheets_hit"] == 1 and kept["saved_files"] == 1
    keep_file = next(out.glob("*_按批注打印（保留源格式）.xlsx"))
    wb = load_workbook(keep_file)
    ws = wb["报表"]
    assert ws["B1"].fill.fgColor.rgb in ("00FFFF00", "FFFFFF00")
    assert ws.column_dimensions["A"].width == 18
    assert ws["A1"].comment is not None
    wb.close()


def test_print_config_loader_precheck_and_all_modes(tmp_path: Path) -> None:
    cfg = tmp_path / "config" / "config.xlsx"
    initialize_or_repair_config(cfg)
    source = _print_source(tmp_path / "source.xlsx")
    rows = []
    for mode in (1, 2, 3):
        rows.append(["是", mode, str(source), "报表", "A1:B2", f"out/mode{mode}.xlsx",
                     f"输出{mode}", 1, 1, "横向", "是", "否", "是", "是", "测试"])
    with pd.ExcelWriter(cfg, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        pd.DataFrame(rows, columns=PRINT_CONFIG_COLS).to_excel(
            writer, sheet_name=SHEET_PRINT_CONFIG, index=False
        )
    tasks = load_print_tasks(cfg)
    assert [task["mode"] for task in tasks] == [1, 2, 3]
    assert tasks[0]["drop_comments"] is True and tasks[0]["hide_zero"] is True
    checked = run_print_config_precheck(tasks, tmp_path / "logs")
    assert checked == {"ok": 3, "warn": 0, "fail": 0}

    ran = run_print_config_all_modes(tasks, tmp_path / "logs")
    assert ran["task_ok"] == 3 and ran["task_skip"] == 0
    assert ran["written_sheets"] == 3 and ran["written_rows"] == 6
    for mode in (1, 2, 3):
        out_path = tmp_path / "out" / f"mode{mode}.xlsx"
        wb = load_workbook(out_path, data_only=True)
        ws = wb[f"输出{mode}"]
        assert ws["A1"].value is None  # configured hide_zero
        assert ws["B1"].value == "标题"
        assert ws["A1"].comment is None  # configured drop_comments
        assert ws.page_setup.orientation == "landscape"
        wb.close()

    bad = dict(tasks[0], mode=9, source_ranges="not-a-range")
    checked = run_print_config_precheck([bad], tmp_path / "logs")
    assert checked["fail"] == 1


def test_summary_usedrange_and_comment_template(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "资产"
    ws.append(["机构", "数值", "数值"])
    ws.append(["甲行", 10, 11])
    _save(wb, source)

    used = run_summary_by_usedrange([source], tmp_path / "used", tmp_path / "logs")
    assert used["saved"] is True and used["rows"] == 2
    wb = load_workbook(used["path"], data_only=True)
    ws = wb["汇总"]
    assert [ws.cell(1, c).value for c in range(1, 7)] == [
        "工作簿", "工作表", "行号", "列1", "列2", "列3"
    ]
    assert [ws.cell(3, c).value for c in range(1, 7)] == ["source", "资产", 2, "甲行", 10, 11]
    wb.close()

    template = tmp_path / "template.xlsx"
    comment_source = tmp_path / "comment-source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "资产"
    ws["A1"], ws["D1"] = "标题", "广州"
    ws["A3"], ws["B3"], ws["C3"] = "甲行", 10, 11
    _save(wb, comment_source)

    wb = Workbook()
    ws = wb.active
    ws.title = "资产"
    ws["A3"].comment = Comment("行区域1", "test")
    ws["B3"].comment = Comment("行区域#1", "test")
    ws["B1"] = "机构"
    ws["B1"].comment = Comment("列区域1", "test")
    ws["C1"] = "金额"
    ws["C1"].comment = Comment("列区域#1", "test")
    ws["D1"] = "广州"
    ws["D1"].comment = Comment("set(地区)", "test")
    _save(wb, template)
    summary = run_summary_by_comment(template, [comment_source], tmp_path / "comment", tmp_path / "logs")
    assert summary["saved"] is True and summary["rows"] == 1
    wb = load_workbook(summary["path"], data_only=True)
    ws = wb["资产"]
    assert [ws.cell(1, c).value for c in range(1, 6)] == ["工作簿", "工作表", "地区", "机构", "金额"]
    assert [ws.cell(2, c).value for c in range(1, 7)] == ["comment-source", "资产", "广州", 10, 11, 3]
    wb.close()


def test_split_normalize_and_remove_foreign_bank(tmp_path: Path) -> None:
    current, previous, target = (tmp_path / name for name in ("current.xlsx", "previous.xlsx", "target.xlsx"))
    for path, values in ((current, (100, 80)), (previous, (80, 64))):
        wb = Workbook()
        ws = wb.active
        ws.title = "村镇银行本外币"
        ws["C5"], ws["D5"], ws["E5"] = values[0], "存款D", "存款E"
        ws["CI5"], ws["CJ5"], ws["CK5"] = values[1], "贷款J", "贷款K"
        _save(wb, path)
    wb = Workbook()
    ws = wb.active
    ws.title = "本外币分机构"
    ws.append(["机构", "c1", "c2", "c3", "c4", "c5", "c6", "c7", "c8"])
    ws.append(["村镇银行", 0, 0, 0, 0, 0, 0, 0, 0])
    _save(wb, target)
    split = run_split_village_bank(current, previous, [target], tmp_path / "logs")
    assert split["filled_sheets"] == 1 and split["failed_files"] == 0
    wb = load_workbook(target, data_only=True)
    ws = wb["本外币分机构"]
    assert ws["A2"].value == "博罗长江村镇银行"
    assert ws["A3"].value == "惠东惠民村镇银行"
    assert ws["E2"].value == 25.0 and ws["I2"].value == 25.0
    wb.close()

    institutions = tmp_path / "institutions.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "惠州分机构"
    ws.append(["机构名称", "备注"])
    ws.append(["简称甲", "x"])
    ws_other = wb.create_sheet("另一分机构")
    ws_other.append([None, "机构名称"])
    ws_other.append([None, "机构乙（分行）"])
    _save(wb, institutions)
    mapped = run_normalize_institution([institutions], {"简称甲": "标准甲", "机构乙": "标准乙"}, tmp_path / "logs")
    assert mapped["mapped_rows"] == 2 and mapped["failed_files"] == 0
    wb = load_workbook(institutions, data_only=True)
    assert wb["惠州分机构"]["A2"].value == "标准甲"
    assert wb["另一分机构"]["B2"].value == "标准乙"
    wb.close()

    foreign_source = tmp_path / "foreign.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "本外币分机构"
    ws.append(["外资甲", 1])
    ws.append(["本地行", 2])
    ws.append(["外资乙", 3])
    ws.append(["外资乙", 4])
    _save(wb, foreign_source)
    removed = run_remove_foreign_bank([foreign_source], {"外资甲", "外资乙"}, tmp_path / "logs")
    assert removed["deleted_rows"] == 3 and removed["ok_files"] == 1
    wb = load_workbook(foreign_source, data_only=True)
    assert list(wb["本外币分机构"].values) == [("外资甲", 1), ("本地行", 2), ("外资乙", 3), ("外资乙", 4)]
    wb.close()
    copied = load_workbook(tmp_path / "foreign(金融局).xlsx", data_only=True)
    assert list(copied["本外币分机构"].values) == [("本地行", 2)]
    copied.close()


def test_fx_header_region_sum_and_rural_loan_adjustment(tmp_path: Path) -> None:
    fx_file = tmp_path / "fx.xlsx"
    wb = Workbook()
    fx = wb.active
    fx.title = "外汇报表"
    fx.oddHeader.left.text = "单位：万元"
    other = wb.create_sheet("普通报表")
    other.oddHeader.left.text = "单位：万元"
    _save(wb, fx_file)
    fixed = run_fx_header_fix([fx_file], tmp_path / "logs")
    assert fixed["modified_sheets"] == 1 and fixed["saved_files"] == 1
    wb = load_workbook(fx_file)
    assert wb["外汇报表"].oddHeader.left.text == "单位：万美元"
    assert wb["普通报表"].oddHeader.left.text == "单位：万元"
    wb.close()

    correct = _region_book(tmp_path / "correct.xlsx")
    mismatch = _region_book(tmp_path / "mismatch.xlsx", mismatch=True)
    checked = run_region_sum_check([correct, mismatch], tmp_path / "region-out", tmp_path / "logs")
    assert checked["sheets"] == 2 and checked["errors"] == 1
    wb = load_workbook(checked["path"], data_only=True)
    assert wb["核对结果"]["D2"].value == "正确"
    assert wb["核对结果"]["D3"].value == "错误"
    wb.close()

    current, previous = tmp_path / "loan-current.xlsx", tmp_path / "loan-previous.xlsx"
    for path, amount in ((current, 110), (previous, 100)):
        wb = Workbook()
        ws = wb.active
        ws.title = "惠州市涉农贷款分机构表"
        ws.append(["序号", "机构名称", "涉农贷款", "涉农贷款比上月"])
        ws.append([])
        ws.append([1, "机构甲", amount, None])
        area = wb.create_sheet("本外币涉农贷款分地区")
        area.append(["地区", "涉农贷款", "涉农贷款比上月"])
        area.append(["惠州市", 999, 10])
        _save(wb, path)
    adjusted = run_adjust_rural_loan(current, previous, tmp_path / "logs")
    assert adjusted["error"] is None and adjusted["consistent"] is True
    assert adjusted["matched"] == 1 and adjusted["fill_total"] == adjusted["area_total"] == 10
    wb = load_workbook(current, data_only=True)
    assert wb["惠州市涉农贷款分机构表"]["D3"].value == 10
    wb.close()
