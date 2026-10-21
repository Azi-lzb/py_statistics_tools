from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, Side

from nocompytools.core import print_310


def _save(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def _source(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Marked"
    ws.append([0, "title", "extra"])
    ws.append(["row", 2, 3])
    ws["A1"].comment = Comment("打印区域 1；打印横向", "test")
    ws["B2"].comment = Comment("打印区域#1", "test")
    ws["C2"].comment = Comment("打印区域 bad", "test")
    ws["C1"].comment = Comment("打印区域 2", "test")
    ws["A1"].font = Font(bold=True)
    fallback = wb.create_sheet("Output")
    fallback.append(["output", 1])
    fallback["A1"].comment = Comment("输出纵向", "test")
    fallback.merge_cells("A1:B1")
    wb.create_sheet("Unmarked").append(["ignore"])
    ws["A1"].border = Border(bottom=Side(style="thin", color="FF000000"))
    return _save(wb, path)


def _task(source: Path, target: Path, **overrides):
    task = {
        "enabled": True,
        "row_no": 2,
        "mode": 1,
        "source_wb": source,
        "source_ws": "Marked",
        "source_ranges": "A1:B2",
        "target_wb": target,
        "target_ws": "Output",
        "fit_wide": 1,
        "fit_tall": 1,
        "orientation": "横向",
        "center_h": False,
        "center_v": False,
        "hide_zero": False,
        "drop_comments": False,
    }
    task.update(overrides)
    return task


def test_print_helpers_cover_orientation_bounds_ranges_zero_and_fallback_save(tmp_path: Path, monkeypatch) -> None:
    assert print_310._norm(None) == ""
    assert print_310._parse_orientation("") is None
    assert print_310._parse_orientation("横向") == "landscape"
    assert print_310._parse_orientation("纵向") == "portrait"
    assert print_310._parse_orientation("横向纵向") is None
    assert print_310._auto_orientation(5, 3) == "portrait"
    assert print_310._auto_orientation(2, 5) == "landscape"
    assert print_310._parse_ranges_text("") == []
    assert print_310._parse_ranges_text(" A1:B2；C3:D4 ") == [(1, 1, 2, 2), (3, 3, 4, 4)]
    with pytest.raises(ValueError):
        print_310._parse_ranges_text("not a range")
    assert print_310._is_zero_like(None) is False
    assert print_310._is_zero_like(True) is False
    assert print_310._is_zero_like(0) and print_310._is_zero_like("0.00")
    assert print_310._is_zero_like("1,000") is False
    assert print_310._is_zero_like("bad") is False
    assert print_310._is_zero_like(" ") is False
    assert print_310._normalize_out_value("0", True) is None
    assert print_310._normalize_out_value(0, False) == 0
    orientation_wb = Workbook()
    assert print_310._auto_orientation_by_range(orientation_wb.active, (1, 1, 1, 2)) == "landscape"
    assert print_310._auto_orientation_by_range(orientation_wb.active, (1, 1, 3, 1)) == "portrait"
    orientation_wb.close()

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "value"
    assert print_310._used_bounds(ws) == (1, 1, 1, 1)
    monkeypatch.setattr(ws, "calculate_dimension", lambda: (_ for _ in ()).throw(RuntimeError("bad dimension")))
    assert print_310._used_bounds(ws) == (1, 1, 1, 1)
    monkeypatch.setattr(ws, "calculate_dimension", lambda: "B2:A1")
    assert print_310._used_bounds(ws) == (1, 1, 1, 1)
    print_310._apply_print_setup(ws, 2, 3, 0, -1, center_h=True, center_v=True)
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
    assert ws.print_options.horizontalCentered and ws.print_options.verticalCentered
    print_310._apply_print_setup(ws, 1, 1)
    ws.sheet_properties.pageSetUpPr = None
    print_310._apply_print_setup(ws, 1, 1)
    wb.close()
    print_310._apply_sheet_order_footer(SimpleNamespace(worksheets=[]))

    preferred = tmp_path / "result.xlsx"
    preferred.write_text("occupied", encoding="utf-8")
    (tmp_path / "result_1.xlsx").write_text("occupied", encoding="utf-8")
    wb = Workbook()
    save = wb.save
    calls = 0

    def locked_once(path):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise PermissionError("locked")
        save(path)

    monkeypatch.setattr(wb, "save", locked_once)
    actual = print_310._save_with_suffix_fallback(wb, preferred)
    assert actual.name == "result_2_1.xlsx" and actual.is_file()
    wb.close()

    busy = Workbook()
    attempts = 0

    def always_locked(_path):
        nonlocal attempts
        attempts += 1
        raise PermissionError("still locked")

    monkeypatch.setattr(busy, "save", always_locked)
    with pytest.raises(PermissionError):
        print_310._save_with_suffix_fallback(busy, tmp_path / "busy.xlsx")
    assert attempts == 1000
    busy.close()


def test_comment_print_modes_skip_bad_sources_and_use_multiple_ranges(tmp_path: Path, monkeypatch) -> None:
    source = _source(tmp_path / "source.xlsx")
    no_mark = _save(Workbook(), tmp_path / "unmarked.xlsx")
    invalid = tmp_path / "unsupported.txt"
    read_sheets = []
    original_read_sheet = print_310.read_sheet_2d

    def read_with_calamine(path, sheet_name):
        read_sheets.append(sheet_name)
        return original_read_sheet(path, sheet_name)

    monkeypatch.setattr(print_310, "read_sheet_2d", read_with_calamine)

    fast = print_310.run_print_fast_by_comment(
        [source, no_mark, invalid, tmp_path / "missing.xlsx"], tmp_path / "fast", tmp_path / "logs", hide_zero=True
    )
    assert fast == {"workbooks_hit": 1, "sheets_hit": 2, "saved_files": 1, "skipped": 2}
    assert read_sheets == ["Marked", "Output"]
    fast_file = next((tmp_path / "fast").glob("*.xlsx"))
    wb = load_workbook(fast_file, data_only=True)
    try:
        assert wb.sheetnames == ["Marked", "Output"]
        assert wb["Marked"]["A1"].value is None
        assert wb["Marked"]["B1"].value == "title"
        assert wb["Marked"]["A1"].border.bottom.style == "thin"
        assert wb["Output"]["A1"].value == "output"
        assert "A1:B1" in {str(item) for item in wb["Output"].merged_cells.ranges}
    finally:
        wb.close()

    keep = print_310.run_print_keep_by_comment([source], tmp_path / "keep", tmp_path / "logs")
    assert keep["workbooks_hit"] == 1 and keep["sheets_hit"] == 2
    keep_skips = print_310.run_print_keep_by_comment(
        [no_mark, invalid, tmp_path / "missing.xlsx"], tmp_path / "keep-skips", tmp_path / "logs"
    )
    assert keep_skips == {"workbooks_hit": 0, "sheets_hit": 0, "saved_files": 0, "skipped": 2}


def test_print_config_precheck_accumulates_failures_and_warnings(tmp_path: Path) -> None:
    source = _source(tmp_path / "source.xlsx")
    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_text("not xlsx", encoding="utf-8")
    good = _task(source, tmp_path / "out.xlsx", fit_wide=0, fit_tall=-1, orientation="sideways")
    tasks = [
        dict(good, enabled=False),
        dict(good, mode=9),
        dict(good, source_wb=tmp_path / "missing.xlsx"),
        dict(good, source_wb=tmp_path / "unsupported.txt"),
        dict(good, source_ws=""),
        dict(good, target_wb=None),
        dict(good, target_wb=tmp_path / "out.csv"),
        dict(good, target_ws=""),
        dict(good, source_ranges="invalid"),
        dict(good, source_wb=corrupt),
        dict(good, source_ws="absent"),
        good,
    ]

    result = print_310.run_print_config_precheck(tasks, tmp_path / "logs")

    assert result == {"ok": 1, "warn": 6, "fail": 10}


def test_print_config_run_skips_bad_tasks_reuses_target_and_clears_old_sheet(tmp_path: Path, capsys) -> None:
    source = _source(tmp_path / "source.xlsx")
    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_text("not xlsx", encoding="utf-8")
    target = tmp_path / "existing.xlsx"
    old = Workbook()
    old.active.title = "Output"
    old.active.append(["old data"])
    _save(old, target)
    tasks = [
        dict(_task(source, target, row_no=0), enabled=False),
        _task(source, target, row_no=1, source_ranges="A1:B2", hide_zero=True),
        _task(source, target, row_no=2, mode=2, source_ws="Output", target_ws="Second", source_ranges="", orientation="sideways"),
        _task(source, target, row_no=3, source_wb=tmp_path / "missing.xlsx"),
        _task(source, tmp_path / "unsupported.csv", row_no=4),
        _task(source, target, row_no=5, source_ws="Missing"),
        _task(source, target, row_no=6, source_ranges="bad range"),
        _task(corrupt, target, row_no=7),
    ]

    result = print_310.run_print_config_all_modes(tasks, tmp_path / "logs")

    assert result["task_ok"] == 2 and result["task_skip"] == 5
    assert result["written_sheets"] == 2 and result["written_rows"] == 3
    assert len(result["slow_top"]) == 3
    wb = load_workbook(target, data_only=True)
    try:
        assert wb.sheetnames == ["Output", "Second"]
        assert wb["Output"]["A1"].value is None
        assert wb["Output"].max_row == 2
        assert wb["Second"]["A1"].value == "output"
    finally:
        wb.close()
    assert "skip=config_invalid" in capsys.readouterr().out


def test_config_print_mode_three_reads_values_with_calamine(tmp_path: Path, monkeypatch) -> None:
    source = _source(tmp_path / "source.xlsx")
    target = tmp_path / "target.xlsx"
    calls = []
    original_read_sheet = print_310.read_sheet_2d

    def read_with_calamine(path, sheet_name):
        calls.append((path, sheet_name))
        return original_read_sheet(path, sheet_name)

    monkeypatch.setattr(print_310, "read_sheet_2d", read_with_calamine)
    result = print_310.run_print_config_all_modes(
        [_task(source, target, mode=3, source_ws="Marked", target_ws="Printed", source_ranges="A1:B2")],
        tmp_path / "logs",
    )

    assert result["task_ok"] == 1
    assert calls == [(source, "Marked")]
    wb = load_workbook(target, data_only=True)
    try:
        assert wb["Printed"]["B2"].value == 2
        assert wb["Printed"]["A1"].border.bottom.style == "thin"
        assert "A1:B2" not in {str(item) for item in wb["Printed"].merged_cells.ranges}
    finally:
        wb.close()


def test_copy_block_preserves_merge_style_and_handles_merge_errors(monkeypatch) -> None:
    src_wb = Workbook()
    src = src_wb.active
    src["A1"] = "merged"
    src["A1"].font = Font(bold=True)
    src["A1"].comment = Comment("comment", "test")
    src.row_dimensions[1].height = 25
    src.merge_cells("A1:B1")
    val_wb = Workbook()
    values = val_wb.active
    values["A1"] = 0
    dst_wb = Workbook()
    dst = dst_wb.active
    rows, cols = print_310._copy_block(
        src, dst, (1, 1, 1, 2), 1, 2, value_ws=values, hide_zero=True, drop_comments=True
    )
    assert (rows, cols) == (1, 2)
    assert dst["A1"].value is None and dst["A1"].font.bold
    assert "A1:B1" in {str(r) for r in dst.merged_cells.ranges}
    assert dst["A1"].comment is None
    formatted = dst_wb.create_sheet("Formatted")
    print_310._copy_block(src, formatted, (1, 1, 1, 2), 1, 1)
    assert formatted.row_dimensions[1].height == 25
    assert formatted["A1"].comment.text == "comment"

    failing = dst_wb.create_sheet("Failing")
    monkeypatch.setattr(failing, "merge_cells", lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError("invalid merge")))
    print_310._copy_block(src, failing, (1, 1, 1, 2), 1, 2)
    print_310._copy_block(src, failing, (1, 1, 1, 2), 1, 3)
    assert failing["A1"].value == "merged"
    print_310._autofit_block_columns(failing, 1, 0, 2)
    print_310._autofit_block_columns(failing, 1, 1, 0)
    print_310._compact_block_row_heights(failing, 1, 0)
    src_wb.close()
    val_wb.close()
    dst_wb.close()
