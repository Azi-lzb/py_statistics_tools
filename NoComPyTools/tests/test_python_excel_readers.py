from pathlib import Path

import pandas as pd
import xlwt
from openpyxl import Workbook

from nocompytools.core.convert_33_34_35_37 import _read_cells, _read_cells_openpyxl
from nocompytools.core.io_excel import _excel_engine_for_path, list_sheet_names, read_sheet_2d
from nocompytools.core.rule_engine import _load_merge_ranges_file


def _xls(path: Path) -> Path:
    book = xlwt.Workbook()
    sheet = book.add_sheet("Data")
    sheet.write_merge(0, 0, 0, 1, "Merged Header")
    sheet.write(1, 0, "Value")
    sheet.write(1, 1, 17)
    book.save(str(path))
    return path


def test_calamine_reads_xls_values_sheets_cells_and_merged_ranges(tmp_path: Path) -> None:
    source = _xls(tmp_path / "legacy.xls")

    assert _excel_engine_for_path(str(source)) == "calamine"
    assert list_sheet_names(source) == ["Data"]
    assert read_sheet_2d(source, "Data").iloc[1].tolist() == ["Value", 17]
    assert _read_cells_openpyxl(source, [("Data", "A1"), ("Data", "B2")]) == ["Merged Header", "17"]
    assert _load_merge_ranges_file(str(source)) == {"Data": ((0, 0, 0, 1),)}


def test_calamine_reads_xlsb_and_ooxml_excel_formats(tmp_path: Path) -> None:
    assert _excel_engine_for_path("legacy.xlsb") == "calamine"
    assert _excel_engine_for_path("report.xlsx") == "calamine"
    assert _excel_engine_for_path("report.xlsm") == "calamine"

    path = tmp_path / "report.xlsx"
    book = Workbook()
    book.active.title = "Data"
    book.active.merge_cells("A1:B1")
    book.active["A1"] = "Merged Header"
    book.active.append(["Value", 21])
    book.save(path)
    book.close()
    assert pd.read_excel(path, engine="calamine", header=None).iloc[0, 0] == "Merged Header"
    assert list_sheet_names(path) == ["Data"]
    assert read_sheet_2d(path, "Data").iloc[1].tolist() == ["Value", 21]
    assert _read_cells(path, [("Data", "A1"), ("Data", "B2"), ("missing", "C3")]) == [
        "Merged Header", "21", "missing!C3",
    ]
    assert _read_cells_openpyxl(path, [("Data", "A1")]) == ["Merged Header"]
    assert _load_merge_ranges_file(str(path)) == {"Data": ((0, 0, 0, 1),)}

    macro_path = tmp_path / "report.xlsm"
    book = Workbook()
    book.active.title = "Data"
    book.active["A1"] = "Macro-enabled values"
    book.save(macro_path)
    book.close()
    assert pd.read_excel(macro_path, engine="calamine", header=None).iloc[0, 0] == "Macro-enabled values"
