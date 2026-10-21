from __future__ import annotations

from pathlib import Path

import pytest
import xlwt
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from nocompytools.core import summary_321_322 as summary_module
from nocompytools.core.summary_321_322 import run_summary_by_comment, run_summary_by_usedrange


def _summary_template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "模板"
    for address, value, text in (
        ("A1", "行开始", "行区域1"),
        ("A3", "行结束", "行区域#1"),
        ("B1", "指标", "列区域1"),
        ("B2", "数值", "列区域#1"),
    ):
        ws[address] = value
        ws[address].comment = Comment(text, "test")
    wb.save(path)
    wb.close()
    return path


def _summary_source(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "回复"
    ws.append(["单位", "指标值"])
    ws.append(["甲", 10])
    ws.append(["乙", 20])
    short = wb.create_sheet("短表")
    short["A1"] = "只有一列"
    wb.save(path)
    wb.close()
    return path


def _xls_summary_source(path: Path) -> Path:
    wb = xlwt.Workbook()
    ws = wb.add_sheet("回复")
    for row, values in enumerate((("单位", "指标值"), ("甲", 10), ("乙", 20))):
        for col, value in enumerate(values):
            ws.write(row, col, value)
    wb.save(str(path))
    return path


def test_usedrange_summary_accepts_valid_files_and_skips_missing_inputs(tmp_path: Path) -> None:
    source = _summary_source(tmp_path / "source.xlsx")
    missing = tmp_path / "missing.xlsx"
    unsupported = tmp_path / "ignored.csv"
    unsupported.touch()

    result = run_summary_by_usedrange(
        [missing, unsupported, source], tmp_path / "out", tmp_path / "logs"
    )

    assert result["saved"] is True and result["rows"] == 4
    wb = load_workbook(result["path"], data_only=True)
    try:
        ws = wb["汇总"]
        assert ws.max_column == 5
        assert ws["A2"].value == "source" and ws["B2"].value == "回复"
        assert ws["E3"].value == 10
        assert ws["E5"].value is None
    finally:
        wb.close()


def test_usedrange_summary_reads_xls_with_calamine(tmp_path: Path, monkeypatch) -> None:
    source = _xls_summary_source(tmp_path / "source.xls")
    monkeypatch.setattr(
        summary_module, "load_workbook",
        lambda *_args, **_kwargs: pytest.fail("只读汇总不应使用 openpyxl 读取源文件"),
    )

    result = run_summary_by_usedrange([source], tmp_path / "out", tmp_path / "logs")

    assert result["saved"] is True and result["rows"] == 3
    wb = load_workbook(result["path"], data_only=True)
    try:
        ws = wb["汇总"]
        assert ws["A2"].value == "source" and ws["B2"].value == "回复"
        assert ws["D2"].value == "单位" and ws["E4"].value == 20
    finally:
        wb.close()


def test_comment_summary_uses_template_fallback_for_unmatched_sheet(tmp_path: Path) -> None:
    template = _summary_template(tmp_path / "template.xlsx")
    source = _summary_source(tmp_path / "source.xlsx")

    result = run_summary_by_comment(
        template,
        [tmp_path / "missing.xlsx", tmp_path / "ignored.csv", source],
        tmp_path / "out",
        tmp_path / "logs",
    )

    assert result["saved"] is True and result["rows"] == 6
    wb = load_workbook(result["path"], data_only=True)
    try:
        ws = wb["模板"]
        assert [ws.cell(1, col).value for col in range(1, 5)] == [
            "工作簿", "工作表", "指标_数值", "行号",
        ]
        assert [ws.cell(2, col).value for col in range(1, 5)] == [
            "source", "回复", "指标值", 1,
        ]
        assert ws["C3"].value == 10
        assert ws["A5"].value == "source" and ws["B5"].value == "短表"
    finally:
        wb.close()


def test_comment_summary_reads_source_with_calamine_and_resolves_merges(
    tmp_path: Path, monkeypatch,
) -> None:
    template = Workbook()
    tmpl_ws = template.active
    tmpl_ws.title = "模板"
    for address, value, text in (
        ("A3", "行开始", "行区域1"),
        ("A4", "行结束", "行区域#1"),
        ("B1", "指标", "列区域1"),
        ("B2", "数值", "列区域#1"),
        ("C1", "地区", "set(地区)"),
    ):
        tmpl_ws[address] = value
        tmpl_ws[address].comment = Comment(text, "test")
    template_path = tmp_path / "template.xlsx"
    template.save(template_path)
    template.close()

    source = Workbook()
    src_ws = source.active
    src_ws.title = "回复"
    src_ws["B1"] = "华南"
    src_ws.merge_cells("B1:C1")
    src_ws["A3"], src_ws["B3"] = "甲", 10
    src_ws["A4"], src_ws["B4"] = "乙", 20
    source_path = tmp_path / "source.xlsx"
    source.save(source_path)
    source.close()

    original_load_workbook = summary_module.load_workbook

    def template_only(path, *args, **kwargs):
        assert Path(path) == template_path, "批注汇总只允许 openpyxl 打开配置模板"
        return original_load_workbook(path, *args, **kwargs)

    monkeypatch.setattr(summary_module, "load_workbook", template_only)
    result = run_summary_by_comment(
        template_path, [source_path], tmp_path / "out", tmp_path / "logs",
    )

    assert result["saved"] is True and result["rows"] == 2
    wb = load_workbook(result["path"], data_only=True)
    try:
        ws = wb["模板"]
        assert [ws.cell(2, col).value for col in range(1, 6)] == [
            "source", "回复", "华南", 10, 3,
        ]
        assert [ws.cell(3, col).value for col in range(1, 6)] == [
            "source", "回复", "华南", 20, 4,
        ]
    finally:
        wb.close()
