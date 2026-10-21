"""是否输出行序号：解析开关 + 宽表汇总按开关输出"行序号"列。"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from nocompytools.core.config_xlsx import (
    SHEET_TIMELINE_RULE, TIMELINE_COLS,
    GlobalConfig, load_timeline_rules,
)
from nocompytools.core.wide_397_398 import run_wide_summary


def _make_config(cfg_path: Path, rules: list[dict]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    g_ws = wb.create_sheet("全局配置")
    for r, (k, v) in enumerate([
        ("输入目录", "in"), ("输出目录", "out"), ("日志目录", "log"),
        ("错误策略", "continue"), ("源文件扩展名", ".xlsx"),
    ], start=1):
        g_ws.cell(r, 1, k)
        g_ws.cell(r, 2, v)
    t_ws = wb.create_sheet(SHEET_TIMELINE_RULE)
    for c, name in enumerate(TIMELINE_COLS, start=1):
        t_ws.cell(1, c, name)
    for r, rule in enumerate(rules, start=2):
        for c, name in enumerate(TIMELINE_COLS, start=1):
            t_ws.cell(r, c, rule.get(name, ""))
    wb.save(cfg_path)


def _rule(name: str, **extra) -> dict:
    rule = {c: "" for c in TIMELINE_COLS}
    rule.update({
        "是否启用": "是",
        "规则名称": name,
        "工作簿关键字": "测试",
        "工作表关键字": "Sheet1",
        "行头列": 1,
        "列表头行": "2",
        "必含列头": "本月余额",
        "数据起始行": 3,
        "数据起始列": 2,
    })
    rule.update(extra)
    return rule


def _make_source(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    rows = [
        ["", ""],
        ["", "本月余额"],
        ["甲", 1],
        ["其他情况导致", 10],
        ["乙", 2],
        ["其他情况导致", 20],
    ]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    wb.save(path)


@pytest.fixture
def env(tmp_path: Path):
    src = tmp_path / "测试机构.xlsx"
    _make_source(src)
    g = GlobalConfig(input_dir=tmp_path, output_dir=tmp_path / "out", log_dir=tmp_path / "log")
    return tmp_path, src, g


def test_parse_output_row_no_flag(tmp_path: Path):
    cfg = tmp_path / "config.xlsx"
    _make_config(cfg, [
        _rule("开关开", **{"是否输出行序号": "是"}),
        _rule("开关关"),
    ])
    rules = load_timeline_rules(cfg)
    assert [r.output_row_no for r in rules] == [True, False]


def test_wide_row_no_column_on(env):
    tmp_path, src, g = env
    cfg = tmp_path / "config.xlsx"
    _make_config(cfg, [_rule("规则A", **{"是否输出行序号": "是"})])
    rules = load_timeline_rules(cfg)

    out = run_wide_summary(rules, [], g, row_suffix_enabled=False, source_paths=[src])
    assert out is not None
    df = pd.read_excel(out, sheet_name="规则A")
    assert "行序号" in df.columns
    dup = df[df["行头路径"] == "其他情况导致"]
    assert len(dup) == 2
    assert sorted(dup["行序号"].tolist()) == [4, 6]


def test_wide_row_no_column_off(env):
    tmp_path, src, g = env
    cfg = tmp_path / "config.xlsx"
    _make_config(cfg, [_rule("规则B")])
    rules = load_timeline_rules(cfg)

    out = run_wide_summary(rules, [], g, row_suffix_enabled=False, source_paths=[src])
    assert out is not None
    df = pd.read_excel(out, sheet_name="规则B")
    assert "行序号" not in df.columns
    assert len(df[df["行头路径"] == "其他情况导致"]) == 2
