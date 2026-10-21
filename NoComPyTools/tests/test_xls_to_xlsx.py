from pathlib import Path
import shutil
import sys
import types

import pytest
import xlwt
from openpyxl import load_workbook

from nocompytools.core import xls_to_xlsx as convert


def _has_libreoffice() -> bool:
    return bool(
        shutil.which("soffice")
        or shutil.which("libreoffice")
        or any(Path("/opt").glob("libreoffice*/program/soffice"))
    )


def _xls(path: Path) -> Path:
    book = xlwt.Workbook()
    sheet = book.add_sheet("Data")
    sheet.write_merge(0, 0, 0, 1, "Merged Header")
    sheet.write(1, 0, "Value")
    sheet.write(1, 1, 17)
    book.save(str(path))
    return path


def test_linux_uses_libreoffice_and_duplicate_outputs_get_suffixes(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(convert, "sys", types.SimpleNamespace(platform="linux"))
    monkeypatch.setattr(convert, "_libreoffice_executable", lambda: "soffice")
    source = _xls(tmp_path / "source.xls")
    xlsx = tmp_path / "other.xlsx"
    xlsx.touch()
    calls = []

    def fake_convert(_exe, src, out):
        calls.append(src)
        out.write_bytes(b"xlsx")

    monkeypatch.setattr(convert, "_convert_one", fake_convert)
    result = convert.run_xls_to_xlsx([source, xlsx], tmp_path / "logs")

    assert result == {"ok": 1, "skip": 1, "target_ext": "xlsx", "engine": "LibreOffice"}
    assert calls == [source]
    assert (tmp_path / "xlsx" / "source.xlsx").read_bytes() == b"xlsx"

    result = convert.run_xls_to_xlsx([source], tmp_path / "logs")
    assert result["ok"] == 1
    assert (tmp_path / "xlsx" / "source_1.xlsx").is_file()


def test_empty_input_does_not_require_libreoffice(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(convert, "_libreoffice_executable", lambda: pytest.fail("unexpected LibreOffice discovery"))
    result = convert.run_xls_to_xlsx([tmp_path / "missing.xls"], tmp_path / "logs")
    assert result == {"ok": 0, "skip": 1, "target_ext": "xlsx", "engine": ""}


def test_windows_uses_excel_or_wps_com(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(convert, "sys", types.SimpleNamespace(platform="win32"))
    source = _xls(tmp_path / "source.xls")
    opened = []
    saves = []

    class FakeWorkbook:
        def SaveAs(self, output, **kwargs):
            saves.append((Path(output), kwargs))
            Path(output).write_bytes(b"converted")

        def Close(self, **kwargs):
            assert kwargs == {"SaveChanges": False}

    class FakeApp:
        Workbooks = types.SimpleNamespace(
            Open=lambda path, **kwargs: opened.append((Path(path), kwargs)) or FakeWorkbook(),
        )

        def Quit(self):
            self.quit = True

    app = FakeApp()
    monkeypatch.setattr(convert, "_start_windows_office", lambda: (app, "ket.Application"))

    result = convert.run_xls_to_xlsx([source], tmp_path / "logs")

    assert result == {"ok": 1, "skip": 0, "target_ext": "xlsx", "engine": "ket.Application"}
    assert opened == [(source.resolve(), {"UpdateLinks": 0, "ReadOnly": True})]
    assert saves == [(tmp_path / "xlsx" / "source.xlsx", {"FileFormat": 51, "CreateBackup": False})]
    assert app.quit


def test_windows_office_dispatch_falls_back_from_excel_to_wps(monkeypatch) -> None:
    calls = []
    app = types.SimpleNamespace()

    def dispatch(prog_id):
        calls.append(prog_id)
        if prog_id == "Excel.Application":
            raise OSError("Excel not installed")
        return app

    package = types.ModuleType("win32com")
    package.__path__ = []
    client = types.ModuleType("win32com.client")
    client.DispatchEx = dispatch
    package.client = client
    monkeypatch.setitem(sys.modules, "win32com", package)
    monkeypatch.setitem(sys.modules, "win32com.client", client)

    started, engine = convert._start_windows_office()

    assert started is app
    assert engine == "ket.Application"
    assert calls == ["Excel.Application", "ket.Application"]
    assert app.Visible is False and app.DisplayAlerts is False
    assert app.ScreenUpdating is False and app.EnableEvents is False


@pytest.mark.skipif(sys.platform != "linux" or not _has_libreoffice(), reason="Linux LibreOffice not installed")
def test_libreoffice_converts_real_xls_values_and_merged_headers(tmp_path: Path) -> None:
    source = _xls(tmp_path / "actual.xls")

    result = convert.run_xls_to_xlsx([source], tmp_path / "logs")

    output = tmp_path / "xlsx" / "actual.xlsx"
    assert result["ok"] == 1
    book = load_workbook(output, data_only=True)
    try:
        sheet = book["Data"]
        assert sheet["A1"].value == "Merged Header"
        assert sheet["B2"].value == 17
        assert "A1:B1" in {str(item) for item in sheet.merged_cells.ranges}
    finally:
        book.close()
