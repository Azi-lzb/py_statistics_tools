from pathlib import Path

from openpyxl import Workbook, load_workbook

from PyCode.core.feature_1_9_adjust_rural_loan_quarterly import run_adjust_rural_loan_quarterly


def _add_institution_sheet(wb: Workbook, title: str, rows: list[tuple[str, float]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(["序号", "机构名称", "涉农贷款", "涉农贷款比上月"])
    ws.append([])
    for index, (name, value) in enumerate(rows, start=1):
        ws.append([index, name, value, None])


def _make_workbook(path: Path, city_values, district_values, area_values=None) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    _add_institution_sheet(wb, "惠州市-涉农贷款分机构表", city_values)
    _add_institution_sheet(wb, "惠阳区涉农贷款分机构表", district_values)
    if area_values is not None:
        ws = wb.create_sheet("涉农贷款分地区统计表")
        ws.append(["地区", "涉农贷款", "涉农贷款比上月"])
        for region, value in area_values.items():
            ws.append([region, None, value])
    wb.save(path)
    wb.close()


def test_adjust_quarterly_rural_loan_by_region(tmp_path):
    current = tmp_path / "本期.xlsx"
    previous = tmp_path / "上期.xlsx"
    _make_workbook(
        current,
        [("机构甲", 100), ("机构乙", 30)],
        [("机构丙", 80)],
        {"惠州市": 7, "惠阳区": 2},
    )
    _make_workbook(previous, [("机构甲", 95), ("机构乙", 28)], [("机构丙", 78)])

    result = run_adjust_rural_loan_quarterly(current, previous, tmp_path / "logs")

    assert result["error"] is None
    assert result["consistent"] is True
    assert result["matched"] == 3
    assert result["fill_total"] == 9
    assert result["area_total"] == 9
    assert [item["region"] for item in result["regions"]] == ["惠州市", "惠阳区"]

    wb = load_workbook(current, data_only=True)
    assert wb["惠州市-涉农贷款分机构表"].cell(3, 4).value == 5
    assert wb["惠州市-涉农贷款分机构表"].cell(4, 4).value == 2
    assert wb["惠阳区涉农贷款分机构表"].cell(3, 4).value == 2
    wb.close()
