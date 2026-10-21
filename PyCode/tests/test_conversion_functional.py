from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from PyCode.core.config_xlsx import CONFIG_RENAME_COLS
from PyCode.core.convert_33_34_35_37 import (
    run_batch_rename_33,
    run_batch_rename_sheet_37,
)


def _make_rename_config(path: Path) -> Path:
    rows = [
        ["建行", "中国建设银行", "", "0001", "中国建设银行", "", "数据日期", "2026-01-31", "",
         "旧表", "新表", "", "旧片段", "新片段", "", "", ""],
        ["", "", "", "", "", "", "报表名称", "存款表", "", "", "", "", "", "", "", "", ""],
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(rows, columns=CONFIG_RENAME_COLS).to_excel(
            writer, sheet_name="重命名配置", index=False
        )
    return path


def test_batch_file_rename_and_sheet_rename_from_shared_config(tmp_path: Path) -> None:
    config = _make_rename_config(tmp_path / "config.xlsx")
    workbook = tmp_path / "建行_报表.xlsx"
    book = Workbook()
    book.active.title = "旧表"
    book.active["A1"] = "data"
    book.save(workbook)
    book.close()

    renamed = run_batch_rename_33(config, [workbook], tmp_path / "logs")
    target = tmp_path / "2026-01-31 0001 中国建设银行 存款表.xlsx"
    assert renamed == {"ok": 1, "skip": 0}
    assert target.exists() and not workbook.exists()

    changed = run_batch_rename_sheet_37(config, [target], tmp_path / "logs")
    assert changed["workbooks"] == 1 and changed["rename_ok"] == 1
    book = load_workbook(target, data_only=True)
    assert book.sheetnames == ["新表"]
    assert book["新表"]["A1"].value == "data"
    book.close()


def test_sheet_rename_skips_duplicate_or_invalid_names(tmp_path: Path) -> None:
    config = _make_rename_config(tmp_path / "config.xlsx")
    workbook = tmp_path / "sheets.xlsx"
    book = Workbook()
    book.active.title = "旧表"
    book.create_sheet("新表")
    book.save(workbook)
    book.close()

    result = run_batch_rename_sheet_37(config, [workbook], tmp_path / "logs")
    assert result["rename_ok"] == 0 and result["skip"] == 1
    book = load_workbook(workbook, read_only=True)
    assert book.sheetnames == ["旧表", "新表"]
    book.close()


def test_rename_config_requires_all_batch_rename_inputs(tmp_path: Path) -> None:
    config = tmp_path / "invalid.xlsx"
    with pd.ExcelWriter(config, engine="openpyxl") as writer:
        pd.DataFrame([["简称", "全称"]], columns=["简称", "全称"]).to_excel(
            writer, sheet_name="重命名配置", index=False
        )
    with pytest.raises(ValueError, match="缺少必要配置"):
        run_batch_rename_33(config, [], tmp_path / "logs")


def test_nonoffice_excel_xlsx_roundtrip_preserves_all_sheets(tmp_path: Path, monkeypatch) -> None:
    from PyCode.core.convert_33_34_35_37 import run_excel_convert_34

    monkeypatch.setenv("PYTOOLS_MODE", "python")
    source = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "第一表"
    wb.active.append(["name", "value"])
    wb.active.append(["alpha", 7])
    second = wb.create_sheet("第二表")
    second.append(["code", "enabled"])
    second.append(["B", True])
    wb.save(source)
    wb.close()

    result = run_excel_convert_34([source], "xlsx", tmp_path / "logs")
    output = tmp_path / "xlsx" / "book.xlsx"
    assert result["engine"] == "Python" and result["ok"] == 1
    converted = load_workbook(output, data_only=True)
    assert converted.sheetnames == ["第一表", "第二表"]
    assert converted["第一表"]["A2"].value == "alpha"
    assert converted["第二表"]["B2"].value is True
    converted.close()


def _make_minimal_docx(path: Path) -> None:
    content_types = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>'''
    relationships = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>'''
    document = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
<w:p><w:r><w:t>pytools office conversion regression</w:t></w:r></w:p><w:sectPr/>
</w:body></w:document>'''
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document)


@pytest.mark.skipif(sys.platform != "win32", reason="Windows Word COM integration")
def test_office_word_conversion_uses_word_com_when_available(tmp_path: Path, monkeypatch) -> None:
    from PyCode.core import convert_33_34_35_37 as convert

    monkeypatch.setenv("PYTOOLS_MODE", "office")
    source = tmp_path / "source.docx"
    _make_minimal_docx(source)
    try:
        stat = convert.run_word_convert_36([source], "doc", tmp_path / "logs")
    except RuntimeError as exc:
        pytest.skip(f"Microsoft Word/WPS COM is unavailable: {exc}")
    assert stat["ok"] == 1 and stat["skip"] == 0
    assert stat["engine"] in {"Word.Application", "kwps.Application", "KWPS.Application", "wps.Application"}
    assert (tmp_path / "doc" / "source.doc").is_file()
