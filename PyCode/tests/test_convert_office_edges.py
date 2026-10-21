from __future__ import annotations

import sys
import types
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from PyCode.core import convert_33_34_35_37 as convert


def _workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["name", "value"])
    ws.append(["alpha", 7])
    wb.save(path)
    wb.close()


def test_python_mode_converts_supported_sources_and_handles_names(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("PYTOOLS_MODE", "python")
    xlsx = tmp_path / "book.xlsx"
    xlsm = tmp_path / "macro.xlsm"
    csv = tmp_path / "table.csv"
    _workbook(xlsx)
    _workbook(xlsm)
    csv.write_text("name,value\nbeta,9\n", encoding="utf-8")

    result_xlsx = convert.run_excel_convert_34([xlsx, xlsm, csv], "xlsx", tmp_path / "logs")
    result_csv = convert.run_excel_convert_34([xlsx, csv], "csv", tmp_path / "logs")

    assert (result_xlsx["ok"], result_xlsx["skip"]) == (3, 0)
    assert load_workbook(tmp_path / "xlsx" / "book.xlsx", data_only=True)["Data"]["A2"].value == "alpha"
    assert (tmp_path / "xlsx" / "macro.xlsx").is_file()
    assert (tmp_path / "xlsx" / "table.xlsx").is_file()
    assert (result_csv["ok"], result_csv["skip"]) == (2, 0)
    assert pd.read_csv(tmp_path / "csv" / "book.csv").iloc[0].to_dict() == {"name": "alpha", "value": 7}
    assert (tmp_path / "csv" / "table.csv").read_text(encoding="utf-8") == csv.read_text(encoding="utf-8")


@pytest.mark.parametrize("source_ext", [".xls", ".xlsb"])
def test_python_mode_imports_legacy_excel_sheets_with_unique_safe_names(
    tmp_path, monkeypatch, source_ext
) -> None:
    monkeypatch.setenv("PYTOOLS_MODE", "python")
    source = tmp_path / f"legacy{source_ext}"
    source.write_bytes(b"fixture")

    class FakeExcelFile:
        sheet_names = ["Bad/Name", "bad:name"]

        def __init__(self, _path):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def parse(self, sheet_name):
            return pd.DataFrame({"sheet": [sheet_name]})

    monkeypatch.setattr(convert.pd, "ExcelFile", FakeExcelFile)
    result = convert.run_excel_convert_34([source], "xlsx", tmp_path / "logs")

    wb = load_workbook(tmp_path / "xlsx" / "legacy.xlsx", data_only=True)
    try:
        assert wb.sheetnames == ["Bad_Name", "bad_name_1"]
        assert wb["bad_name_1"]["A2"].value == "bad:name"
    finally:
        wb.close()
    assert result["ok"] == 1


def test_python_conversion_skips_bad_sources_and_reserves_duplicate_targets(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("PYTOOLS_MODE", "python")
    source = tmp_path / "source.xlsx"
    invalid = tmp_path / "unsupported.et"
    missing = tmp_path / "missing.xlsx"
    _workbook(source)
    invalid.write_text("fixture", encoding="utf-8")
    target_dir = tmp_path / "xlsx"
    target_dir.mkdir()
    (target_dir / "source.xlsx").write_text("existing", encoding="utf-8")

    result = convert.run_excel_convert_34([source, invalid, missing], "unknown", tmp_path / "logs")

    assert result == {"ok": 1, "skip": 2, "target_ext": "xlsx", "engine": "Python"}
    assert (target_dir / "source_1.xlsx").is_file()
    with pytest.raises(RuntimeError, match="仅支持 .xlsx 和 .csv"):
        convert.run_excel_convert_34([source], "xls", tmp_path / "logs")


def test_empty_excel_conversion_does_not_start_office(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("PYTOOLS_MODE", "office")
    monkeypatch.setattr(convert, "_start_excel_app", lambda **_kwargs: pytest.fail("unexpected COM startup"))
    assert convert.run_excel_convert_34([tmp_path / "missing.xlsx"], "xlsx", tmp_path / "logs") == {
        "ok": 0, "skip": 1, "target_ext": "xlsx", "engine": ""
    }


def test_dispatch_first_tries_each_registered_program_id(monkeypatch) -> None:
    calls = []
    client = types.ModuleType("win32com.client")

    def dispatch(progid):
        calls.append(progid)
        if progid == "broken.Application":
            raise OSError("not registered")
        return SimpleNamespace()

    client.DispatchEx = dispatch
    package = types.ModuleType("win32com")
    package.client = client
    monkeypatch.setitem(sys.modules, "win32com", package)
    monkeypatch.setitem(sys.modules, "win32com.client", client)

    app, engine = convert._dispatch_first(["broken.Application", "good.Application"])

    assert engine == "good.Application"
    assert calls == ["broken.Application", "good.Application"]
    assert app is not None

    with pytest.raises(RuntimeError, match="无法创建 COM 应用"):
        convert._dispatch_first(["broken.Application"])


def test_dispatch_first_reports_missing_pywin32(monkeypatch) -> None:
    monkeypatch.setitem(sys.modules, "win32com", None)
    monkeypatch.setitem(sys.modules, "win32com.client", None)
    with pytest.raises(RuntimeError, match="缺少 pywin32"):
        convert._dispatch_first(["Excel.Application"])


def test_start_excel_app_sets_quiet_defaults_and_honors_wps_preference(monkeypatch) -> None:
    app = SimpleNamespace()
    calls = []
    monkeypatch.setattr(convert, "_dispatch_first", lambda ids: (calls.append(ids) or (app, ids[0])))

    returned, engine = convert._start_excel_app(prefer_wps=True)

    assert returned is app and engine == "ket.Application"
    assert calls[0][0:2] == ["ket.Application", "KET.Application"]
    assert app.Visible is False and app.DisplayAlerts is False
    assert app.ScreenUpdating is False and app.EnableEvents is False
    assert app.Calculation == -4135

    class PartlyReadOnlyApp:
        Visible = True
        DisplayAlerts = True

        def __setattr__(self, name, value):
            if name in {"ScreenUpdating", "Calculation"}:
                raise OSError("unsupported property")
            object.__setattr__(self, name, value)

    fallback = PartlyReadOnlyApp()
    monkeypatch.setattr(convert, "_dispatch_first", lambda _ids: (fallback, "Excel.Application"))
    assert convert._start_excel_app()[0] is fallback
    assert fallback.Visible is False and fallback.DisplayAlerts is False


def test_com_excel_conversion_closes_each_workbook_and_quits(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("PYTOOLS_MODE", "office")
    monkeypatch.setattr(convert.sys, "platform", "win32")
    source = tmp_path / "source.et"
    failing = tmp_path / "bad.xlsx"
    source.write_text("wps", encoding="utf-8")
    failing.write_text("excel", encoding="utf-8")
    opened = []

    class FakeBook:
        def __init__(self, path):
            self.path = path

        def SaveAs(self, out, **kwargs):
            if self.path.name == "bad.xlsx":
                raise OSError("save failed")
            Path(out).write_text("converted", encoding="utf-8")
            assert kwargs["CreateBackup"] is False

        def Close(self, **kwargs):
            assert kwargs == {"SaveChanges": False}

    class FakeBooks:
        def Open(self, path, **_kwargs):
            book = FakeBook(Path(path))
            opened.append(book.path.name)
            return book

    app = SimpleNamespace(Workbooks=FakeBooks(), Quit=lambda: opened.append("quit"))
    preferences = []
    monkeypatch.setattr(convert, "_start_excel_app", lambda prefer_wps=False: (preferences.append(prefer_wps) or (app, "KET.Application")))

    result = convert.run_excel_convert_34([source, failing], "xlsx", tmp_path / "logs")

    assert result == {"ok": 1, "skip": 1, "target_ext": "xlsx", "engine": "KET.Application"}
    assert preferences == [True]
    assert opened == ["source.et", "bad.xlsx", "quit"]


def test_libreoffice_discovery_and_conversion_failures(tmp_path, monkeypatch) -> None:
    source = tmp_path / "book.xlsx"
    source.write_text("source", encoding="utf-8")
    executable = tmp_path / "soffice"
    executable.touch()
    monkeypatch.setattr(convert.shutil, "which", lambda name: str(executable) if name == "soffice" else None)
    assert convert._libreoffice_executable() == str(executable)

    monkeypatch.setattr(convert.shutil, "which", lambda _name: None)
    original_glob = Path.glob

    def no_install(path, pattern):
        if str(path) == "/opt" and pattern == "libreoffice*/program/soffice":
            return iter([])
        return original_glob(path, pattern)

    monkeypatch.setattr(Path, "glob", no_install)
    with pytest.raises(RuntimeError, match="需要 LibreOffice"):
        convert._libreoffice_executable()
    monkeypatch.setattr(convert.shutil, "which", lambda name: str(executable) if name == "soffice" else None)

    def successful_run(command, **_kwargs):
        output = Path(command[command.index("--outdir") + 1]) / "book.csv"
        output.write_text("name,value\nalpha,7\n", encoding="utf-8")
        return SimpleNamespace(returncode=0, stderr="", stdout="")

    monkeypatch.setattr(convert.subprocess, "run", successful_run)
    output = tmp_path / "converted.csv"
    convert._libreoffice_convert_one(str(executable), source, output, "csv", "Text - txt - csv (StarCalc)")
    assert output.read_text(encoding="utf-8").startswith("name,value")

    for result, message in (
        (SimpleNamespace(returncode=1, stderr="bad input", stdout=""), "bad input"),
        (SimpleNamespace(returncode=0, stderr="", stdout=""), "未生成 book.csv"),
    ):
        monkeypatch.setattr(convert.subprocess, "run", lambda *_args, _result=result, **_kwargs: _result)
        with pytest.raises(RuntimeError, match=message):
            convert._libreoffice_convert_one(str(executable), source, tmp_path / "failed.csv", "csv", "filter")


def test_linux_office_excel_uses_libreoffice_and_rejects_unmapped_formats(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("PYTOOLS_MODE", "office")
    monkeypatch.setattr(convert.sys, "platform", "linux")
    monkeypatch.setattr(convert, "_libreoffice_executable", lambda: "soffice")
    source = tmp_path / "source.xlsx"
    _workbook(source)
    calls = []

    def fake_convert(executable, src, out, target_ext, output_filter):
        calls.append((executable, src.suffix, target_ext, output_filter))
        if src.stem == "bad":
            raise RuntimeError("fixture conversion failure")
        out.write_bytes(b"converted")

    monkeypatch.setattr(convert, "_libreoffice_convert_one", fake_convert)
    bad = tmp_path / "bad.xlsx"
    _workbook(bad)
    result = convert.run_excel_convert_34([source, bad], "xlsx", tmp_path / "logs")

    assert result == {"ok": 1, "skip": 1, "target_ext": "xlsx", "engine": "LibreOffice"}
    assert calls[0] == ("soffice", ".xlsx", "xlsx", "Calc MS Excel 2007 XML")
    with pytest.raises(RuntimeError, match="不支持导出 .xlsb"):
        convert.run_excel_convert_34([source], "xlsb", tmp_path / "logs")


def test_word_conversion_com_fallback_and_cleanup(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("PYTOOLS_MODE", "office")
    monkeypatch.setattr(convert.sys, "platform", "win32")
    source = tmp_path / "report.docx"
    bad = tmp_path / "bad.doc"
    source.write_bytes(b"docx")
    bad.write_bytes(b"doc")
    events = []

    class FakeDocument:
        def __init__(self, path):
            self.path = Path(path)

        def SaveAs2(self, out, **_kwargs):
            events.append("SaveAs2")
            raise OSError("fallback required")

        def SaveAs(self, out, **_kwargs):
            events.append("SaveAs")
            if self.path.name == "bad.doc":
                raise OSError("cannot save")
            Path(out).write_bytes(b"converted")

        def Close(self, **_kwargs):
            events.append("close")

    class FakeDocuments:
        def Open(self, path, **_kwargs):
            return FakeDocument(path)

    app = SimpleNamespace(Documents=FakeDocuments(), Quit=lambda: events.append("quit"))
    monkeypatch.setattr(convert, "_start_word_app", lambda: (app, "Word.Application"))

    result = convert.run_word_convert_36([source, bad], "doc", tmp_path / "logs")

    assert result == {"ok": 1, "skip": 1, "target_ext": "doc", "engine": "Word.Application"}
    assert (tmp_path / "doc" / "report.doc").is_file()
    assert events.count("SaveAs2") == 2 and events.count("SaveAs") == 2
    assert events.count("close") == 2 and events[-1] == "quit"


def test_word_conversion_modes_and_libreoffice_error_path(tmp_path, monkeypatch) -> None:
    source = tmp_path / "report.docx"
    source.write_bytes(b"docx")
    monkeypatch.setenv("PYTOOLS_MODE", "python")
    assert convert.run_word_convert_36([tmp_path / "missing.txt"], "bad", tmp_path / "logs") == {
        "ok": 0, "skip": 1, "target_ext": "docx", "engine": ""
    }
    with pytest.raises(RuntimeError, match="非 Office 模式"):
        convert.run_word_convert_36([source], "docx", tmp_path / "logs")

    monkeypatch.setenv("PYTOOLS_MODE", "office")
    monkeypatch.setattr(convert.sys, "platform", "linux")
    monkeypatch.setattr(convert, "_libreoffice_executable", lambda: "soffice")
    monkeypatch.setattr(convert, "_libreoffice_convert_one", lambda *_args: (_ for _ in ()).throw(RuntimeError("no output")))
    result = convert.run_word_convert_36([source], "doc", tmp_path / "logs")
    assert result == {"ok": 0, "skip": 1, "target_ext": "doc", "engine": "LibreOffice"}


def test_word_app_discovery_and_libreoffice_success(tmp_path, monkeypatch) -> None:
    app = SimpleNamespace(Visible=True, DisplayAlerts=1)
    seen = []
    monkeypatch.setattr(convert, "_dispatch_first", lambda ids: (seen.append(ids) or (app, ids[0])))
    returned, engine = convert._start_word_app()
    assert returned is app and engine == "Word.Application"
    assert seen[0] == ["Word.Application", "kwps.Application", "KWPS.Application", "wps.Application"]
    assert app.Visible is False and app.DisplayAlerts == 0

    source = tmp_path / "report.docx"
    source.write_bytes(b"docx")
    monkeypatch.setenv("PYTOOLS_MODE", "office")
    monkeypatch.setattr(convert.sys, "platform", "linux")
    monkeypatch.setattr(convert, "_libreoffice_executable", lambda: "soffice")
    calls = []

    def convert_doc(executable, src, out, ext, output_filter):
        calls.append((executable, src.suffix, ext, output_filter))
        out.write_bytes(b"doc")

    monkeypatch.setattr(convert, "_libreoffice_convert_one", convert_doc)
    result = convert.run_word_convert_36([source], "doc", tmp_path / "logs")
    assert result == {"ok": 1, "skip": 0, "target_ext": "doc", "engine": "LibreOffice"}
    assert calls == [("soffice", ".docx", "doc", "MS Word 97")]
