from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook
from openpyxl.comments import Comment

from PyCode.core import compare_393, config_init, summary_321_322, wide_397_398
from PyCode.core.config_xlsx import GlobalConfig, PathMapRule, TimelineRule
from PyCode.core.date_parse import normalize_data_date_value, parse_data_date


def _rule(**changes) -> TimelineRule:
    fields = {
        "enabled": True,
        "name": "rule",
        "wb_keyword": "book",
        "sheet_keyword": "Data",
        "row_header_col": 1,
        "col_header_row": 1,
        "col_header_rows": [1],
        "data_row_start": 2,
        "data_col_start": 2,
    }
    fields.update(changes)
    return TimelineRule(**fields)


def test_config_init_small_repair_helpers() -> None:
    wb = Workbook()
    ws = wb.active
    assert config_init._ensure_sheet_headers(ws, ["键", "值"]) is True
    assert config_init._ensure_sheet_headers(ws, ["键", "值"]) is False

    assert config_init._ensure_header_comments(ws, config_init.SHEET_GLOBAL, ["键", "值"]) is True
    assert config_init._ensure_header_comments(ws, config_init.SHEET_GLOBAL, ["键", "值"]) is False

    rename = wb.create_sheet(config_init.SHEET_CONFIG_RENAME)
    rename["L1"].comment = Comment("remove", "test")
    rename["L2"].comment = Comment("remove", "test")
    assert config_init._ensure_rename_sheet_extra_cells(rename) is True
    assert rename["O2"].value == "Y" and rename["L1"].comment is None
    assert config_init._ensure_rename_sheet_extra_cells(rename) is False

    defaults = wb.create_sheet("defaults")
    defaults.append(["键", "值"])
    defaults.append(["输入目录", "custom"])
    defaults.append(["输出目录", ""])
    assert config_init._ensure_global_defaults(defaults) is True
    assert defaults["B2"].value == "custom" and defaults["B3"].value == "output"
    assert config_init._ensure_global_defaults(defaults) is False

    for helper in (
        config_init._ensure_archive_type_examples,
        config_init._ensure_submission_check_examples,
        config_init._ensure_file_sort_examples,
    ):
        example_ws = wb.create_sheet()
        assert helper(example_ws) is True
        assert helper(example_ws) is False
    wb.close()


def test_date_parse_invalid_and_fallback_cases() -> None:
    assert normalize_data_date_value(None) is None
    assert normalize_data_date_value("") is None
    assert normalize_data_date_value("account-201413") is None
    assert normalize_data_date_value("2026-02-30") is None
    assert normalize_data_date_value(datetime(2026, 2, 3, 12)) == date(2026, 2, 3)
    assert normalize_data_date_value(date(2026, 2, 3)) == date(2026, 2, 3)
    assert parse_data_date("file_202613.xlsx", "2026年4月") == ("2026-04-30", "A1")
    assert parse_data_date("no-date.xlsx", None) == ("", "")


def test_summary_helpers_cover_ranges_merges_comments_and_headers() -> None:
    assert summary_321_322._norm(None) == ""
    assert summary_321_322._supported(Path("a.XLSM"))
    assert not summary_321_322._supported(Path("a.csv"))
    assert summary_321_322._safe_sheet_name("A/B", set()) == "A_B"
    assert summary_321_322._safe_sheet_name("A/B", {"A_B"}) == "A_B_2"
    assert summary_321_322._disambiguate_headers(["x", "x", "x"]) == ["x", "x_2", "x_3"]

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "merged title"
    ws.merge_cells("A1:B1")
    ws["C2"] = "value"
    ws["C2"].comment = Comment(" note ", "test")
    assert summary_321_322._gcv(ws, 1, 2) == "merged title"
    assert summary_321_322._extract_comments(ws) == {"C2": "note"}
    assert summary_321_322._used_bounds(ws) == (1, 1, 2, 3)

    fake = SimpleNamespace(
        calculate_dimension=lambda: (_ for _ in ()).throw(ValueError("bad dimension")),
        max_row=4,
        max_column=3,
    )
    assert summary_321_322._used_bounds(fake) == (1, 1, 4, 3)
    assert summary_321_322._split_addr("$AB$12") == ("AB", 12)
    assert summary_321_322._split_addr("A") == ("A", 0)
    assert summary_321_322._col2num("AB") == 28
    assert summary_321_322._col2num("A1") == 0

    regions = summary_321_322._extract_regions({"A1": "行区域1起", "B3": "行区域#1终", "C1": "行区域2"}, "行区域")
    assert regions == [(1, 3, 1, 2)]
    names, addrs = summary_321_322._extract_set_info({"A1": "Set(机构)", "B2": "set（月份）"})
    assert names == ["机构", "月份"] and addrs == ["A1", "B2"]
    assert summary_321_322._extract_set_info({"A1": "no set"}) == ([], [])

    ws["A2"] = "group"
    ws["A3"] = "item"
    assert summary_321_322._build_col_headers(ws, [(1, 2, 1, 3)]) == ["merged title_group", "merged title", "value"]
    assert summary_321_322._build_col_headers(ws, [(2, 2, 1, 3)]) == ["group", "B", "value"]
    wb.close()


def test_summary_empty_and_invalid_workbooks(tmp_path: Path) -> None:
    result = summary_321_322.run_summary_by_usedrange(
        [tmp_path / "missing.xlsx", tmp_path / "unsupported.csv"], tmp_path / "out", tmp_path / "logs"
    )
    assert result == {"saved": False, "rows": 0, "path": ""}
    with pytest.raises(ValueError, match="模板仅支持"):
        summary_321_322.run_summary_by_comment(tmp_path / "missing.xlsx", [], tmp_path / "out", tmp_path / "logs")

    template = tmp_path / "empty-template.xlsx"
    wb = Workbook()
    wb.active["A1"] = "without annotations"
    wb.save(template)
    wb.close()
    with pytest.raises(ValueError, match="未识别到可用批注区域"):
        summary_321_322.run_summary_by_comment(template, [], tmp_path / "out", tmp_path / "logs")


def test_compare_helper_range_and_diff_edges() -> None:
    rule = _rule(data_row_end=5, data_col_end=5)
    target = pd.DataFrame([["row", "A", "B"], ["x", 1, 2]])
    source = pd.DataFrame([["row", "A", "C"], ["y", 1, 3], ["z", 2, 4]])
    assert compare_393._resolve_range(target, rule) == (0, [0], 1, 1, 1, 2)
    assert compare_393._build_col_path(target, 2, [0]) == "B"
    assert compare_393._build_col_path(target, 3, [0]) == ""
    assert compare_393._build_row_path(target, 1, 0) == "x"
    assert compare_393._build_row_path(target, 5, 0) == ""
    assert compare_393._collect_path_set(target, rule) == ({"x"}, {"A", "B"})
    assert compare_393._position_diff_rows(target, source, rule)
    assert compare_393._path_diff_rows(target, source, rule)
    assert compare_393._resolve_template_sheet(["Data"], "Other") == "Data"
    assert compare_393._resolve_template_sheet(["Data", "Other"], "Missing") is None


def test_wide_helpers_parse_keys_merges_and_compatibility() -> None:
    log = SimpleNamespace(warning=lambda *_args: None, info=lambda *_args: None)
    assert wide_397_398._row_no_from_addr("AB123") == "123"
    assert wide_397_398._row_no_from_addr("AB") == ""
    assert wide_397_398._row_no_from_addr("") == ""
    assert wide_397_398._disambiguate_headers(["x", "x", "x"]) == ["x", "x_2", "x_3"]

    rule = _rule(row_header_col_specified=False, data_row_start=0, data_col_start=0)
    assert not wide_397_398._is_wide_rule_compatible(rule, log)
    rule.data_row_start, rule.data_col_start = 2, 2
    rule.required_row_headers = ["row"]
    assert wide_397_398._is_wide_rule_compatible(rule, log)

    explicit = _rule(target_dedup_cols=["日期", "C", "日期"])
    assert wide_397_398._resolve_target_dedup_idx(explicit, ["日期", "name", "value"], [], log) == [0, 2]
    assert wide_397_398._resolve_target_dedup_idx(
        _rule(target_dedup_cols=["missing"]), ["日期"], [], log
    ) is None
    assert wide_397_398._resolve_target_dedup_idx(
        _rule(), ["工作簿名", "工作表名", "数据日期", "机构", "行头路径"], ["机构"], log
    ) == [2, 3, 4]
    assert wide_397_398._resolve_target_dedup_idx(
        _rule(), ["工作簿名", "工作表名", "数据日期", "value"], [], log
    ) == [2, 3]
    assert wide_397_398._resolve_target_dedup_idx(_rule(), ["工作簿名", "工作表名"], [], log) is None

    frame = pd.DataFrame([["anchor"], [None]])
    assert wide_397_398._merged_value(frame, [(0, 1, 0, 0)], 1, 0) == "anchor"
    assert wide_397_398._merged_value(frame, [], 1, 0) is None
    assert wide_397_398._merged_value(frame, [], -1, 0) is None


def test_empty_source_timeline_and_wide_do_not_create_outputs(tmp_path: Path) -> None:
    g = GlobalConfig(tmp_path, tmp_path / "out", tmp_path / "logs")
    pmap = PathMapRule(True, "map", "*", "", "", "两者", "精确", "a", "b")
    rule = _rule()
    assert __import__("PyCode.core.timeline_396", fromlist=["run_timeline_slim"]).run_timeline_slim(
        [rule], [pmap], g, source_paths=[]
    ) is None
    assert wide_397_398.run_wide_summary([rule], [pmap], g, True, source_paths=[]) is None

