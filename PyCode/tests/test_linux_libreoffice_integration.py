from __future__ import annotations

import shutil
import sys
import zipfile
from pathlib import Path
from shutil import copy2

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from PyCode.core import convert_33_34_35_37 as convert


_LO_AVAILABLE = bool(shutil.which("soffice") or shutil.which("libreoffice") or list(Path("/opt").glob("libreoffice*/program/soffice")))
pytestmark = pytest.mark.skipif(
    sys.platform != "linux" or not _LO_AVAILABLE,
    reason="Linux LibreOffice integration test",
)


def _excel_book(path: Path) -> Path:
    wb = Workbook()
    first = wb.active
    first.title = "Data"
    first.append(["name", "value"])
    first.append(["alpha", 7])
    second = wb.create_sheet("Second")
    second.append(["code", "enabled"])
    second.append(["B", True])
    wb.save(path)
    wb.close()
    return path


def _minimal_docx(path: Path) -> Path:
    content_types = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>'''
    relationships = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>'''
    document = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
<w:p><w:r><w:t>LibreOffice integration content</w:t></w:r></w:p><w:sectPr/>
</w:body></w:document>'''
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document)
    return path


def test_libreoffice_excel_conversion_and_python_mode_xls_import(tmp_path: Path, monkeypatch) -> None:
    source = _excel_book(tmp_path / "source.xlsx")
    monkeypatch.setenv("PYTOOLS_MODE", "office")

    xls_stat = convert.run_excel_convert_34([source], "xls", tmp_path / "office-xls")
    xls_path = tmp_path / "xls" / "source.xls"
    assert xls_stat == {"ok": 1, "skip": 0, "target_ext": "xls", "engine": "LibreOffice"}
    assert pd.read_excel(xls_path, sheet_name=0, engine="xlrd").iloc[0].tolist() == ["alpha", 7]

    xlsx_stat = convert.run_excel_convert_34([xls_path], "xlsx", tmp_path / "office-xlsx")
    office_xlsx = tmp_path / "xls" / "xlsx" / "source.xlsx"
    assert xlsx_stat["ok"] == 1 and xlsx_stat["engine"] == "LibreOffice"
    wb = load_workbook(office_xlsx, data_only=True)
    try:
        assert wb["Data"]["A2"].value == "alpha"
        assert wb["Second"]["A2"].value == "B"
    finally:
        wb.close()

    monkeypatch.setenv("PYTOOLS_MODE", "python")
    python_source = tmp_path / "python-input" / "source.xls"
    python_source.parent.mkdir()
    copy2(xls_path, python_source)
    python_stat = convert.run_excel_convert_34([python_source], "xlsx", tmp_path / "python-xlsx")
    python_xlsx = tmp_path / "python-input" / "xlsx" / "source.xlsx"
    assert python_stat == {"ok": 1, "skip": 0, "target_ext": "xlsx", "engine": "Python"}
    wb = load_workbook(python_xlsx, data_only=True)
    try:
        assert wb.sheetnames == ["Data", "Second"]
        assert wb["Data"]["A2"].value == "alpha"
        assert wb["Second"]["A2"].value == "B"
    finally:
        wb.close()

    monkeypatch.setenv("PYTOOLS_MODE", "office")
    csv_stat = convert.run_excel_convert_34([source], "csv", tmp_path / "office-csv")
    csv_path = tmp_path / "csv" / "source.csv"
    assert csv_stat["ok"] == 1 and csv_stat["engine"] == "LibreOffice"
    assert pd.read_csv(csv_path, encoding="utf-8-sig").iloc[0].tolist() == ["alpha", 7]


def test_libreoffice_word_docx_doc_roundtrip(tmp_path: Path, monkeypatch) -> None:
    source = _minimal_docx(tmp_path / "report.docx")
    monkeypatch.setenv("PYTOOLS_MODE", "office")

    doc_stat = convert.run_word_convert_36([source], "doc", tmp_path / "office-doc")
    doc_path = tmp_path / "doc" / "report.doc"
    assert doc_stat == {"ok": 1, "skip": 0, "target_ext": "doc", "engine": "LibreOffice"}
    assert doc_path.is_file() and doc_path.stat().st_size > 0

    docx_stat = convert.run_word_convert_36([doc_path], "docx", tmp_path / "office-docx")
    docx_path = tmp_path / "doc" / "docx" / "report.docx"
    assert docx_stat["ok"] == 1 and docx_stat["engine"] == "LibreOffice"
    with zipfile.ZipFile(docx_path) as archive:
        assert "word/document.xml" in archive.namelist()
        assert b"LibreOffice integration content" in archive.read("word/document.xml")
