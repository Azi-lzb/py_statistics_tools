"""Windows Office-mode integration tests; each COM result is checked against Python output."""
from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from PyCode.core.feature_1_1_split_village_bank import run_split_village_bank
from PyCode.core.feature_1_1_split_village_bank_com import run_split_village_bank_com
from PyCode.core.feature_1_2_normalize_institution import run_normalize_institution
from PyCode.core.feature_1_2_normalize_institution_com import run_normalize_institution_com
from PyCode.core.feature_1_3_remove_foreign_bank import run_remove_foreign_bank
from PyCode.core.feature_1_3_remove_foreign_bank_com import run_remove_foreign_bank_com
from PyCode.core.feature_1_4_fx_header import run_fx_header_fix
from PyCode.core.feature_1_4_fx_header_com import run_fx_header_fix_com
from PyCode.core.feature_1_5_region_sum_check import run_region_sum_check
from PyCode.core.feature_1_5_region_sum_check_com import run_region_sum_check_com
from PyCode.core.feature_1_8_adjust_rural_loan import run_adjust_rural_loan
from PyCode.core.feature_1_8_adjust_rural_loan_com import run_adjust_rural_loan_com
from PyCode.core.feature_1_9_adjust_rural_loan_quarterly import run_adjust_rural_loan_quarterly
from PyCode.core.feature_1_9_adjust_rural_loan_quarterly_com import run_adjust_rural_loan_quarterly_com
from PyCode.core.convert_33_34_35_37 import run_batch_rename_sheet_37, run_batch_rename_sheet_37_com
from PyCode.tests.test_conversion_functional import _make_rename_config


pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="Office COM is Windows-only")


def _save(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def _workbook_values(path: Path) -> dict[str, list[tuple]]:
    wb = load_workbook(path, data_only=True)
    try:
        return {ws.title: list(ws.values) for ws in wb.worksheets}
    finally:
        wb.close()


def _split_sources(root: Path) -> tuple[Path, Path]:
    current, previous = root / "current.xlsx", root / "previous.xlsx"
    for path, values in ((current, (100, 80)), (previous, (80, 64))):
        wb = Workbook()
        ws = wb.active
        ws.title = "村镇银行本外币"
        ws["C5"], ws["D5"], ws["E5"] = values[0], "存款D", "存款E"
        ws["CI5"], ws["CJ5"], ws["CK5"] = values[1], "贷款J", "贷款K"
        _save(wb, path)
    return current, previous


def _split_target(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "本外币分机构"
    ws.append(["机构"])
    ws.append(["村镇银行"])
    return _save(wb, path)


def test_com_village_bank_split_matches_python(tmp_path: Path) -> None:
    current, previous = _split_sources(tmp_path)
    py_target = _split_target(tmp_path / "python" / "target.xlsx")
    com_target = _split_target(tmp_path / "com" / "target.xlsx")

    py_stat = run_split_village_bank(current, previous, [py_target], tmp_path / "logs-python")
    com_stat = run_split_village_bank_com(current, previous, [com_target], tmp_path / "logs-com")

    assert com_stat["filled_sheets"] == py_stat["filled_sheets"] == 1
    assert com_stat["failed_files"] == py_stat["failed_files"] == 0
    assert _workbook_values(com_target) == _workbook_values(py_target)


def test_com_normalize_and_remove_foreign_match_python(tmp_path: Path) -> None:
    mapping = {"简称甲": "标准甲", "机构乙": "标准乙"}
    py_file, com_file = tmp_path / "python" / "institutions.xlsx", tmp_path / "com" / "institutions.xlsx"
    for path in (py_file, com_file):
        wb = Workbook()
        ws = wb.active
        ws.title = "惠州分机构"
        ws.append(["机构名称", "备注"])
        ws.append(["简称甲", "x"])
        ws_b = wb.create_sheet("另一分机构")
        ws_b.append([None, "机构名称"])
        ws_b.append([None, "机构乙（分行）"])
        _save(wb, path)

    py_stat = run_normalize_institution([py_file], mapping, tmp_path / "logs-python")
    com_stat = run_normalize_institution_com([com_file], mapping, tmp_path / "logs-com")
    assert com_stat["mapped_rows"] == py_stat["mapped_rows"] == 2
    assert _workbook_values(com_file) == _workbook_values(py_file)

    py_foreign = tmp_path / "python" / "foreign.xlsx"
    com_foreign = tmp_path / "com" / "foreign.xlsx"
    for path in (py_foreign, com_foreign):
        wb = Workbook()
        ws = wb.active
        ws.title = "本外币分机构"
        for row in (("外资甲", 1), ("本地行", 2), ("外资乙", 3), ("外资乙", 4)):
            ws.append(row)
        _save(wb, path)

    py_stat = run_remove_foreign_bank([py_foreign], {"外资甲", "外资乙"}, tmp_path / "logs-python")
    com_stat = run_remove_foreign_bank_com([com_foreign], {"外资甲", "外资乙"}, tmp_path / "logs-com")
    assert com_stat["deleted_rows"] == py_stat["deleted_rows"] == 3
    py_result = py_foreign.with_name("foreign(金融局).xlsx")
    com_result = com_foreign.with_name("foreign(金融局).xlsx")
    assert _workbook_values(com_result) == _workbook_values(py_result)


def _fx_source(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "外汇报表"
    ws.oddHeader.left.text = "单位：万元"
    return _save(wb, path)


def _region_source(path: Path, total: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "资产分地区"
    ws.append(["序号", "地区", "指标甲"])
    ws.append([None, "地区", "指标甲"])
    ws.append([None, "惠州市", total])
    for name, value in (("惠阳区", 40), ("博罗县", 30), ("惠东县", 10),
                        ("龙门县", 5), ("惠州市本部", 10), ("大亚湾", 5)):
        ws.append([None, name, value])
    return _save(wb, path)


def test_com_fx_header_and_region_sum_match_python(tmp_path: Path) -> None:
    py_fx = _fx_source(tmp_path / "python" / "fx.xlsx")
    com_fx = _fx_source(tmp_path / "com" / "fx.xlsx")
    py_fx_stat = run_fx_header_fix([py_fx], tmp_path / "logs-python")
    com_fx_stat = run_fx_header_fix_com([com_fx], tmp_path / "logs-com")
    assert com_fx_stat["modified_sheets"] == py_fx_stat["modified_sheets"] == 1
    for path in (py_fx, com_fx):
        wb = load_workbook(path)
        try:
            assert "万美元" in wb["外汇报表"].oddHeader.left.text
        finally:
            wb.close()

    py_source = _region_source(tmp_path / "python" / "region.xlsx", 101)
    com_source = _region_source(tmp_path / "com" / "region.xlsx", 101)
    py_stat = run_region_sum_check([py_source], tmp_path / "python" / "report", tmp_path / "logs-python")
    com_stat = run_region_sum_check_com([com_source], tmp_path / "com" / "report", tmp_path / "logs-com")
    assert com_stat["errors"] == py_stat["errors"] == 1
    py_wb = load_workbook(py_stat["path"], data_only=True)
    com_wb = load_workbook(com_stat["path"], data_only=True)
    try:
        assert [py_wb["核对结果"].cell(2, col).value for col in (4, 5, 6)] == [
            com_wb["核对结果"].cell(2, col).value for col in (4, 5, 6)
        ]
    finally:
        py_wb.close()
        com_wb.close()


def _rural_workbook(path: Path, amount: int, area_value: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "惠州市涉农贷款分机构表"
    ws.append(["序号", "机构名称", "涉农贷款", "涉农贷款比上月"])
    ws.append([])
    ws.append([1, "机构甲", amount, None])
    area = wb.create_sheet("本外币涉农贷款分地区")
    area.append(["地区", "涉农贷款", "涉农贷款比上月"])
    area.append(["惠州市", 1000, area_value])
    return _save(wb, path)


def _quarterly_workbook(
    path: Path, city_amount: int, district_amount: int, area_values=None,
    city_second_amount: int = 30,
) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in (
        ("惠州市-涉农贷款分机构表", [("机构甲", city_amount), ("机构乙", city_second_amount)]),
        ("惠阳区涉农贷款分机构表", [("机构丙", district_amount)]),
    ):
        ws = wb.create_sheet(title)
        ws.append(["序号", "机构名称", "涉农贷款", "涉农贷款比上月"])
        ws.append([])
        for idx, (name, amount) in enumerate(rows, start=1):
            ws.append([idx, name, amount, None])
    if area_values is not None:
        area = wb.create_sheet("涉农贷款分地区统计表")
        area.append(["地区", "涉农贷款", "涉农贷款比上月"])
        for region, value in area_values.items():
            area.append([region, None, value])
    return _save(wb, path)


def test_com_rural_loan_and_quarterly_adjustments_match_python(tmp_path: Path) -> None:
    py_cur = _rural_workbook(tmp_path / "python" / "rural-current.xlsx", 110, 10)
    py_prev = _rural_workbook(tmp_path / "python" / "rural-previous.xlsx", 100, 0)
    com_cur = _rural_workbook(tmp_path / "com" / "rural-current.xlsx", 110, 10)
    com_prev = _rural_workbook(tmp_path / "com" / "rural-previous.xlsx", 100, 0)
    py_result = run_adjust_rural_loan(py_cur, py_prev, tmp_path / "logs-python")
    com_result = run_adjust_rural_loan_com(com_cur, com_prev, tmp_path / "logs-com")
    assert (com_result["matched"], com_result["fill_total"], com_result["consistent"], com_result["error"]) == (
        py_result["matched"], py_result["fill_total"], py_result["consistent"], py_result["error"]
    )
    assert _workbook_values(com_cur) == _workbook_values(py_cur)

    py_cur_q = _quarterly_workbook(
        tmp_path / "python" / "quarterly-current.xlsx", 100, 80,
        {"惠州市": 7, "惠阳区": 2}, city_second_amount=30,
    )
    py_prev_q = _quarterly_workbook(
        tmp_path / "python" / "quarterly-previous.xlsx", 95, 78, city_second_amount=28,
    )
    com_cur_q = _quarterly_workbook(
        tmp_path / "com" / "quarterly-current.xlsx", 100, 80,
        {"惠州市": 7, "惠阳区": 2}, city_second_amount=30,
    )
    com_prev_q = _quarterly_workbook(
        tmp_path / "com" / "quarterly-previous.xlsx", 95, 78, city_second_amount=28,
    )
    py_result = run_adjust_rural_loan_quarterly(py_cur_q, py_prev_q, tmp_path / "logs-python")
    com_result = run_adjust_rural_loan_quarterly_com(com_cur_q, com_prev_q, tmp_path / "logs-com")
    assert com_result["consistent"] == py_result["consistent"] is True
    assert com_result["matched"] == py_result["matched"] == 3
    assert _workbook_values(com_cur_q) == _workbook_values(py_cur_q)


def test_com_sheet_rename_matches_python(tmp_path: Path) -> None:
    config = _make_rename_config(tmp_path / "config.xlsx")
    py_path = tmp_path / "python" / "book.xlsx"
    com_path = tmp_path / "com" / "book.xlsx"
    for path in (py_path, com_path):
        wb = Workbook()
        wb.active.title = "旧表"
        wb.active["A1"] = "preserved"
        wb.create_sheet("保持")
        _save(wb, path)

    py_stat = run_batch_rename_sheet_37(config, [py_path], tmp_path / "logs-python")
    com_stat = run_batch_rename_sheet_37_com(config, [com_path], tmp_path / "logs-com")

    assert py_stat == {"workbooks": 1, "rename_ok": 1, "skip": 0}
    assert com_stat["workbooks"] == 1 and com_stat["rename_ok"] == 1 and com_stat["skip"] == 0
    assert com_stat["engine"]
    assert _workbook_values(py_path) == _workbook_values(com_path)
    assert list(_workbook_values(com_path)) == ["新表", "保持"]
