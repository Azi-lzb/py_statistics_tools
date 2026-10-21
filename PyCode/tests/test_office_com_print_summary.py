from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from PyCode.core.print_310 import (
    run_print_config_all_modes,
    run_print_fast_by_comment,
    run_print_keep_by_comment,
)
from PyCode.core.print_310_com import (
    run_print_config_all_modes_com,
    run_print_config_to_pdf_com,
    run_print_fast_by_comment_com,
    run_print_keep_by_comment_com,
)
from PyCode.core.summary_321_322 import run_summary_by_comment, run_summary_by_usedrange
from PyCode.core.summary_321_322_com import (
    run_summary_by_comment_com,
    run_summary_by_usedrange_com,
)


pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="Office COM is Windows-only")


def _save(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def _values(path: Path) -> dict[str, list[tuple]]:
    wb = load_workbook(path, data_only=True)
    try:
        return {ws.title: list(ws.values) for ws in wb.worksheets}
    finally:
        wb.close()


def _print_source(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "报表"
    ws.append([0, "标题"])
    ws.append(["甲行", 5])
    ws["A1"].comment = Comment("打印区域 1；打印 保留源格式", "test")
    ws["B2"].comment = Comment("打印区域#1", "test")
    ws["B1"].comment = Comment("内容批注", "test")
    return _save(wb, path)


def test_com_comment_print_and_config_print_match_python(tmp_path: Path) -> None:
    source = _print_source(tmp_path / "source.xlsx")
    python_output, com_output = tmp_path / "python", tmp_path / "com"

    py_keep = run_print_keep_by_comment([source], python_output, tmp_path / "logs-python")
    com_keep = run_print_keep_by_comment_com([source], com_output, tmp_path / "logs-com")
    assert com_keep["sheets_hit"] == py_keep["sheets_hit"] == 1
    py_file = next(python_output.glob("*_按批注打印（保留源格式）*.xlsx"))
    com_file = next(com_output.glob("*_按批注打印（保留源格式）_COM.xlsx"))
    assert _values(com_file) == _values(py_file)

    py_fast = run_print_fast_by_comment([source], python_output, tmp_path / "logs-python", hide_zero=True)
    com_fast = run_print_fast_by_comment_com([source], com_output, tmp_path / "logs-com", hide_zero=True)
    assert com_fast["sheets_hit"] == py_fast["sheets_hit"] == 1
    py_fast_file = next(python_output.glob("*_按批注打印（快速复制）*.xlsx"))
    com_fast_file = next(com_output.glob("*_按批注打印（快速复制）_COM.xlsx"))
    assert _values(com_fast_file) == _values(py_fast_file)
    wb = load_workbook(com_fast_file, data_only=True)
    assert wb["报表"]["A1"].value is None
    wb.close()

    py_tasks, com_tasks = [], []
    for mode in (1, 2, 3):
        common = {
            "row_no": mode + 1, "enabled": True, "mode": mode, "source_wb": source,
            "source_ws": "报表", "source_ranges": "A1:B2", "target_ws": f"输出{mode}",
            "fit_wide": 1, "fit_tall": 1, "orientation": "横向", "center_h": True,
            "center_v": False, "drop_comments": True, "hide_zero": True,
        }
        py_tasks.append(dict(common, target_wb=tmp_path / "python" / f"mode{mode}.xlsx"))
        com_tasks.append(dict(common, target_wb=tmp_path / "com" / f"mode{mode}.xlsx"))

    py_stat = run_print_config_all_modes(py_tasks, tmp_path / "logs-python")
    com_stat = run_print_config_all_modes_com(com_tasks, tmp_path / "logs-com")
    assert com_stat["task_ok"] == py_stat["task_ok"] == 3
    assert com_stat["written_rows"] == py_stat["written_rows"] == 6
    for mode in (1, 2, 3):
        assert _values(tmp_path / "com" / f"mode{mode}.xlsx") == _values(
            tmp_path / "python" / f"mode{mode}.xlsx"
        )

    pdf = run_print_config_to_pdf_com([com_tasks[0]], tmp_path / "pdf", tmp_path / "logs-com")
    assert pdf["task_ok"] == 1 and len(pdf["pdf_files"]) == 1
    assert Path(pdf["pdf_files"][0]).read_bytes().startswith(b"%PDF-")


def _summary_source(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "资产"
    ws.append(["机构", "数值", "数值"])
    ws.append(["甲行", 10, 11])
    return _save(wb, path)


def _comment_fixture(root: Path) -> tuple[Path, Path]:
    source = root / "comment-source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "资产"
    ws["A1"], ws["D1"] = "标题", "广州"
    ws["A3"], ws["B3"], ws["C3"] = "甲行", 10, 11
    _save(wb, source)

    template = root / "template.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "资产"
    ws["A3"].comment = Comment("行区域1", "test")
    ws["B3"].comment = Comment("行区域#1", "test")
    ws["B1"] = "机构"
    ws["B1"].comment = Comment("列区域1", "test")
    ws["C1"] = "金额"
    ws["C1"].comment = Comment("列区域#1", "test")
    ws["D1"] = "地区"
    ws["D1"].comment = Comment("set(地区)", "test")
    _save(wb, template)
    return source, template


def test_com_summary_modes_match_python(tmp_path: Path) -> None:
    source = _summary_source(tmp_path / "used" / "source.xlsx")
    py_used = run_summary_by_usedrange([source], tmp_path / "python" / "used", tmp_path / "logs-python")
    com_used = run_summary_by_usedrange_com([source], tmp_path / "com" / "used", tmp_path / "logs-com")
    assert com_used["saved"] is py_used["saved"] is True
    assert com_used["rows"] == py_used["rows"] == 2
    assert _values(com_used["path"]) == _values(py_used["path"])

    comment_source, template = _comment_fixture(tmp_path / "comment")
    py_comment = run_summary_by_comment(
        template, [comment_source], tmp_path / "python" / "comment", tmp_path / "logs-python"
    )
    com_comment = run_summary_by_comment_com(
        template, [comment_source], tmp_path / "com" / "comment", tmp_path / "logs-com"
    )
    assert com_comment["saved"] is py_comment["saved"] is True
    assert com_comment["rows"] == py_comment["rows"] == 1
    assert _values(com_comment["path"]) == _values(py_comment["path"])
