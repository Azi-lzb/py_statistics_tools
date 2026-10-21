from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook

from PyCode.core import precheck_399
from PyCode.core.config_xlsx import (
    PATH_MAP_COLS,
    REQUIRED_GLOBAL_KEYS,
    SHEET_GLOBAL,
    SHEET_PATH_MAP,
    SHEET_TIMELINE_RULE,
    TIMELINE_COLS,
)


def _config(path: Path, *, timeline_rows=(), path_map_rows=()) -> None:
    wb = Workbook()
    global_ws = wb.active
    global_ws.title = SHEET_GLOBAL
    global_ws.append(["键", "值"])
    for key, value in zip(REQUIRED_GLOBAL_KEYS, ("input", "output", "logs")):
        global_ws.append([key, value])
    timeline = wb.create_sheet(SHEET_TIMELINE_RULE)
    timeline.append(TIMELINE_COLS)
    for row in timeline_rows:
        timeline.append([row.get(col, "") for col in TIMELINE_COLS])
    path_map = wb.create_sheet(SHEET_PATH_MAP)
    path_map.append(PATH_MAP_COLS)
    for row in path_map_rows:
        path_map.append([row.get(col, "") for col in PATH_MAP_COLS])
    wb.save(path)
    wb.close()


def test_precheck_early_errors_for_missing_or_unreadable_config(tmp_path, capsys):
    assert precheck_399.run_precheck(tmp_path / "missing.xlsx") == 1
    assert "配置文件不存在" in capsys.readouterr().out

    broken = tmp_path / "broken.xlsx"
    broken.write_text("not an excel file", encoding="utf-8")
    assert precheck_399.run_precheck(broken) == 1
    assert "无法打开配置" in capsys.readouterr().out


def test_precheck_reports_missing_sheets_columns_and_global_keys(tmp_path, capsys):
    cfg = tmp_path / "incomplete.xlsx"
    wb = Workbook()
    global_ws = wb.active
    global_ws.title = SHEET_GLOBAL
    global_ws.append(["键"])
    timeline = wb.create_sheet(SHEET_TIMELINE_RULE)
    timeline.append(["规则名称"])
    wb.save(cfg)
    wb.close()

    errors = precheck_399.run_precheck(cfg)
    report = capsys.readouterr().out
    assert errors > 10
    assert "至少需要 键/值 两列" in report
    assert "缺少 Sheet" in report
    assert "缺少列" in report


def test_precheck_reports_schema_directory_and_rule_errors(tmp_path, monkeypatch, capsys):
    cfg = tmp_path / "config.xlsx"
    timeline_rows = [
        {
            "是否启用": "Y", "规则名称": "", "工作簿关键字": "", "工作表关键字": "",
            "行头列": "bad-column", "列表头行": "", "数据起始行": "bad-row", "数据结束行": "bad-end",
            "数据起始列": "B", "数据结束列": "A", "启用目标写入": "Y", "set区域": "无地址",
        },
        {
            "是否启用": "Y", "规则名称": "重复", "工作簿关键字": "book", "工作表关键字": "sheet",
            "行头列": "", "列表头行": " ,；，;; ", "必含行头": "将被忽略", "数据起始行": "",
            "数据起始列": "", "启用目标写入": "N",
        },
        {
            "是否启用": "Y", "规则名称": "重复", "工作簿关键字": "book", "工作表关键字": "sheet",
            "行头列": "A", "列表头行": "0", "数据起始行": "0", "数据结束行": "1",
            "数据起始列": "bad", "数据结束列": "bad", "启用目标写入": "N",
        },
        {
            "是否启用": "Y", "规则名称": "range", "工作簿关键字": "book", "工作表关键字": "sheet",
            "行头列": "A", "列表头行": "1", "数据起始行": "10", "数据结束行": "2",
            "数据起始列": "C", "数据结束列": "A", "启用目标写入": "N",
        },
        {"是否启用": "N", "规则名称": "disabled"},
    ]
    path_rows = [
        {
            "是否启用": "Y", "映射名称": "", "适用规则名": "missing-rule", "作用对象": "列",
            "匹配方式": "fuzzy", "原始路径": "", "标准路径": "",
        },
        {"是否启用": "N", "映射名称": "ignored"},
    ]
    _config(cfg, timeline_rows=timeline_rows, path_map_rows=path_rows)
    missing_input = tmp_path / "missing-input"
    output = tmp_path / "out"
    logs = tmp_path / "logs"
    monkeypatch.setattr(precheck_399, "load_global", lambda _path: SimpleNamespace(
        input_dir=missing_input, output_dir=output, log_dir=logs
    ))

    errors = precheck_399.run_precheck(cfg)
    report = capsys.readouterr().out
    assert errors >= 15
    assert "输入目录不存在" in report
    assert "行头列非法" in report
    assert "数据结束行非法" in report
    assert "目标写入" in report
    assert "路径标准化映射" in report
    assert "启用规则名称重复" in report
    assert output.is_dir() and logs.is_dir()


def test_precheck_catches_global_config_errors_and_continues(tmp_path, monkeypatch, capsys):
    cfg = tmp_path / "config.xlsx"
    _config(cfg)
    monkeypatch.setattr(precheck_399, "load_global", lambda _path: (_ for _ in ()).throw(ValueError("bad global")))

    errors = precheck_399.run_precheck(cfg)
    report = capsys.readouterr().out
    assert errors == 1
    assert "解析全局配置失败: bad global" in report
    assert "时序提取规则: 启用 0 / 共 0 条" in report


def test_precheck_valid_config_with_warnings_returns_zero(tmp_path, monkeypatch, capsys):
    cfg = tmp_path / "config.xlsx"
    _config(
        cfg,
        timeline_rows=[{
            "是否启用": "Y", "规则名称": "valid", "工作簿关键字": "book", "工作表关键字": "sheet",
            "行头列": "A", "列表头行": "1", "启用目标写入": "Y",
        }],
        path_map_rows=[{
            "是否启用": "Y", "映射名称": "valid", "适用规则名": "valid", "作用对象": "两者",
            "匹配方式": "包含", "原始路径": "old", "标准路径": "new",
        }],
    )
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    output = tmp_path / "output"
    logs = tmp_path / "logs"
    monkeypatch.setattr(precheck_399, "load_global", lambda _path: SimpleNamespace(
        input_dir=input_dir, output_dir=output, log_dir=logs
    ))

    assert precheck_399.run_precheck(cfg) == 0
    report = capsys.readouterr().out
    assert "存在=True" in report
    assert "目标工作簿路径/目标工作表未完整填写" in report
    assert output.is_dir() and logs.is_dir()
