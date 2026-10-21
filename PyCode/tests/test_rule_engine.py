from __future__ import annotations

import builtins
import sys
import types
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook

from PyCode.core import rule_engine as engine
from PyCode.core.config_xlsx import PathMapRule, TimelineRule


def _rule(**changes) -> TimelineRule:
    values = {
        "enabled": True,
        "name": "rule",
        "wb_keyword": "",
        "sheet_keyword": "",
        "row_header_col": 1,
        "col_header_row": 2,
        "row_header_cols": [1],
        "row_header_col_specified": True,
        "col_header_rows": [2],
        "required_col_headers": [],
        "required_row_headers": [],
        "data_row_start": 0,
        "data_row_end": None,
        "data_col_start": 0,
        "data_col_end": None,
        "skip_keywords": [],
    }
    values.update(changes)
    return TimelineRule(**values)


def _map(**changes) -> PathMapRule:
    values = {
        "enabled": True,
        "name": "map",
        "applicable_rule": "",
        "wb_keyword": "",
        "sheet_keyword": "",
        "target": "两者",
        "match_mode": "精确",
        "original": "old",
        "standard": "new",
    }
    values.update(changes)
    return PathMapRule(**values)


def test_column_keyword_and_header_helpers():
    assert [engine._col_letter(n) for n in (0, 1, 26, 27, 52, 703)] == ["", "A", "Z", "AA", "AZ", "AAA"]
    assert engine._cell_addr(12, 28) == "AB12"
    assert engine._fallback_col_path(26) == "列_AA"
    assert engine.match_all_keywords("Monthly Report 2026", "monthly；report,2026")
    assert engine.match_all_keywords("anything", "；，，")
    assert engine.match_all_keywords("anything", "")
    assert not engine.match_all_keywords("", "required")
    assert not engine.match_all_keywords("Monthly Report", "monthly;balance")
    rules = [_rule(wb_keyword="2026"), _rule(wb_keyword="Report"), _rule(wb_keyword="missing")]
    assert len(engine.pick_rules_for_workbook(rules, "2026 Report", "all_match")) == 2
    assert engine.pick_rules_for_workbook(rules, "2026 Report", "first_match") == rules[:1]
    assert engine.pick_rules_for_workbook(rules, "none", "unknown") == []

    assert engine.parse_col_spec(None) == 0
    assert engine.parse_col_spec(float("nan")) == 0
    assert engine.parse_col_spec(" ") == 0
    assert engine.parse_col_spec("2.0") == 2
    assert engine.parse_col_spec("AB") == 28
    assert engine.parse_col_spec("A1") == 0
    assert engine.parse_col_specs(None) == []
    assert engine.parse_col_specs(float("nan")) == []
    assert engine.parse_col_specs("") == []
    assert engine.parse_col_specs("A,C;3，AA,A") == [1, 3, 27]
    assert engine.parse_col_specs("A,,B") == [1, 2]
    assert engine.parse_col_specs("A1,0,-1") == []
    assert engine._skip_match([], "x") is False
    assert engine._skip_match(["draft"], "draft sheet") is True
    assert engine._skip_match(["draft"], "final") is False
    assert engine._required_match([], []) is True
    assert engine._required_match(["bank", "total"], ["bank name", None, "total amount"])
    assert not engine._required_match(["bank", "missing"], ["bank name"])
    assert engine._val_str(None) == ""
    assert engine._val_str(float("nan")) == ""
    assert engine._val_str("  x  ") == "x"


def test_merge_header_and_noise_helpers():
    df = pd.DataFrame([["Merged", None, None], ["first", None, "third"]])
    merged = ((0, 0, 0, 1),)
    assert engine._resolve_merged_top_left(df, 0, 1, merged) == "Merged"
    assert engine._resolve_merged_top_left(df, 1, 1, merged) == ""
    assert engine._header_cell_text(df, -1, 0, ()) == ""
    assert engine._header_cell_text(df, 0, 1, merged) == "Merged"
    assert engine._header_cell_text(df, 1, 1, ()) == ""
    assert engine._header_cell_text_with_left_fill(df, 1, 1, ()) == "first"
    assert engine._header_cell_text_with_left_fill(pd.DataFrame([[None, None, None]]), 0, 2, ()) == ""
    assert engine._header_cell_text_with_left_fill(df, 1, 2, ()) == "third"
    assert engine._header_cell_text_with_left_fill(df, 0, 1, merged) == "Merged"
    assert engine._row_header_text(df, 1, 2, ()) == "third"
    assert engine._row_header_text(df, 9, 0, ()) == ""
    assert engine._row_header_text(df, 0, 1, merged) == "Merged"

    assert engine._normalize_header_noise("deposit1", "deposit") == "deposit"
    assert engine._normalize_header_noise("deposit", "deposit") == "deposit"
    assert engine._normalize_header_noise("deposit11", "other") == "deposit11"
    assert engine._normalize_header_noise("", "prev") == ""
    assert engine._strip_tail_digits("deposit21") == "deposit"
    assert engine._strip_tail_digits("123") == "123"
    assert engine._strip_tail_digits("deposit") == "deposit"
    assert engine._strip_tail_digits("") == ""
    assert engine._normalize_header_row_noise([]) == []
    assert engine._normalize_header_row_noise(["deposit", "deposit1", "deposit11", "other", "other2"]) == [
        "deposit", "deposit", "deposit", "other", "other"
    ]
    assert engine._normalize_header_row_noise(["", "deposit", "deposit2", "deposit21"]) == [
        "", "deposit", "deposit", "deposit"
    ]
    assert engine._normalize_header_row_noise(["first", "other"]) == ["first", "other"]


def test_column_path_builder_disambiguates_fills_maps_and_overrides():
    df = pd.DataFrame([
        ["date", "Group", None, "Other"],
        ["", "Amount", "Amount", "Count"],
        ["row", 1, 2, 3],
    ])
    rule = _rule(col_header_rows=[1, 2], data_col_start=2)
    maps = [
        _map(applicable_rule="other", target="列头"),
        _map(target="行头"),
        _map(target="列头", wb_keyword="another"),
        _map(target="列头", sheet_keyword="another"),
        _map(target="列头", match_mode="包含", original="Group_Amount", standard="Standard"),
        _map(target="列头", original="Other_Count", standard="Count standard"),
    ]
    cache, ordered = engine.compute_wide_col_paths_from_arrays(
        df, (), rule, "Assets", maps, "report.xlsx", c_start=1, c_end=3, ch_rows=[0, 1]
    )
    assert ordered == ["Standard_1", "Standard_2", "Count standard"]
    assert cache == {1: "Standard_1", 2: "Standard_2", 3: "Count standard"}

    df_empty = pd.DataFrame([["", "", ""], ["row", 1, 2]])
    _, fallback = engine.compute_wide_col_paths_from_arrays(
        df_empty, (), rule, "Assets", [], "report.xlsx", c_start=1, c_end=2, ch_rows=[0]
    )
    assert fallback == ["列_B", "列_C"]
    _, merged_no_fill = engine.compute_wide_col_paths_from_arrays(
        df, ((0, 0, 1, 2),), rule, "Assets", [], "report.xlsx", c_start=1, c_end=3, ch_rows=[0, 1]
    )
    assert merged_no_fill == ["Group_Amount_1", "Group_Amount_2", "Other_Count"]
    overrides = {(0, 1): "GroupOverride", (1, 3): "CountOverride"}
    _, over = engine.compute_wide_col_paths_from_arrays(
        df, (), rule, "Assets", [], "report.xlsx", c_start=1, c_end=3, ch_rows=[0, 1],
        cell_text_overrides=overrides,
    )
    assert over == ["GroupOverride_Amount_1", "GroupOverride_Amount_2", "Other_CountOverride"]


def test_extract_cells_handles_paths_dates_filters_and_boundaries():
    df = pd.DataFrame([
        ["", "", "", ""],
        ["指标", "存款", "存款", "贷款"],
        ["单位", "本币", "本币", "外币"],
        ["甲行", 10, 20, None],
        ["甲行", 30, 40, 50],
        ["停用甲行", 1, 2, 3],
    ])
    rule = _rule(
        col_header_rows=[2, 3], required_col_headers=["存款"], required_row_headers=["甲行"],
        data_row_start=4, data_row_end=6, data_col_start=2, data_col_end=4, skip_keywords=["停用"],
    )
    maps = [
        _map(target="行头", match_mode="包含", original="甲行", standard="标准甲", applicable_rule="rule"),
        _map(target="列头", match_mode="包含", original="存款", standard="存款余额", applicable_rule="rule"),
        _map(target="行头", applicable_rule="other"),
        _map(target="行头", wb_keyword="other workbook"),
        _map(target="行头", sheet_keyword="other sheet"),
    ]
    cells = engine.extract_cells_from_arrays(df, (), Path("report_20260131.xlsx"), rule, "Assets", maps, True)
    assert len(cells) == 5
    assert [c.row_path for c in cells] == ["标准甲_1"] * 2 + ["标准甲_2"] * 3
    assert [c.col_path for c in cells[:2]] == ["存款余额_本币_1", "存款余额_本币_2"]
    assert cells[0].data_date == "2026-01-31"
    assert cells[0].cell_addr == "B4" and cells[-1].cell_addr == "D5"

    exact_rows = engine.extract_cells_from_arrays(
        df, (), Path("report_20260131.xlsx"), rule, "Assets",
        [_map(target="行头", original="甲行_1", standard="Exact bank")], True,
    )
    assert exact_rows[0].row_path == "Exact bank"

    blank_column_cells = engine.extract_cells_from_arrays(
        df, (), Path("report_20260131.xlsx"), rule, "Assets",
        maps + [_map(target="列头", original="贷款_外币", standard="")], True,
    )
    assert len(blank_column_cells) == 4
    assert all(cell.col_path != "" for cell in blank_column_cells)

    merged_blank_header = pd.DataFrame([["202601", "", ""], ["", None, None], ["甲行", 1, 2]])
    merged_blank_result = engine.extract_cells_from_arrays(
        merged_blank_header, ((1, 1, 1, 2),), Path("x.xlsx"),
        _rule(col_header_rows=[2], data_row_start=3, data_col_start=2), "Merged", [], False,
    )
    assert [c.col_path for c in merged_blank_result] == ["列_B", "列_C"]

    left_fill_df = pd.DataFrame([
        ["202601", "", ""], ["指标", "存款", None], ["甲行", 1, 2],
    ])
    left_fill = engine.extract_cells_from_arrays(
        left_fill_df, (), Path("x.xlsx"),
        _rule(col_header_rows=[2], data_row_start=3, data_col_start=2), "Fill", [], False,
    )
    assert [c.col_path for c in left_fill] == ["存款_1", "存款_2"]

    override_df = pd.DataFrame([
        ["202601", "", ""], ["指标", "存款", None], ["甲行", 1, 2],
    ])
    override_result = engine.extract_cells_from_arrays(
        override_df, (), Path("x.xlsx"),
        _rule(col_header_rows=[2], data_row_start=3, data_col_start=2), "Fill", [], False,
        cell_text_overrides={(1, 1): "列覆盖", (2, 0): "行覆盖"},
    )
    assert {c.row_path for c in override_result} == {"行覆盖"}
    assert {c.col_path for c in override_result} == {"列覆盖_1", "列覆盖_2"}

    blank_rows = pd.DataFrame([["date", "header"], ["", "value"], ["", None]])
    assert engine.extract_cells_from_arrays(
        blank_rows, (), Path("x.xlsx"), _rule(col_header_rows=[2], data_row_start=3, data_col_start=2),
        "Blank", [], False,
    ) == []

    assert engine.extract_cells_from_arrays(pd.DataFrame(), (), Path("x.xlsx"), rule, "Assets", [], False) == []
    assert engine.extract_cells_from_arrays(
        df, (), Path("x.xlsx"), _rule(row_header_col=0, row_header_cols=[]), "Assets", [], False
    ) == []
    assert engine.extract_cells_from_arrays(
        df, (), Path("x.xlsx"), _rule(row_header_cols=[0]), "Assets", [], False
    ) == []
    assert engine.extract_cells_from_arrays(df, (), Path("x.xlsx"), _rule(row_header_cols=[99]), "Assets", [], False) == []
    assert engine.extract_cells_from_arrays(df, (), Path("x.xlsx"), _rule(col_header_rows=[99]), "Assets", [], False) == []
    assert engine.extract_cells_from_arrays(df, (), Path("x.xlsx"), _rule(data_row_start=9, data_row_end=3), "Assets", [], False) == []
    assert engine.extract_cells_from_arrays(df, (), Path("x.xlsx"), _rule(data_col_start=9, data_col_end=2), "Assets", [], False) == []
    assert engine.extract_cells_from_arrays(
        df, (), Path("x.xlsx"), _rule(required_row_headers=["missing"]), "Assets", [], False
    ) == []
    assert engine.extract_cells_from_arrays(
        df, (), Path("x.xlsx"), _rule(required_col_headers=["missing"]), "Assets", [], False
    ) == []

    wide_rule = _rule(
        row_header_col_specified=False, row_header_cols=[], col_header_rows=[2],
        data_row_start=4, data_row_end=5, data_col_start=2, data_col_end=3,
        required_row_headers=["ignored"],
    )
    wide = engine.extract_cells_from_arrays(
        df, (), Path("report_without_date.xlsx"), wide_rule, "Assets", [], False,
        cell_text_overrides={(1, 1): "OVERRIDE", (3, 0): "ignored row header"},
    )
    assert [c.row_path for c in wide] == ["ROW#4", "ROW#4", "ROW#5", "ROW#5"]
    assert [c.col_path for c in wide] == ["OVERRIDE", "存款", "OVERRIDE", "存款"]
    assert all(c.data_date == "" for c in wide)


def test_merge_ranges_file_openpyxl_cache_and_excel_fallback(tmp_path, monkeypatch):
    path = tmp_path / "merged.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Merged"
    ws.merge_cells("B2:C3")
    wb.create_sheet("Empty")
    wb.save(path)
    wb.close()
    engine._load_merge_ranges_file.cache_clear()
    assert engine._load_merge_ranges_file(str(path)) == {
        "Merged": ((1, 2, 1, 2),), "Empty": ()
    }
    assert engine._load_merge_ranges(str(path), "Missing") == ()

    engine._load_merge_ranges_file.cache_clear()
    xls = tmp_path / "old.xls"
    assert engine._load_merge_ranges_file(str(xls)) == {}


class _Collection(list):
    @property
    def Count(self):
        return len(self)


class _FakeRange:
    def __init__(self, *, row=1, column=1, row_count=1, col_count=1, areas=None, rows=(), fail_areas=False):
        self.Row, self.Column = row, column
        self.Rows = _Collection(rows or [None] * row_count)
        self.Columns = _Collection([None] * col_count)
        self._areas, self._rows, self._fail_areas = areas or [], rows, fail_areas

    @property
    def MergeAreas(self):
        if self._fail_areas:
            raise RuntimeError("MergeAreas unavailable")
        return self._areas

    @property
    def Rows_iter(self):
        return self._rows

    def __iter__(self):
        return iter(self._rows)


class _FakeCell:
    def __init__(self, merge=None, *, fail=False):
        self._merge, self.fail = merge, fail

    @property
    def MergeCells(self):
        if self.fail:
            raise RuntimeError("cell unavailable")
        return self._merge is not None

    @property
    def MergeArea(self):
        return self._merge


def _install_fake_win32com(monkeypatch, dispatch):
    win32com = types.ModuleType("win32com")
    win32com.__path__ = []
    client = types.ModuleType("win32com.client")
    client.DispatchEx = dispatch
    win32com.client = client
    monkeypatch.setitem(sys.modules, "win32com", win32com)
    monkeypatch.setitem(sys.modules, "win32com.client", client)


def test_com_merge_loader_uses_merged_areas_and_closes_resources(monkeypatch):
    ma = _FakeRange(row=2, column=3, row_count=2, col_count=4)
    used = _FakeRange(areas=[ma])
    ws = SimpleNamespace(Name="Data", UsedRange=used)

    class FakeBook:
        Worksheets = [ws]
        closed = False

        def Close(self, **_kwargs):
            self.closed = True
            raise RuntimeError("close failed")

    class FakeApp:
        Visible = True
        DisplayAlerts = True
        quit = False

        def __init__(self):
            self.Workbooks = SimpleNamespace(Open=lambda *_args, **_kwargs: book)

        def Quit(self):
            self.quit = True
            raise RuntimeError("quit failed")

    book, app = FakeBook(), None
    def dispatch(_prog_id):
        nonlocal app
        app = FakeApp()
        return app

    _install_fake_win32com(monkeypatch, dispatch)
    got = engine._load_merge_ranges_file_by_com("book.xls")
    assert got == {"Data": ((1, 2, 2, 5),)}
    assert book.closed and app.quit


def test_com_merge_loader_falls_back_to_cell_scan_and_handles_failures(monkeypatch):
    merged = _FakeRange(row=2, column=2, row_count=2, col_count=2)
    cells = [_FakeCell(merged), _FakeCell(), _FakeCell(fail=True)]
    row = SimpleNamespace(Cells=cells)
    used = _FakeRange(rows=[row], fail_areas=True)
    ws = SimpleNamespace(Name="Fallback", UsedRange=used)
    book = SimpleNamespace(Worksheets=[ws], Close=lambda **_kw: None)
    app = SimpleNamespace(
        Workbooks=SimpleNamespace(Open=lambda *_args, **_kwargs: book),
        Quit=lambda: None,
    )
    _install_fake_win32com(monkeypatch, lambda _prog_id: app)
    assert engine._load_merge_ranges_file_by_com("book.xls") == {"Fallback": ((1, 2, 1, 2),)}

    _install_fake_win32com(monkeypatch, lambda _prog_id: (_ for _ in ()).throw(RuntimeError("no app")))
    assert engine._load_merge_ranges_file_by_com("book.xls") == {}


def test_com_merge_loader_returns_empty_when_pywin32_import_fails(monkeypatch):
    original_import = builtins.__import__

    def no_com(name, *args, **kwargs):
        if name.startswith("win32com"):
            raise ImportError("no COM runtime")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", no_com)
    assert engine._load_merge_ranges_file_by_com("book.xls") == {}


def test_iter_matching_sheets_and_excel_reader_dispatch(monkeypatch, tmp_path):
    rules = [_rule(sheet_keyword="data"), _rule(sheet_keyword="missing")]
    monkeypatch.setattr(engine, "list_sheet_names", lambda _path: ["Data sheet", "Other"])
    assert [(rule.sheet_keyword, name) for rule, name in engine.iter_matching_sheets(tmp_path / "src.xlsx", rules)] == [
        ("data", "Data sheet")
    ]
    df = pd.DataFrame([["x"]])
    monkeypatch.setattr(engine, "read_sheet_2d", lambda *_args: df)
    monkeypatch.setattr(engine, "_load_merge_ranges", lambda *_args: ())
    monkeypatch.setattr(engine, "extract_cells_from_arrays", lambda *args: [(args[0], args[2])])
    result = engine.extract_cells(tmp_path / "src.xlsx", _rule(), "Data", [], False)
    assert result == [(df, tmp_path / "src.xlsx")]
