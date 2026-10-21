from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from PyCode.core.summary_321_322 import run_summary_by_comment, run_summary_by_usedrange


def _summary_template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "模板"
    for address, value, text in (
        ("A1", "行开始", "行区域1"),
        ("A3", "行结束", "行区域#1"),
        ("B1", "指标", "列区域1"),
        ("B2", "数值", "列区域#1"),
    ):
        ws[address] = value
        ws[address].comment = Comment(text, "test")
    wb.save(path)
    wb.close()
    return path


def _summary_source(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "回复"
    ws.append(["单位", "指标值"])
    ws.append(["甲", 10])
    ws.append(["乙", 20])
    short = wb.create_sheet("短表")
    short["A1"] = "只有一列"
    wb.save(path)
    wb.close()
    return path


def test_usedrange_summary_accepts_valid_files_and_skips_missing_inputs(tmp_path: Path) -> None:
    source = _summary_source(tmp_path / "source.xlsx")
    missing = tmp_path / "missing.xlsx"
    unsupported = tmp_path / "ignored.csv"
    unsupported.touch()

    result = run_summary_by_usedrange(
        [missing, unsupported, source], tmp_path / "out", tmp_path / "logs"
    )

    assert result["saved"] is True and result["rows"] == 4
    wb = load_workbook(result["path"], data_only=True)
    try:
        ws = wb["汇总"]
        assert ws.max_column == 5
        assert ws["A2"].value == "source" and ws["B2"].value == "回复"
        assert ws["E3"].value == 10
        assert ws["E5"].value is None
    finally:
        wb.close()


def test_comment_summary_uses_template_fallback_for_unmatched_sheet(tmp_path: Path) -> None:
    template = _summary_template(tmp_path / "template.xlsx")
    source = _summary_source(tmp_path / "source.xlsx")

    result = run_summary_by_comment(
        template,
        [tmp_path / "missing.xlsx", tmp_path / "ignored.csv", source],
        tmp_path / "out",
        tmp_path / "logs",
    )

    assert result["saved"] is True and result["rows"] == 6
    wb = load_workbook(result["path"], data_only=True)
    try:
        ws = wb["模板"]
        assert [ws.cell(1, col).value for col in range(1, 5)] == [
            "工作簿", "工作表", "指标_数值", "行号",
        ]
        assert [ws.cell(2, col).value for col in range(1, 5)] == [
            "source", "回复", "指标值", 1,
        ]
        assert ws["C3"].value == 10
        assert ws["A5"].value == "source" and ws["B5"].value == "短表"
    finally:
        wb.close()
