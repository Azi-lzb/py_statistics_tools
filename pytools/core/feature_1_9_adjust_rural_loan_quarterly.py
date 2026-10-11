from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from .feature_1_8_adjust_rural_loan import (
    HEADER_SCAN_ROWS,
    HEADER_TEXT,
    INST_DATA_START_ROW,
    INST_DIFF_COL,
    INST_NAME_COL,
    INST_VALUE_COL,
    NAME_SCAN_COLS,
    TOLERANCE,
    _normalize,
    _to_num,
)
from .logger import get_logger


INSTITUTION_SHEET_KEY = "涉农贷款分机构"
AREA_SHEET_KEY = "涉农贷款分地区统计"
_REGION_SUFFIX_RE = re.compile(r"[\s\-—_－]+$")


def _last_row(ws) -> int:
    for row_no in range(ws.max_row or 0, 0, -1):
        if any(ws.cell(row=row_no, column=col).value not in (None, "") for col in range(1, (ws.max_column or 0) + 1)):
            return row_no
    return 0


def _last_col(ws) -> int:
    for col_no in range(ws.max_column or 0, 0, -1):
        if any(ws.cell(row=row, column=col_no).value not in (None, "") for row in range(1, (ws.max_row or 0) + 1)):
            return col_no
    return 0


def _region_from_sheet_name(title: str) -> str:
    position = title.find(INSTITUTION_SHEET_KEY)
    if position < 0:
        return ""
    return _REGION_SUFFIX_RE.sub("", _normalize(title[:position]))


def _institution_sheets(wb) -> dict[str, object]:
    sheets: dict[str, object] = {}
    for ws in wb.worksheets:
        if INSTITUTION_SHEET_KEY not in ws.title:
            continue
        region = _region_from_sheet_name(ws.title)
        if not region:
            raise ValueError(f"无法从分机构Sheet名识别地区：{ws.title}")
        if region in sheets:
            raise ValueError(f"同一工作簿存在重复地区分机构Sheet：{region}")
        sheets[region] = ws
    return sheets


def _find_area_sheet(wb):
    hits = [ws for ws in wb.worksheets if AREA_SHEET_KEY in ws.title]
    if len(hits) > 1:
        raise ValueError(f"匹配到多个“{AREA_SHEET_KEY}”工作表，请保留唯一目标表。")
    return hits[0] if hits else None


def _build_institution_map(ws) -> tuple[dict[str, float], dict[str, int], list[str]]:
    values: dict[str, float] = {}
    rows: dict[str, int] = {}
    duplicates: list[str] = []
    for row_no in range(INST_DATA_START_ROW, _last_row(ws) + 1):
        name = _normalize(ws.cell(row=row_no, column=INST_NAME_COL).value)
        if not name:
            continue
        if name in values:
            duplicates.append(name)
            continue
        values[name] = _to_num(ws.cell(row=row_no, column=INST_VALUE_COL).value)
        rows[name] = row_no
    return values, rows, duplicates


def _find_header_col(ws, header: str) -> int:
    for row_no in range(1, min(_last_row(ws), HEADER_SCAN_ROWS) + 1):
        for col_no in range(1, _last_col(ws) + 1):
            if _normalize(ws.cell(row=row_no, column=col_no).value) == header:
                return col_no
    return 0


def _find_region_row(ws, region: str) -> int:
    for row_no in range(1, _last_row(ws) + 1):
        for col_no in range(1, NAME_SCAN_COLS + 1):
            if _normalize(ws.cell(row=row_no, column=col_no).value) == region:
                return row_no
    return 0


def run_adjust_rural_loan_quarterly(current_path: Path, previous_path: Path, log_dir: Path) -> dict:
    """修正全部地区涉农贷款分机构表，并逐地区核对分地区统计表。"""
    logger = get_logger("feature_1_9_adjust_rural_loan_quarterly", log_dir)
    result = {
        "regions": [], "matched": 0, "fill_total": 0.0, "area_total": 0.0,
        "consistent": False, "error": None,
    }
    keep_vba = current_path.suffix.lower() == ".xlsm"
    wb_current = load_workbook(current_path, keep_vba=keep_vba)
    wb_previous = load_workbook(previous_path, data_only=True, read_only=True)
    try:
        current_sheets = _institution_sheets(wb_current)
        previous_sheets = _institution_sheets(wb_previous)
        if not current_sheets:
            raise ValueError("本期文件未找到包含“涉农贷款分机构”的工作表。")
        if not previous_sheets:
            raise ValueError("上期文件未找到包含“涉农贷款分机构”的工作表。")
        area_sheet = _find_area_sheet(wb_current)
        if area_sheet is None:
            raise ValueError("本期文件未找到包含“涉农贷款分地区统计”的工作表。")
        header_col = _find_header_col(area_sheet, HEADER_TEXT)
        if not header_col:
            raise ValueError(f'未在“{area_sheet.title}”中找到表头“{HEADER_TEXT}”。')

        for region, current_sheet in current_sheets.items():
            region_result = {
                "region": region, "sheet": current_sheet.title, "matched": 0,
                "fill_total": 0.0, "area_total": None, "consistent": False,
                "unmatched_current": [], "previous_only": [], "error": None,
            }
            previous_sheet = previous_sheets.get(region)
            if previous_sheet is None:
                region_result["error"] = f"上期文件缺少{region}分机构表"
                result["regions"].append(region_result)
                continue

            current_values, current_rows, current_duplicates = _build_institution_map(current_sheet)
            previous_values, _, previous_duplicates = _build_institution_map(previous_sheet)
            if current_duplicates or previous_duplicates:
                duplicates = current_duplicates + previous_duplicates
                region_result["error"] = "分机构表存在重复机构名称：" + "、".join(duplicates)
                result["regions"].append(region_result)
                continue

            for name, row_no in current_rows.items():
                if name not in previous_values:
                    current_sheet.cell(row=row_no, column=INST_DIFF_COL, value=None)
                    region_result["unmatched_current"].append(name)
                    continue
                difference = current_values[name] - previous_values[name]
                current_sheet.cell(row=row_no, column=INST_DIFF_COL, value=difference)
                region_result["matched"] += 1
                region_result["fill_total"] += difference

            region_result["previous_only"] = [name for name in previous_values if name not in current_values]
            area_row = _find_region_row(area_sheet, region)
            if not area_row:
                region_result["error"] = f'“{area_sheet.title}”未找到地区行：{region}'
            else:
                region_result["area_total"] = _to_num(area_sheet.cell(row=area_row, column=header_col).value)
                region_result["consistent"] = abs(region_result["fill_total"] - region_result["area_total"]) <= TOLERANCE

            result["matched"] += region_result["matched"]
            result["fill_total"] += region_result["fill_total"]
            if region_result["area_total"] is not None:
                result["area_total"] += region_result["area_total"]
            result["regions"].append(region_result)

        wb_current.save(current_path)
        result["consistent"] = bool(result["regions"]) and all(
            item["consistent"] and not item["error"] for item in result["regions"]
        )
        logger.info(
            "地区=%s 匹配机构=%s 分机构合计=%.4f 分地区合计=%.4f %s",
            len(result["regions"]), result["matched"], result["fill_total"], result["area_total"],
            "一致" if result["consistent"] else "存在不一致",
        )
        for item in result["regions"]:
            if item["error"]:
                logger.warning("%s: %s", item["region"], item["error"])
    except Exception as exc:
        result["error"] = str(exc)
        logger.error("失败: %s", exc)
    finally:
        wb_previous.close()
        wb_current.close()
    return result
